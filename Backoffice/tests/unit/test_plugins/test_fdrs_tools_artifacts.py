"""Tests for durable FDRS Tools Excel artifacts."""

from __future__ import annotations

import json

import pytest

from plugins.fdrs.services import fdrs_tools_artifacts as artifacts

pytestmark = [pytest.mark.unit]


def test_persist_and_load_verify_latest(app, tmp_path):
    import openpyxl

    path = tmp_path / "verify.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data_points"
    ws.append(["year", "ISO3", "status"])
    ws.append([2024, "KEN", "matched"])
    ws.append([2024, "UGA", "mismatch"])
    wb.save(path)
    wb.close()

    with app.app_context():
        stored = artifacts.persist_verify_latest(
            template_id=999021,
            local_xlsx_path=str(path),
            stats={"total": 2, "matched": 1, "mismatch": 1},
            extra={"fdrs_years": [2024]},
        )
        assert stored["template_id"] == 999021
        assert stored["stats"]["matched"] == 1
        assert stored["completed_at"]
        assert artifacts.verify_latest_exists(999021)

        meta = artifacts.load_verify_latest_meta(999021)
        assert meta["fdrs_years"] == [2024]
        assert meta["kind"] == "fdrs.sync_verify"

        raw = artifacts.download_bytes(artifacts.verify_xlsx_rel(999021))
        table = artifacts.read_xlsx_sheet(
            raw,
            sheet="data_points",
            aliases={"data_points": "data_points"},
            filters={"status": "mismatch"},
        )
        assert table["total_rows"] == 2
        assert table["filtered_rows"] == 1
        assert table["rows"][0]["ISO3"] == "UGA"


def test_persist_documents_latest_and_search(app, tmp_path):
    import openpyxl

    path = tmp_path / "docs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "documents"
    ws.append(["iso3", "year", "document_type", "downloadable", "is_public", "http_status", "name"])
    ws.append(["KEN", 2024, "Annual Report", True, True, 200, "Kenya AR"])
    ws.append(["UGA", 2023, "Audited Financial Statement", False, False, 403, "Uganda AFS"])
    wb.create_sheet("summary").append(["http_status", "count"])
    wb.save(path)
    wb.close()

    with app.app_context():
        artifacts.persist_documents_latest(
            local_xlsx_path=str(path),
            stats={"total_documents": 2, "downloadable_count": 1},
        )
        assert artifacts.documents_latest_exists()
        raw = artifacts.download_bytes(artifacts.documents_xlsx_rel())
        kenya = artifacts.read_xlsx_sheet(
            raw,
            sheet="documents",
            search="ken",
            search_columns=["iso3", "name"],
        )
        assert kenya["filtered_rows"] == 1
        assert kenya["rows"][0]["iso3"] == "KEN"

        blocked = artifacts.read_xlsx_sheet(
            raw,
            sheet="documents",
            filters={"downloadable": "false", "http_status": "403"},
        )
        assert blocked["filtered_rows"] == 1
        assert blocked["rows"][0]["iso3"] == "UGA"

        all_rows = artifacts.read_xlsx_sheet(raw, sheet="documents", per_page=0)
        assert all_rows["filtered_rows"] == 2
        assert len(all_rows["rows"]) == 2
        assert all_rows["page"] == 1


def test_load_missing_meta_returns_none(app):
    with app.app_context():
        assert artifacts.load_verify_latest_meta(999001) is None
        # documents latest may exist from a prior test in this process; missing
        # verify for an unused template id is the contract we care about here.
        sidecar = artifacts.load_json(artifacts.verify_meta_rel(999001))
        assert sidecar is None


def test_persist_json_roundtrip(app):
    with app.app_context():
        artifacts.persist_json("scratch/meta.json", {"ok": True, "n": 3})
        loaded = artifacts.load_json("scratch/meta.json")
        assert loaded == {"ok": True, "n": 3}
        raw = artifacts.download_bytes("scratch/meta.json")
        assert json.loads(raw.decode("utf-8"))["n"] == 3
