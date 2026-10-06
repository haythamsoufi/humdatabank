"""Read and write the IFRC FDRS data-collection workbook (template 21).

Question labels are Excel formulas over the Translations sheet. Input cells are
the blank cells next to those labels. Sex-age grids use the FDRS age bands,
which are summed into the form's age groups on import.
"""

from __future__ import annotations

import re
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

import base64
import json
import zlib

from app.services.forms.processing_service import slugify_age_group

# Same labels as FDRS_INCOME_KPI_TO_MATRIX_ROW. Kept here so workbook parsing
# does not import the FDRS sync script path.
INCOME_ROW_LABELS: Tuple[str, ...] = (
    "Home Government",
    "Foreign Government",
    "Individuals",
    "Corporations",
    "Foundations",
    "UN Agencies & other Multilateral Agencies",
    "Pooled funds",
    "Non-governmental organizations",
    "Service income",
    "Income generating activity",
    "Other National Society",
    "IFRC",
    "ICRC",
    "Other",
)

_NEW_AGE_GROUPS_B64 = (
    "XY1LCoAwDETvknULSf89hQcoZVaCa/H+mFaE6ibJe0yY1sgzfCVDUmyo1E3TC+5r+BHRSp4cGXEKxuTACP+PqKxjQGKkNb5dx34qP3uYzMhrQjy0amlMELdy4Tfdbw=="
)

INDICATOR_SHEET = "Indicator Data"
TRANSLATION_SHEET = "Translations"
YEAR_CELL = "E10"
NS_NAME_CELL = "E9"
INCOME_HEADER_ROW = 112
INCOME_VALUE_ROW = 113

# English question text (or a distinctive fragment) -> indicator bank KPI.
SEX_AGE_QUESTIONS: Tuple[Tuple[str, str], ...] = (
    ("Number of people on the National Society Governing Board", "KPI_GB"),
    ("Number of people volunteering their time", "KPI_PeopleVol"),
    ("Number of volunteers covered by accident insurance", "KPI_noVolCoveredAI"),
    ("Number of paid staff", "KPI_PStaff"),
    ("Number of staff covered by accident insurance", "KPI_PStaffCoveredAI"),
    ("People donating blood", "KPI_DonBlood"),
    ("People trained in First Aid", "KPI_TrainFA"),
    ("disaster response and early recovery", "KPI_ReachDRER"),
    ("long term services and development", "KPI_ReachLTSPD"),
    ("disaster risk reduction", "KPI_ReachDRR"),
    ("People reached by shelter", "KPI_ReachS"),
    ("People reached by livelihoods", "KPI_ReachL"),
    ("People reached by health", "KPI_ReachH"),
    ("mental health and psychosocial", "KPI_ReachHPM"),
    ("immunisation", "KPI_ReachHI"),
    ("water, sanitation and hygiene", "KPI_ReachWASH"),
    ("Migrants and displaced persons reached", "KPI_ReachM"),
    ("climate and environmental activities", "KPI_Climate"),
    ("heatwave", "KPI_ClimateHeat"),
    ("cash transfer programming", "KPI_ReachCTP"),
    ("protection, gender and inclusion", "KPI_ReachSI"),
    ("educational programmes", "KPI_ReachRCRCEd"),
)

SCALAR_QUESTIONS: Tuple[Tuple[str, str], ...] = (
    ("Number of branches", "KPI_noBranches"),
    ("Number of local units", "KPI_noLocalUnits"),
    ("Reporting currency", "KPI_CUR_Code"),
    ("Financial reporting start date", "KPI_StartDate"),
    ("Financial reporting end date", "KPI_EndDate"),
)

SEX_KEYS = {
    "male": "male",
    "female": "female",
    "non-binary": "non_binary",
    "non binary": "non_binary",
    "unknown": "unknown",
}

_MATCH_ID = re.compile(r"MATCH\(\s*(\d+)")
_CELL_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _decode_new_age_groups() -> Dict[str, str]:
    raw = base64.b64decode(_NEW_AGE_GROUPS_B64)
    data = json.loads(zlib.decompress(raw, -zlib.MAX_WBITS).decode("utf-8"))
    return {str(row[0]).strip(): str(row[1]).strip() for row in data if len(row) >= 2}


_NEW_AGE_GROUPS = _decode_new_age_groups()


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _is_formula(value: Any) -> bool:
    if isinstance(value, str) and value.startswith("="):
        return True
    return type(value).__name__ == "ArrayFormula"


def translation_labels(wb) -> Dict[int, str]:
    ws = wb[TRANSLATION_SHEET]
    labels: Dict[int, str] = {}
    for row in range(2, (ws.max_row or 1) + 1):
        key = ws.cell(row, 1).value
        english = ws.cell(row, 5).value
        if key is None or not english:
            continue
        try:
            labels[int(key)] = str(english).replace("\n", " ").strip()
        except (TypeError, ValueError):
            continue
    return labels


def _label_at(ws, cell, labels: Mapping[int, str]) -> str:
    value = cell.value
    if not isinstance(value, str):
        return ""
    match = _MATCH_ID.search(value)
    if not match:
        return ""
    return labels.get(int(match.group(1)), "")


def _sheet_labels(ws, labels: Mapping[int, str]) -> Dict[Tuple[int, int], str]:
    found: Dict[Tuple[int, int], str] = {}
    for row in ws.iter_rows(max_row=ws.max_row or 1, max_col=min(ws.max_column or 1, 22)):
        for cell in row:
            text = _label_at(ws, cell, labels)
            if text:
                found[(cell.row, cell.column)] = text
    return found


def _find_label(cells: Mapping[Tuple[int, int], str], fragment: str) -> Optional[Tuple[int, int]]:
    needle = _norm(fragment)
    for (row, col), text in cells.items():
        if needle in _norm(text):
            return row, col
    return None


def _input_cell_for_label(ws, row: int, col: int):
    """Return the blank or typed cell a Required-check formula points at."""
    for scan_row in range(max(1, row - 2), row + 2):
        for scan_col in range(1, min(ws.max_column or 1, 22) + 1):
            formula = ws.cell(scan_row, scan_col).value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            for column, ref_row in _CELL_REF.findall(formula):
                ref_row_i = int(ref_row)
                if ref_row_i != row:
                    continue
                from openpyxl.utils import column_index_from_string

                ref_col = column_index_from_string(column)
                if ref_col == col:
                    continue
                target = ws.cell(row, ref_col)
                if not _is_formula(target.value):
                    return target
    return None


def _age_token(header: str) -> Optional[str]:
    text = str(header or "").strip().lower().replace("≤", "0-").replace("–", "-").replace("—", "-")
    text = text.replace("6t-", "6-").replace("+", "")
    text = re.sub(r"\s+", "", text)
    if not text or not re.search(r"\d", text):
        return None
    token = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    if token == "80_":
        return "80"
    return token or None


def _form_age_slug(header: str) -> Optional[str]:
    token = _age_token(header)
    if not token:
        return None
    mapped = _NEW_AGE_GROUPS.get(token)
    if not mapped:
        return None
    if mapped.lower() == "other":
        return "unknown"
    return slugify_age_group(mapped)


def _sex_key(label: str) -> Optional[str]:
    return SEX_KEYS.get(_norm(label).replace(" ", " ").strip()) or SEX_KEYS.get(_norm(label))


def _number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value is None or _is_formula(value):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            value = float(text)
        except ValueError:
            return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _json_number(value: float) -> int | float:
    if float(value).is_integer():
        return int(value)
    return float(value)


def _read_sex_age_grid(ws, labels: Mapping[int, str], start_row: int) -> Dict[str, Any]:
    header_row = None
    age_columns: List[Tuple[int, str]] = []
    for row in range(start_row, min(start_row + 8, (ws.max_row or start_row) + 1)):
        columns = []
        for col in range(4, min(ws.max_column or 1, 20) + 1):
            value = ws.cell(row, col).value
            if _is_formula(value) or value is None:
                continue
            slug = _form_age_slug(str(value))
            if slug:
                columns.append((col, slug))
        if len(columns) >= 3:
            header_row = row
            age_columns = columns
            break
    direct: Dict[str, int] = {}
    if header_row is None:
        return {"mode": "sex_age", "values": {"direct": direct, "indirect": None}}
    for row in range(header_row + 1, header_row + 8):
        sex_label = _label_at(ws, ws.cell(row, 2), labels)
        sex = _sex_key(sex_label)
        if not sex:
            if _norm(sex_label) == "total":
                break
            continue
        for col, age_slug in age_columns:
            number = _number(ws.cell(row, col).value)
            if number is None:
                continue
            key = f"{sex}_{age_slug}"
            direct[key] = direct.get(key, 0) + int(number) if float(number).is_integer() else direct.get(key, 0) + number
    return {"mode": "sex_age", "values": {"direct": direct, "indirect": None}}


def _write_sex_age_grid(ws, labels: Mapping[int, str], start_row: int, disagg: Mapping[str, Any]) -> None:
    values = (disagg or {}).get("values") or {}
    direct = values.get("direct") or {}
    if not isinstance(direct, dict):
        return
    header_row = None
    age_columns: List[Tuple[int, str]] = []
    for row in range(start_row, min(start_row + 8, (ws.max_row or start_row) + 1)):
        columns = []
        for col in range(4, min(ws.max_column or 1, 20) + 1):
            value = ws.cell(row, col).value
            if _is_formula(value) or value is None:
                continue
            slug = _form_age_slug(str(value))
            if slug:
                columns.append((col, slug))
        if len(columns) >= 3:
            header_row = row
            age_columns = columns
            break
    if header_row is None:
        return
    for row in range(header_row + 1, header_row + 8):
        sex = _sex_key(_label_at(ws, ws.cell(row, 2), labels))
        if not sex:
            continue
        # Several workbook bands collapse to one form band. Write the form
        # value once, into the first matching column of this sex row.
        written_ages = set()
        for col, age_slug in age_columns:
            if age_slug in written_ages:
                continue
            key = f"{sex}_{age_slug}"
            if key not in direct:
                continue
            ws.cell(row, col).value = direct[key]
            written_ages.add(age_slug)


def _income_label(text: str) -> Optional[str]:
    normalized = _norm(text).replace("goverment", "government")
    for label in INCOME_ROW_LABELS:
        if _norm(label) == normalized:
            return label
    return None


def read_fdrs_workbook(wb) -> Dict[str, Any]:
    ws = wb[INDICATOR_SHEET]
    labels = translation_labels(wb)
    placed = _sheet_labels(ws, labels)
    scalars: Dict[str, Any] = {}
    disagg: Dict[str, Any] = {}
    warnings: List[str] = []

    for fragment, kpi in SCALAR_QUESTIONS:
        located = _find_label(placed, fragment)
        if not located:
            warnings.append(f"Could not find the {fragment} question in the FDRS workbook.")
            continue
        target = _input_cell_for_label(ws, located[0], located[1])
        if target is None:
            continue
        if kpi in {"KPI_StartDate", "KPI_EndDate"} and target.value not in (None, ""):
            scalars[kpi] = target.value.isoformat()[:10] if hasattr(target.value, "isoformat") else str(target.value)
            continue
        if kpi == "KPI_CUR_Code" and target.value not in (None, ""):
            scalars[kpi] = str(target.value).strip()
            continue
        number = _number(target.value)
        if number is not None:
            scalars[kpi] = _json_number(number)

    for fragment, kpi in SEX_AGE_QUESTIONS:
        located = _find_label(placed, fragment)
        if not located:
            warnings.append(f"Could not find the {fragment} question in the FDRS workbook.")
            continue
        payload = _read_sex_age_grid(ws, labels, located[0])
        if (payload.get("values") or {}).get("direct"):
            disagg[kpi] = payload

    income: Dict[str, Any] = {}
    for col in range(1, 21):
        label = _income_label(_label_at(ws, ws.cell(INCOME_HEADER_ROW, col), labels))
        if not label:
            continue
        number = _number(ws.cell(INCOME_VALUE_ROW, col).value)
        if number is not None:
            income[label] = _json_number(number)

    given = _read_support_table(ws, labels, "National Societies to which support was given")
    received = _read_support_table(ws, labels, "National Societies from which support was received")
    return {
        "scalars": scalars,
        "disagg": disagg,
        "income": income,
        "support_given": given,
        "support_received": received,
        "warnings": warnings,
    }


def write_fdrs_answers(wb, answers: Mapping[str, Any]) -> None:
    ws = wb[INDICATOR_SHEET]
    labels = translation_labels(wb)
    placed = _sheet_labels(ws, labels)
    scalars = answers.get("scalars") or {}
    disagg = answers.get("disagg") or {}
    for fragment, kpi in SCALAR_QUESTIONS:
        if kpi not in scalars:
            continue
        located = _find_label(placed, fragment)
        if not located:
            continue
        target = _input_cell_for_label(ws, located[0], located[1])
        if target is not None:
            target.value = scalars[kpi]
    for fragment, kpi in SEX_AGE_QUESTIONS:
        if kpi not in disagg:
            continue
        located = _find_label(placed, fragment)
        if located:
            _write_sex_age_grid(ws, labels, located[0], disagg[kpi])
    income = answers.get("income") or {}
    for col in range(1, 21):
        label = _income_label(_label_at(ws, ws.cell(INCOME_HEADER_ROW, col), labels))
        if label and label in income:
            ws.cell(INCOME_VALUE_ROW, col).value = income[label]
    _write_support_table(ws, labels, "National Societies to which support was given", answers.get("support_given") or [])
    _write_support_table(
        ws, labels, "National Societies from which support was received", answers.get("support_received") or []
    )


def write_fdrs_identity(wb, *, ns_name: str, year: Optional[int]) -> None:
    ws = wb[INDICATOR_SHEET]
    if ns_name and not _is_formula(ws[NS_NAME_CELL].value):
        ws[NS_NAME_CELL] = ns_name
    if year is not None and not _is_formula(ws[YEAR_CELL].value):
        ws[YEAR_CELL] = year


def _support_amount_column(ws, labels: Mapping[int, str], header_row: int, name_col: int) -> Optional[int]:
    for col in range(name_col + 1, min(ws.max_column or 1, 16) + 1):
        text = _norm(_label_at(ws, ws.cell(header_row, col), labels))
        if any(token in text for token in ("fund", "amount", "chf")):
            return col
    return None


def _read_support_table(ws, labels: Mapping[int, str], title: str) -> List[Dict[str, Any]]:
    placed = _sheet_labels(ws, labels)
    located = _find_label(placed, title)
    if not located:
        return []
    header_row = None
    name_col = None
    for row in range(located[0] + 1, located[0] + 4):
        for col in range(1, 12):
            if _norm(_label_at(ws, ws.cell(row, col), labels)) == "national society":
                header_row = row
                name_col = col
                break
        if header_row:
            break
    if not header_row or not name_col:
        return []
    amount_col = _support_amount_column(ws, labels, header_row, name_col)
    rows: List[Dict[str, Any]] = []
    blank = 0
    for row in range(header_row + 1, header_row + 12):
        name = ws.cell(row, name_col).value
        if _is_formula(name):
            name = ""
        name = str(name or "").strip()
        amount = _number(ws.cell(row, amount_col).value) if amount_col else None
        if not name and amount is None:
            blank += 1
            if blank >= 2:
                break
            continue
        blank = 0
        if name and amount is not None:
            rows.append({"name": name, "amount": _json_number(amount)})
    return rows


def _write_support_table(ws, labels: Mapping[int, str], title: str, rows: Sequence[Mapping[str, Any]]) -> None:
    placed = _sheet_labels(ws, labels)
    located = _find_label(placed, title)
    if not located:
        return
    header_row = None
    name_col = None
    for row in range(located[0] + 1, located[0] + 4):
        for col in range(1, 12):
            if _norm(_label_at(ws, ws.cell(row, col), labels)) == "national society":
                header_row = row
                name_col = col
                break
        if header_row:
            break
    if not header_row or not name_col:
        return
    amount_col = _support_amount_column(ws, labels, header_row, name_col)
    for offset, row in enumerate(rows[:10]):
        excel_row = header_row + 1 + offset
        if not _is_formula(ws.cell(excel_row, name_col).value):
            ws.cell(excel_row, name_col).value = row.get("name") or ""
        if amount_col and not _is_formula(ws.cell(excel_row, amount_col).value):
            ws.cell(excel_row, amount_col).value = row.get("amount")


def workbook_has_fdrs_structure(wb) -> bool:
    return INDICATOR_SHEET in wb.sheetnames and TRANSLATION_SHEET in wb.sheetnames
