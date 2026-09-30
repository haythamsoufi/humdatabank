"""Compare appealgroupchild with the legacy GO v2 appeal catalogue."""

import os
from collections import Counter
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.auth import HTTPBasicAuth

from plugins.emergency_operations.appeal_group import (
    fetch_appeal_group_records,
    map_subtype_to_appeal_type,
)

OUT = Path(r"c:\Humanitarian Databank\appeals-type-differences.xlsx")
ENV = Path(r"c:\Humanitarian Databank\Backoffice\.env")


def load_env():
    if not ENV.exists():
        return
    for line in ENV.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, _, value = line.partition("=")
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def fetch_go():
    headers = {"Accept": "application/json", "User-Agent": "IFRC-Network-Databank/1.0"}
    url = "https://goadmin.ifrc.org/api/v2/appeal/"
    params = {"format": "json", "limit": 1000}
    rows = []
    while url:
        response = requests.get(url, params=params, headers=headers, timeout=120)
        response.raise_for_status()
        payload = response.json()
        rows.extend(payload.get("results") or [])
        print(f"GO {len(rows)} / {payload.get('count')}", flush=True)
        url = payload.get("next")
        params = None
    return rows


def day(value):
    return str(value or "")[:10]


def write_sheet(wb, title, headers, data_rows):
    ws = wb.create_sheet(title)
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r_idx, row in enumerate(data_rows, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)
    last = max(1, len(data_rows) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for col, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(42, max(14, len(header) + 2))


def main():
    load_env()
    user = (os.environ.get("IFRC_API_USER") or "").strip()
    password = (os.environ.get("IFRC_API_PASSWORD") or "").strip()
    if not user or not password:
        raise SystemExit("IFRC API credentials missing")
    # fetch_appeal_group_records reads the env itself
    _ = HTTPBasicAuth(user, password)

    print("Fetching appealgroupchild...", flush=True)
    new_rows = fetch_appeal_group_records("https://go-api.ifrc.org/api/appealgroupchild", timeout=120)
    new_by_code = {row["code"]: row for row in new_rows}
    print("new", len(new_by_code), flush=True)
    go_rows = fetch_go()
    go_by_code = {}
    for item in go_rows:
        code = (item.get("code") or "").strip().upper()
        if code:
            go_by_code[code] = item
    print("go", len(go_by_code), flush=True)

    both = sorted(set(new_by_code) & set(go_by_code))
    type_rows = []
    other_rows = []
    subtype_conflict_rows = []
    for code in both:
        new = new_by_code[code]
        old = go_by_code[code]
        go_type = (old.get("atype_display") or "").strip()
        mapped = new["atype_display"]
        original_mapped = map_subtype_to_appeal_type(new.get("original_subtype"))
        go_country = ""
        country = old.get("country") if isinstance(old.get("country"), dict) else {}
        go_country = (country.get("iso") or "").upper()
        new_country = ((new.get("country") or {}).get("iso") or "")
        reasons = []
        if mapped.casefold() != go_type.casefold():
            reasons.append("type")
        if (new.get("name") or "").casefold() != (old.get("name") or "").casefold():
            reasons.append("name")
        if (new.get("status_display") or "").casefold() != (old.get("status_display") or "").casefold():
            reasons.append("status")
        if day(new.get("start_date")) != day(old.get("start_date")):
            reasons.append("start_date")
        if day(new.get("end_date")) != day(old.get("end_date")):
            reasons.append("end_date")
        if new_country and go_country and new_country != go_country:
            reasons.append("country")
        record = [
            code,
            "; ".join(reasons),
            new.get("appeal_subtype"),
            new.get("original_subtype"),
            mapped,
            original_mapped if new.get("original_subtype") else "",
            go_type,
            new.get("part_of"),
            new.get("name"),
            old.get("name"),
            new.get("status_display"),
            old.get("status_display"),
            day(new.get("start_date")),
            day(old.get("start_date")),
            day(new.get("end_date")),
            day(old.get("end_date")),
            new_country,
            go_country,
        ]
        if "type" in reasons:
            type_rows.append(record)
        elif reasons:
            other_rows.append(record)
        if new.get("original_subtype") and original_mapped.casefold() != mapped.casefold():
            subtype_conflict_rows.append(record)

    only_new = sorted(set(new_by_code) - set(go_by_code))
    only_go = sorted(set(go_by_code) - set(new_by_code))
    child_missing = [code for code in only_new if new_by_code[code].get("part_of")]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "appealgroupchild vs GO v2 appeal"
    summary["A1"].font = Font(bold=True, size=14)
    notes = [
        ("Generated", datetime.now().isoformat(timespec="seconds")),
        ("New feed", "https://go-api.ifrc.org/api/appealgroupchild"),
        ("Legacy feed", "https://goadmin.ifrc.org/api/v2/appeal/"),
        ("Type map", "Subtype Emergency = Emergency Appeal; Subtype Minor Emergency = DREF"),
        ("appealgroupchild rows", len(new_by_code)),
        ("GO appeal rows", len(go_by_code)),
        ("In both", len(both)),
        ("Inconsistent type (mapped subtype != GO type)", len(type_rows)),
        ("Other field differences, type matches", len(other_rows)),
        ("Subtype and Original_Subtype map to different types", len(subtype_conflict_rows)),
        ("Only in appealgroupchild", len(only_new)),
        ("Of those, child appeals (Part_of set) absent from GO", len(child_missing)),
        ("Only in GO", len(only_go)),
        ("Mapped type counts", str(Counter(row["atype_display"] for row in new_by_code.values()))),
        ("GO type counts", str(Counter((item.get("atype_display") or "") for item in go_by_code.values()))),
    ]
    for idx, (label, value) in enumerate(notes, 3):
        summary.cell(idx, 1, label).font = Font(bold=True)
        summary.cell(idx, 2, value)
        summary.cell(idx, 2).alignment = Alignment(wrap_text=True)
    summary.column_dimensions["A"].width = 62
    summary.column_dimensions["B"].width = 110

    headers = [
        "code", "difference", "subtype", "original_subtype", "mapped_type", "original_subtype_mapped",
        "go_type", "part_of", "new_name", "go_name", "new_status", "go_status",
        "new_start", "go_start", "new_end", "go_end", "new_country", "go_country",
    ]
    write_sheet(wb, "Inconsistent types", headers, type_rows)
    write_sheet(wb, "Other differences", headers, other_rows)
    write_sheet(wb, "Subtype conflicts", headers, subtype_conflict_rows)
    write_sheet(
        wb,
        "Children missing on GO",
        ["code", "name", "part_of", "group_name", "mapped_type", "subtype", "original_subtype", "country", "start", "end", "status"],
        [[
            new_by_code[code]["code"],
            new_by_code[code]["name"],
            new_by_code[code]["part_of"],
            new_by_code[code]["group_appeal_name"],
            new_by_code[code]["atype_display"],
            new_by_code[code]["appeal_subtype"],
            new_by_code[code]["original_subtype"],
            (new_by_code[code].get("country") or {}).get("iso"),
            day(new_by_code[code].get("start_date")),
            day(new_by_code[code].get("end_date")),
            new_by_code[code]["status_display"],
        ] for code in child_missing],
    )
    write_sheet(
        wb,
        "Only in GO",
        ["code", "name", "type", "status", "country", "start", "end"],
        [[
            code,
            go_by_code[code].get("name"),
            go_by_code[code].get("atype_display"),
            go_by_code[code].get("status_display"),
            ((go_by_code[code].get("country") or {}) if isinstance(go_by_code[code].get("country"), dict) else {}).get("iso"),
            day(go_by_code[code].get("start_date")),
            day(go_by_code[code].get("end_date")),
        ] for code in only_go],
    )
    wb.save(OUT)
    print({"type_mismatches": len(type_rows), "other": len(other_rows), "subtype_conflicts": len(subtype_conflict_rows), "children_missing_on_go": len(child_missing), "only_go": len(only_go), "out": str(OUT)})


if __name__ == "__main__":
    main()
