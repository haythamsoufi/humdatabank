"""FDRS sync verification background job (AIJob + single item, same pattern as data sync)."""

from __future__ import annotations

import logging
import os
import sys
import time
import uuid
from typing import Any, Dict, List, Optional

from app.extensions import db
from app.models import AIJob, AIJobItem
from app.services.imports.async_import_job_store import (
    FDRS_SYNC_VERIFY_JOB_TYPE,
    clear_import_job_logging_state,
    get_import_job,
    get_import_job_logging_state,
    update_import_job,
    _isolated_job_session,
)
from app.services.ai.ai_job_runner import (
    ensure_ai_job_running,
    job_cancel_requested,
    run_ai_job,
    signal_job_cancel,
    start_ai_job_thread,
)
from app.utils.datetime_helpers import utcnow

logger = logging.getLogger(__name__)


def _fdrs_imports_dir() -> str:
    from plugins.fdrs.scripts_path import ensure_fdrs_scripts_in_path

    return ensure_fdrs_scripts_in_path()


def _touch_verify_item_heartbeat(item_id: int) -> None:
    try:
        with _isolated_job_session() as session:
            item = session.get(AIJobItem, int(item_id))
            if item:
                item.updated_at = utcnow()
    except Exception as exc:
        logger.debug("FDRS verify item heartbeat failed: item=%s err=%s", item_id, exc)


def _summarize_error(exc: BaseException) -> str:
    err_msg = str(exc).strip() or type(exc).__name__
    if len(err_msg) > 2000:
        return err_msg[:1997] + "..."
    return err_msg


def create_fdrs_sync_verify_job(
    *,
    user_id: int,
    template_id: int,
    fdrs_years: Optional[List[int]],
    skip_imputed: bool,
    fresh_imputed: bool,
    problems_only: bool,
    fdrs_reported_import_states: Optional[List[int]],
    output_path: str,
) -> str:
    """Create AIJob + single AIJobItem for async FDRS sync verification."""
    job_id = uuid.uuid4().hex
    now_ts = time.time()
    meta = {
        "template_id": int(template_id),
        "stage": "queued",
        "message": "Queued",
        "current": 0,
        "total": None,
        "percent": 0.0,
        "stats": None,
        "error": None,
        "preview_path": output_path,
        "download_ready": False,
        "last_logged_pct": None,
        "started_ts": now_ts,
        "updated_ts": now_ts,
    }
    payload = {
        "template_id": int(template_id),
        "fdrs_years": fdrs_years,
        "skip_imputed": bool(skip_imputed),
        "fresh_imputed": bool(fresh_imputed),
        "problems_only": bool(problems_only),
        "fdrs_reported_import_states": fdrs_reported_import_states,
        "output_path": output_path,
    }
    job = AIJob(
        id=job_id,
        job_type=FDRS_SYNC_VERIFY_JOB_TYPE,
        user_id=int(user_id or 0),
        status="queued",
        total_items=1,
        meta=meta,
    )
    item = AIJobItem(
        job_id=job_id,
        item_index=0,
        entity_type="form_template",
        entity_id=int(template_id),
        status="queued",
        payload=payload,
    )
    db.session.add(job)
    db.session.add(item)
    db.session.commit()
    return job_id


def get_active_fdrs_sync_verify_jobs_for_user(
    user_id: int,
    *,
    template_id: Optional[int] = None,
) -> List[Dict[str, Any]]:
    if not user_id:
        return []

    query = AIJob.query.filter(
        AIJob.user_id == int(user_id),
        AIJob.job_type == FDRS_SYNC_VERIFY_JOB_TYPE,
        AIJob.status.in_(("queued", "running", "cancel_requested")),
    )
    jobs = query.order_by(AIJob.created_at.desc()).all()
    out: List[Dict[str, Any]] = []
    for job in jobs:
        meta = dict(job.meta or {})
        if template_id is not None and int(meta.get("template_id") or 0) != int(template_id):
            continue
        out.append(
            {
                "job_id": str(job.id),
                "job_type": job.job_type,
                "status": str(job.status),
                "template_id": int(meta.get("template_id") or 0),
                "stage": meta.get("stage") or "",
                "message": meta.get("message") or "",
                "percent": float(meta.get("percent") or 0.0),
                "total_items": int(job.total_items or 1),
            }
        )
    return out


def start_fdrs_sync_verify_job(app, job_id: str) -> None:
    start_ai_job_thread(app, job_id, _run_fdrs_sync_verify_job)


def ensure_fdrs_sync_verify_job_running(app, job_id: str) -> None:
    ensure_ai_job_running(app, job_id, _run_fdrs_sync_verify_job)


def request_fdrs_sync_verify_cancel(job_id: str) -> str:
    job = AIJob.query.get(str(job_id))
    if not job:
        return "missing"
    if job.status in ("completed", "failed", "cancelled"):
        return str(job.status)

    job.status = "cancel_requested"
    try:
        (
            db.session.query(AIJobItem)
            .filter(
                AIJobItem.job_id == str(job_id),
                AIJobItem.status == "queued",
            )
            .update({"status": "cancelled", "error": None}, synchronize_session=False)
        )
    except Exception as exc:
        logger.debug("FDRS verify cancel item update failed: job=%s err=%s", job_id, exc)
        db.session.rollback()
    db.session.commit()
    signal_job_cancel(str(job_id))
    return "cancel_requested"


def build_fdrs_sync_verify_status_payload(job_id: str, template_id: int) -> Optional[Dict[str, Any]]:
    job = get_import_job(job_id)
    if not job or int(job.get("template_id") or 0) != int(template_id):
        return None
    return {
        "job_id": job_id,
        "status": job.get("status"),
        "stage": job.get("stage"),
        "message": job.get("message"),
        "current": job.get("current"),
        "total": job.get("total"),
        "percent": job.get("percent"),
        "stats": job.get("stats"),
        "error": job.get("error"),
        "started_at": job.get("started_at"),
        "updated_at": job.get("updated_at"),
        "download_ready": bool(job.get("download_ready")),
    }


_VERIFY_SHEET_ALIASES = {
    "data_points": "data_points",
    "datapoints": "data_points",
    "summary": "summary",
    "kpi_coverage": "kpi_coverage",
    "kpi-coverage": "kpi_coverage",
    "coverage": "kpi_coverage",
}


def _cell_to_json(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_fdrs_sync_verify_workbook(
    path: str,
    *,
    sheet: str = "data_points",
    status: Optional[str] = None,
    page: int = 1,
    per_page: int = 200,
) -> Dict[str, Any]:
    """Read one verification Excel sheet as paginated table rows.

    The workbook written by ``verify_fdrs_sync`` has ``data_points``,
    ``summary``, and ``kpi_coverage``. Status filter applies only to
    ``data_points`` (column ``status``).
    """
    import openpyxl

    page = max(1, int(page or 1))
    per_page = min(500, max(1, int(per_page or 200)))
    wanted = _VERIFY_SHEET_ALIASES.get((sheet or "").strip().lower(), "data_points")
    status_filter = (status or "").strip().lower() or None

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        if wanted not in wb.sheetnames:
            wanted = sheet_names[0] if sheet_names else wanted
        ws = wb[wanted] if wanted in wb.sheetnames else None
        if ws is None:
            return {
                "sheet": wanted,
                "sheets": sheet_names,
                "columns": [],
                "rows": [],
                "page": page,
                "per_page": per_page,
                "total_rows": 0,
                "filtered_rows": 0,
            }

        rows_iter = ws.iter_rows(values_only=True)
        raw_header = next(rows_iter, None) or ()
        columns = [str(c) if c is not None else "" for c in raw_header]
        status_idx = None
        for i, col in enumerate(columns):
            if col.lower() == "status":
                status_idx = i
                break

        matched: List[Dict[str, Any]] = []
        total_rows = 0
        for raw in rows_iter:
            total_rows += 1
            if status_filter and status_idx is not None:
                cell = raw[status_idx] if status_idx < len(raw) else None
                if str(cell or "").strip().lower() != status_filter:
                    continue
            row: Dict[str, Any] = {}
            for i, col in enumerate(columns):
                if not col:
                    continue
                row[col] = _cell_to_json(raw[i] if i < len(raw) else None)
            matched.append(row)

        filtered_rows = len(matched)
        start = (page - 1) * per_page
        page_rows = matched[start:start + per_page]
        return {
            "sheet": wanted,
            "sheets": sheet_names,
            "columns": [c for c in columns if c],
            "rows": page_rows,
            "page": page,
            "per_page": per_page,
            "total_rows": total_rows,
            "filtered_rows": filtered_rows,
        }
    finally:
        wb.close()


def _process_fdrs_sync_verify_item(app, *, job_id: str, item_id: int) -> None:
    with app.app_context():
        item = AIJobItem.query.get(int(item_id))
        if not item or str(item.job_id) != str(job_id):
            return

        if job_cancel_requested(job_id):
            item.status = "cancelled"
            item.error = None
            update_import_job(
                job_id,
                force=True,
                status="cancelled",
                stage="cancelled",
                message="Cancelled",
                error="Verification cancelled by user.",
            )
            db.session.commit()
            return

        payload = dict(item.payload or {})
        template_id = int(payload.get("template_id") or 0)
        output_path = payload.get("output_path")
        last_cancel_db_check = 0.0

        def _progress_cb(progress_payload: Dict[str, Any]) -> None:
            stage = progress_payload.get("stage") or ""
            pct = progress_payload.get("percent")
            msg = progress_payload.get("message") or ""
            existing = get_import_job(job_id) or {}
            current = (
                progress_payload.get("current")
                if progress_payload.get("current") is not None
                else existing.get("current")
            )
            total = (
                progress_payload.get("total")
                if progress_payload.get("total") is not None
                else existing.get("total")
            )
            update_import_job(
                job_id,
                status="running",
                stage=progress_payload.get("stage") or existing.get("stage"),
                message=progress_payload.get("message") or existing.get("message"),
                current=current,
                total=total,
                percent=float(progress_payload.get("percent") or existing.get("percent") or 0.0),
                stats=(
                    progress_payload.get("stats")
                    if progress_payload.get("stats") is not None
                    else existing.get("stats")
                ),
            )
            _touch_verify_item_heartbeat(item_id)

            try:
                pct_f = float(pct) if pct is not None else None
            except Exception:
                pct_f = None
            log_state = get_import_job_logging_state(job_id)
            last_logged = log_state.get("last_logged_pct")
            should_log = (
                stage in ("complete", "failed", "cancelled", "load_databank", "write_excel")
                or (
                    pct_f is not None
                    and (last_logged is None or abs(pct_f - float(last_logged)) >= 5.0)
                )
            )
            if should_log and pct_f is not None:
                log_state["last_logged_pct"] = pct_f
                update_import_job(job_id, last_logged_pct=pct_f)
            if should_log:
                app.logger.info(
                    "FDRS verify %s: %s %s%% %s",
                    job_id,
                    stage or "-",
                    f"{pct_f:.1f}" if pct_f is not None else "-",
                    msg,
                )

        def _cancel_check() -> bool:
            nonlocal last_cancel_db_check
            if job_cancel_requested(job_id):
                return True
            now = time.time()
            if now - last_cancel_db_check >= 1.0:
                last_cancel_db_check = now
                if job_cancel_requested(job_id):
                    return True
            return False

        update_import_job(
            job_id,
            force=True,
            status="running",
            stage="starting",
            message="Starting...",
            worker_pid=os.getpid(),
        )
        _touch_verify_item_heartbeat(item_id)
        app.logger.info(
            "FDRS verify %s: starting (template_id=%s, years=%s, skip_imputed=%s)",
            job_id,
            template_id,
            payload.get("fdrs_years"),
            payload.get("skip_imputed"),
        )

        terminal_item_status = "failed"
        FdrsVerifyCancelled = None
        try:
            imports_dir = _fdrs_imports_dir()
            if imports_dir not in sys.path:
                sys.path.insert(0, imports_dir)
            from verify_fdrs_sync import FdrsVerifyCancelled, execute_verification

            stats = execute_verification(
                years=payload.get("fdrs_years") or [],
                output_path=str(output_path),
                skip_imputed=bool(payload.get("skip_imputed")),
                fresh_imputed=bool(payload.get("fresh_imputed")),
                problems_only=bool(payload.get("problems_only")),
                reported_states=payload.get("fdrs_reported_import_states"),
                template_id=template_id,
                progress_cb=_progress_cb,
                cancel_check=_cancel_check,
            )
            update_import_job(
                job_id,
                force=True,
                status="completed",
                stage="complete",
                message="Completed",
                percent=100.0,
                stats=dict(stats or {}),
                download_ready=bool(output_path and os.path.isfile(output_path)),
            )
            item.status = "completed"
            item.error = None
            terminal_item_status = "completed"
            app.logger.info(
                "FDRS verify %s: completed total=%s matched=%s skipped=%s missing=%s mismatch=%s",
                job_id,
                (stats or {}).get("total"),
                (stats or {}).get("matched"),
                (stats or {}).get("skipped_intentionally"),
                (stats or {}).get("missing"),
                (stats or {}).get("mismatch"),
            )
        except FdrsVerifyCancelled:
            update_import_job(
                job_id,
                force=True,
                status="cancelled",
                stage="cancelled",
                message="Cancelled",
                error="Verification cancelled by user.",
            )
            item.status = "cancelled"
            item.error = None
            terminal_item_status = "cancelled"
            app.logger.info("FDRS verify %s: cancelled", job_id)
        except Exception as exc:
            logger.exception("Async FDRS verify job failed: %s", exc)
            err_msg = _summarize_error(exc)
            update_import_job(
                job_id,
                force=True,
                status="failed",
                stage="failed",
                message="Failed",
                error=err_msg,
            )
            item.status = "failed"
            item.error = err_msg
            app.logger.error("FDRS verify %s: failed: %s", job_id, exc, exc_info=True)
        finally:
            if item.status not in ("completed", "failed", "cancelled"):
                item.status = terminal_item_status
            db.session.commit()
            clear_import_job_logging_state(job_id)
            db.session.remove()


def _run_fdrs_sync_verify_job(app, job_id: str) -> None:
    run_ai_job(
        app,
        str(job_id),
        _process_fdrs_sync_verify_item,
        concurrency_config_keys=("FDRS_DATA_SYNC_CONCURRENCY",),
        default_concurrency=1,
    )
