"""Published FDRS snapshot, reshaped for the legacy data.ifrc.org client.

The Angular client talks to ``data-api.ifrc.org`` and parses nested ``data`` trees
keyed by legacy KPI strings. Backoffice stores one cell per indicator-bank item, with
sex/age inside ``published_disagg_data``. This module translates between the two.
See ``docs/FDRS-PUBLIC-FRONTEND-ADAPTER.md`` for the contract.

Reads only the published snapshot (``FormData.published_*``) of template 21 items
marked ``privacy = public``.

Legacy semantics this module reproduces:

- ``X_IP`` is reported **plus** imputed values; plain ``X`` is reported only.
- ``KpiImputedValue`` returns every society with a value and flags imputed rows
  with ``source = "I"``.
- ``globalvalidationdata`` is the volunteer headline: a 3-year rolling mean of
  reported values, not the single-year total.
- A society series with no ``year`` (the NS page line chart) includes only the
  latest published year and the five years before it. 2025 starts at 2020.
- Sex and age series (``X_M``, ``X_F``, ``X_M_age_6_12`` ...) come from the
  disaggregation of the base indicator.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date
from typing import Any, Iterable, Mapping, Optional

from sqlalchemy import and_, func

from app.extensions import db
from app.models import (
    AssignedForm,
    Country,
    FormData,
    FormItem,
    IndicatorBank,
    IndicatorBankType,
    IndicatorBankUnit,
    NationalSociety,
)
from app.models.assignments import AssignmentEntityStatus
from app.models.organization import FDRS_ZONES
from app.services.data_quality.helpers import parse_period_year
from app.services.data_retrieval.shared import form_item_privacy_is_public_expr
from app.utils.data_quality_constants import FDRS_TEMPLATE_ID
from plugins.fdrs.published_facts import published_numeric, published_snapshot_clause

# Same ids as ``plugins/fdrs/scripts/fdrs_sync_constants.py`` (kept here so the web
# process does not import the sync scripts). ``test_fdrs_public_compat`` asserts they
# match.
INCOME_MATRIX_ITEM_ID = 943
INCOME_MATRIX_COLUMN = "Funding"
INCOME_KPI_TO_ROW = {
    "h_gov_CHF": "Home Government",
    "f_gov_CHF": "Foreign Government",
    "ind_CHF": "Individuals",
    "corp_CHF": "Corporations",
    "found_CHF": "Foundations",
    "un_CHF": "UN Agencies & other Multilateral Agencies",
    "pooled_f_CHF": "Pooled funds",
    "ngo_CHF": "Non-governmental organizations",
    "si_CHF": "Service income",
    "iga_CHF": "Income generating activity",
    "KPI_incomeFromNSsLC_CHF": "Other National Society",
    "ifrc_CHF": "IFRC",
    "icrc_CHF": "ICRC",
    "other_CHF": "Other",
}
QUESTION_KPI_TO_ITEM = {"KPI_pr_sex": 924, "KPI_sg_sex": 934}
NETWORK_ITEMS = {
    "supported": (919, "Funding provided"),
    "received_support": (929, "Funding Received"),
}

AGE_BANDS = ("6_12", "13_17", "18_29", "30_39", "40_49", "50_59", "60_69", "70_79", "80")
HEADLINE_SUFFIXES = ("_Tot", "_CPD")
IMPUTED_SUFFIX = "_IP"
VOLUNTEER_BANK_CODE = "KPI_PeopleVol"
MAX_INDICATORS_PER_REQUEST = 200

KIND_TOTAL = "total"
KIND_SEX = "sex"
KIND_AGE = "age"
KIND_INCOME = "income"
KIND_QUESTION = "question"
KIND_NETWORK = "network"

_SEX_AGE_RE = re.compile(r"^(?P<base>.+?)_(?P<sex>[MF])(?:_age_(?P<age>[0-9]+(?:_[0-9]+)?))?$")
_NETWORK_RE = re.compile(r"^(?P<prefix>supported|received_support)(?P<slot>\d+)(?P<amount>_amount)?$")
_STRUCTURE_KEYS = frozenset({"mode", "values", "direct", "indirect", "disability", "total", "total_direct", "total_indirect"})
_AGE_ALIASES = {"80_plus": "80", "age_80": "80"}


@dataclass(frozen=True)
class SeriesSpec:
    """One response series, identified by the legacy code the client expects back."""

    code: str
    kind: str
    candidates: tuple[str, ...] = ()
    reported_only: bool = True
    sex: Optional[str] = None
    age: Optional[str] = None
    income_row: Optional[str] = None
    item_id: Optional[int] = None
    network_prefix: Optional[str] = None
    network_slot: int = 1
    network_amount: bool = False


@dataclass(frozen=True)
class Observation:
    value: Any
    imputed: bool = False


# ---------------------------------------------------------------------------
# Legacy code parsing (no database)
# ---------------------------------------------------------------------------


def _bank_candidates(code: str) -> tuple[str, ...]:
    out = [code]
    for suffix in HEADLINE_SUFFIXES:
        if code.endswith(suffix) and len(code) > len(suffix):
            out.append(code[: -len(suffix)])
    return tuple(out)


def parse_code(token: str) -> Optional[SeriesSpec]:
    """Translate one legacy KPI string. ``None`` when it cannot be a series."""
    token = (token or "").strip()
    if not token:
        return None
    imputed_read = token.endswith(IMPUTED_SUFFIX) and len(token) > len(IMPUTED_SUFFIX)
    core = token[: -len(IMPUTED_SUFFIX)] if imputed_read else token
    reported_only = not imputed_read

    if core in INCOME_KPI_TO_ROW:
        return SeriesSpec(
            token,
            KIND_INCOME,
            reported_only=reported_only,
            item_id=INCOME_MATRIX_ITEM_ID,
            income_row=INCOME_KPI_TO_ROW[core],
        )
    if core in QUESTION_KPI_TO_ITEM:
        return SeriesSpec(token, KIND_QUESTION, reported_only=reported_only, item_id=QUESTION_KPI_TO_ITEM[core])
    network = _NETWORK_RE.match(core)
    if network:
        return SeriesSpec(
            token,
            KIND_NETWORK,
            reported_only=reported_only,
            item_id=NETWORK_ITEMS[network.group("prefix")][0],
            network_prefix=network.group("prefix"),
            network_slot=int(network.group("slot")),
            network_amount=bool(network.group("amount")),
        )

    sex_age = _SEX_AGE_RE.match(core)
    if sex_age:
        age = sex_age.group("age")
        if age is not None and age not in AGE_BANDS:
            return None
        return SeriesSpec(
            token,
            KIND_AGE if age else KIND_SEX,
            candidates=_bank_candidates(sex_age.group("base")),
            reported_only=reported_only,
            sex="male" if sex_age.group("sex") == "M" else "female",
            age=age,
        )

    return SeriesSpec(token, KIND_TOTAL, candidates=_bank_candidates(core), reported_only=reported_only)


def _breakdown_specs(spec: SeriesSpec) -> list[SeriesSpec]:
    """Sex (and, for a bare base code, age) series the profile charts ask for implicitly.

    The sex chart requests ``KPI_PeopleVol_Tot`` and reads ``..._M`` / ``..._F`` out of the
    reply; the age chart strips the suffix and requests ``KPI_PeopleVol``.
    """
    if spec.kind != KIND_TOTAL or not spec.reported_only:
        return []
    out = [
        SeriesSpec(f"{spec.code}_{letter}", KIND_SEX, candidates=spec.candidates, sex=sex)
        for letter, sex in (("M", "male"), ("F", "female"))
    ]
    if not spec.code.endswith(HEADLINE_SUFFIXES):
        for letter, sex in (("M", "male"), ("F", "female")):
            for band in AGE_BANDS:
                out.append(
                    SeriesSpec(
                        f"{spec.code}_{letter}_age_{band}",
                        KIND_AGE,
                        candidates=spec.candidates,
                        sex=sex,
                        age=band,
                    )
                )
    return out


def split_codes(raw: Optional[str]) -> list[str]:
    seen: dict[str, None] = {}
    for part in (raw or "").split(","):
        part = part.strip()
        if part:
            seen.setdefault(part, None)
    return list(seen)


def resolve_requested(raw_indicators: Optional[str], *, society_scoped: bool) -> list[SeriesSpec]:
    """Series to return for a comma-separated ``indicator`` parameter.

    Sex and age series are added only when the request is scoped to a society
    (``KPI_Don_Code``). Adding them to the world-wide map call would triple its size.
    """
    specs: dict[str, SeriesSpec] = {}
    for token in split_codes(raw_indicators)[:MAX_INDICATORS_PER_REQUEST]:
        spec = parse_code(token)
        if spec is None:
            continue
        specs.setdefault(spec.code, spec)
        if society_scoped:
            for extra in _breakdown_specs(spec):
                specs.setdefault(extra.code, extra)
    return list(specs.values())


# ---------------------------------------------------------------------------
# Value extraction from a published cell
# ---------------------------------------------------------------------------


def _number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace(" ", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def json_number(value: Optional[float]) -> Any:
    if value is None:
        return None
    if float(value).is_integer():
        return int(value)
    return value


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(text).strip().lower()).strip("_")


def _payload(disagg: Any) -> Optional[Mapping[str, Any]]:
    if isinstance(disagg, str):
        try:
            disagg = json.loads(disagg)
        except json.JSONDecodeError:
            return None
    if not isinstance(disagg, dict):
        return None
    values = disagg.get("values") if "values" in disagg else disagg
    if not isinstance(values, dict):
        return None
    direct = values.get("direct")
    return direct if isinstance(direct, dict) else values


def _leaves(disagg: Any) -> Iterable[tuple[str, float]]:
    payload = _payload(disagg)
    if not payload:
        return
    for key, raw in payload.items():
        name = str(key).strip()
        if name.lower() in _STRUCTURE_KEYS:
            continue
        if isinstance(raw, dict):
            for sub_key, sub_raw in raw.items():
                number = _number(sub_raw)
                if number is not None:
                    yield f"{name}_{sub_key}", number
            continue
        number = _number(raw)
        if number is not None:
            yield name, number


def _sex_and_age(slug: str) -> tuple[Optional[str], str]:
    for sex in ("female", "male"):
        if slug == sex:
            return sex, ""
        if slug.startswith(sex + "_"):
            return sex, slug[len(sex) + 1 :]
    return None, slug


def breakdown_cells(disagg: Any) -> dict[tuple[str, str], float]:
    """``{(sex, age_band): value}``; the age part is ``""`` for a sex-only cell."""
    out: dict[tuple[str, str], float] = {}
    for key, number in _leaves(disagg):
        sex, rest = _sex_and_age(_slug(key))
        if sex is None:
            continue
        band = _AGE_ALIASES.get(rest, rest)
        slot = (sex, band if band in AGE_BANDS else "")
        out[slot] = out.get(slot, 0.0) + number
    return out


def sex_total(disagg: Any, sex: str) -> Optional[float]:
    cells = breakdown_cells(disagg)
    matching = [value for (cell_sex, _band), value in cells.items() if cell_sex == sex]
    return sum(matching) if matching else None


def age_value(disagg: Any, sex: str, band: str) -> Optional[float]:
    return breakdown_cells(disagg).get((sex, band))


def disagg_total(disagg: Any) -> Optional[float]:
    values = [number for _key, number in _leaves(disagg)]
    return sum(values) if values else None


def matrix_cells(disagg: Any) -> dict[str, float]:
    """Matrix cells keyed ``{row label}_{column}`` (lower-cased)."""
    out: dict[str, float] = {}
    payload = _payload(disagg)
    if not payload:
        return out
    for key, raw in payload.items():
        lowered = str(key).strip().lower()
        if lowered.startswith("_chf|"):
            continue
        number = _number(raw)
        if number is not None:
            out[lowered] = number
    return out


def income_matrix_chf_amount(disagg: Any, label: str, column: str) -> Any:
    """Swiss-franc amount for one income-source row.

    Prefers the companion stored beside the cell. Older rows keep that amount in
    the Funding cell itself.
    """
    payload = _payload(disagg) or {}
    lowered = {str(key).strip().lower(): raw for key, raw in payload.items()}
    base = f"{label}_{column}".strip().lower()
    raw = lowered.get(f"_chf|{base}")
    if raw is None:
        raw = lowered.get(base)
    return json_number(_number(raw))


def _clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _total_value(row: Any) -> Optional[float]:
    number = published_numeric(row.published_numeric_value, row.published_value)
    if number is not None:
        return number
    return disagg_total(row.published_disagg_data)


def observation_for(spec: SeriesSpec, row: Any, *, name_to_don: Mapping[str, str]) -> Optional[Observation]:
    """The value ``row`` contributes to ``spec``, or ``None`` when it has none."""
    imputed = (row.published_source or "").strip().lower() == "imputed"
    if spec.reported_only and imputed:
        return None
    if spec.kind == KIND_TOTAL:
        value = json_number(_total_value(row))
    elif spec.kind == KIND_SEX:
        value = json_number(sex_total(row.published_disagg_data, spec.sex))
    elif spec.kind == KIND_AGE:
        value = json_number(age_value(row.published_disagg_data, spec.sex, spec.age))
    elif spec.kind == KIND_INCOME:
        value = income_matrix_chf_amount(row.published_disagg_data, spec.income_row or "", INCOME_MATRIX_COLUMN)
    elif spec.kind == KIND_QUESTION:
        value = _clean_text(row.published_value)
    elif spec.kind == KIND_NETWORK:
        if spec.network_slot != 1:
            return None
        value = _network_value(spec, row, name_to_don)
    else:
        return None
    if value is None:
        return None
    return Observation(value=value, imputed=imputed)


def _network_value(spec: SeriesSpec, row: Any, name_to_don: Mapping[str, str]) -> Optional[str]:
    """Counterpart DON codes (or their amounts), comma-separated, as the client splits them.

    Backoffice keeps one matrix cell per counterpart, so the legacy slot numbers are gone;
    slot 1 carries every counterpart. A blank amount still keeps that counterpart on the list.
    """
    column = NETWORK_ITEMS[spec.network_prefix][1].lower()
    payload = _payload(row.published_disagg_data) or {}
    codes: list[str] = []
    amounts: list[str] = []
    suffix = f"_{column}"
    for key, amount in payload.items():
        lowered = str(key).strip().lower()
        if lowered.startswith("_") or not lowered.endswith(suffix):
            continue
        don = name_to_don.get(_norm(lowered[: -len(suffix)]))
        if not don:
            continue
        number = _number(amount)
        codes.append(don)
        amounts.append("" if number is None else str(json_number(number)))
    if not codes:
        return None
    if spec.network_amount:
        if not any(part.strip() for part in amounts):
            return None
        return ",".join(amounts)
    return ",".join(codes)


def _norm(text: str) -> str:
    return " ".join(str(text).casefold().split())


# ---------------------------------------------------------------------------
# Database access
# ---------------------------------------------------------------------------


def _with_don_code():
    """Active National Societies the public site can select."""
    return and_(
        NationalSociety.is_active.is_(True),
        NationalSociety.fdrs_don_code.isnot(None),
    )


def society_name_index() -> dict[str, str]:
    """Normalised society name, ISO3 or DON code -> DON code.

    Network-support matrix rows are labelled with the National Society name, so the
    reply needs the reverse lookup the importer used to build them.
    """
    index: dict[str, str] = {}
    for society in NationalSociety.query.filter(_with_don_code()).all():
        index.setdefault(_norm(society.name), society.fdrs_don_code)
        index.setdefault(_norm(society.fdrs_don_code), society.fdrs_don_code)
        iso3 = society.country.iso3 if society.country else None
        if iso3:
            index.setdefault(_norm(iso3), society.fdrs_don_code)
    return index


def _resolve_item_ids(specs: list[SeriesSpec]) -> dict[str, list[int]]:
    """Public template-21 form items behind each spec."""
    wanted = {c.lower() for s in specs if s.kind in (KIND_TOTAL, KIND_SEX, KIND_AGE) for c in s.candidates}
    code_to_banks: dict[str, list[int]] = {}
    if wanted:
        for bank_id, code in (
            IndicatorBank.query.filter(func.lower(IndicatorBank.fdrs_kpi_code).in_(wanted))
            .with_entities(IndicatorBank.id, IndicatorBank.fdrs_kpi_code)
            .all()
        ):
            code_to_banks.setdefault(code.lower(), []).append(bank_id)
    special_ids = {s.item_id for s in specs if s.item_id is not None}
    bank_ids = [b for ids in code_to_banks.values() for b in ids]

    conditions = []
    if bank_ids:
        conditions.append(FormItem.indicator_bank_id.in_(bank_ids))
    if special_ids:
        conditions.append(FormItem.id.in_(list(special_ids)))
    if not conditions:
        return {}
    bank_to_items: dict[int, list[int]] = {}
    public_special: set[int] = set()
    for item_id, bank_id in (
        FormItem.query.filter(FormItem.template_id == FDRS_TEMPLATE_ID)
        .filter(form_item_privacy_is_public_expr())
        .filter(db.or_(*conditions))
        .with_entities(FormItem.id, FormItem.indicator_bank_id)
        .all()
    ):
        public_special.add(item_id)
        if bank_id is not None:
            bank_to_items.setdefault(bank_id, []).append(item_id)

    resolved: dict[str, list[int]] = {}
    for spec in specs:
        if spec.item_id is not None:
            if spec.item_id in public_special:
                resolved[spec.code] = [spec.item_id]
            continue
        for candidate in spec.candidates:
            items = [i for b in code_to_banks.get(candidate.lower(), []) for i in bank_to_items.get(b, [])]
            if items:
                resolved[spec.code] = items
                break
    return resolved


def _fetch_cells(
    item_ids: Iterable[int],
    *,
    year: Optional[int],
    don_codes: Optional[list[str]],
    zone_code: Optional[str],
):
    item_ids = list(item_ids)
    if not item_ids:
        return []
    query = (
        db.session.query(
            FormData.form_item_id,
            FormData.published_value,
            FormData.published_numeric_value,
            FormData.published_disagg_data,
            FormData.published_source,
            AssignedForm.period_name,
            NationalSociety.fdrs_don_code.label("don_code"),
            NationalSociety.fdrs_zone_code.label("zone_code"),
        )
        .select_from(FormData)
        .join(AssignmentEntityStatus, FormData.assignment_entity_status_id == AssignmentEntityStatus.id)
        .join(AssignedForm, AssignmentEntityStatus.assigned_form_id == AssignedForm.id)
        .join(NationalSociety, and_(
            NationalSociety.country_id == AssignmentEntityStatus.entity_id,
            _with_don_code(),
        ))
        .filter(AssignedForm.template_id == FDRS_TEMPLATE_ID)
        .filter(AssignmentEntityStatus.entity_type == "country")
        .filter(FormData.form_item_id.in_(item_ids))
        .filter(published_snapshot_clause())
    )
    if year is not None:
        query = query.filter(AssignedForm.period_name.like(f"%{int(year)}%"))
    if don_codes:
        query = query.filter(NationalSociety.fdrs_don_code.in_(don_codes))
    if zone_code:
        query = query.filter(NationalSociety.fdrs_zone_code == zone_code)
    return query.order_by(FormData.published_at.asc().nullsfirst(), FormData.id.asc()).all()


# NS profile line charts plot this many years before the latest published year,
# and that year itself. Last published 2025 -> 2020 through 2025.
NS_PROFILE_YEAR_LOOKBACK = 5


def ns_profile_min_year(last_published: int) -> int:
    """First year of an NS page series. Inclusive of ``last_published - 5``."""
    return int(last_published) - NS_PROFILE_YEAR_LOOKBACK


def compute_series(
    specs: list[SeriesSpec],
    *,
    year: Optional[int] = None,
    don_codes: Optional[list[str]] = None,
    zone_code: Optional[str] = None,
    min_year: Optional[int] = None,
) -> dict[str, dict[str, dict[int, Observation]]]:
    """``{series code: {don code: {year: Observation}}}`` for the specs that have data."""
    if not specs:
        return {}
    don_codes = [d.upper() for d in don_codes] if don_codes else None
    items_by_spec = _resolve_item_ids(specs)
    all_items = {i for ids in items_by_spec.values() for i in ids}
    rows_by_item: dict[int, list[Any]] = {}
    for row in _fetch_cells(all_items, year=year, don_codes=don_codes, zone_code=zone_code):
        rows_by_item.setdefault(row.form_item_id, []).append(row)

    name_to_don = society_name_index() if any(s.kind == KIND_NETWORK for s in specs) else {}
    result: dict[str, dict[str, dict[int, Observation]]] = {}
    for spec in specs:
        by_don: dict[str, dict[int, Observation]] = {}
        for item_id in items_by_spec.get(spec.code, []):
            for row in rows_by_item.get(item_id, []):
                row_year = parse_period_year(row.period_name)
                if row_year is None or (year is not None and row_year != year):
                    continue
                if min_year is not None and row_year < min_year:
                    continue
                observation = observation_for(spec, row, name_to_don=name_to_don)
                if observation is not None:
                    by_don.setdefault(row.don_code, {})[row_year] = observation
        if by_don:
            result[spec.code] = by_don
    return result


def build_data_tree(
    specs: list[SeriesSpec],
    *,
    year: Optional[int] = None,
    don_codes: Optional[list[str]] = None,
    zone_code: Optional[str] = None,
) -> dict[str, Any]:
    """The ``/api/data`` response: indicator -> society -> ``{year, value}``.

    A society request with no ``year`` is the NS page line chart. Every society
    gets the same window: the latest published year and the five years before it.
    """
    min_year = None
    if year is None and don_codes:
        min_year = ns_profile_min_year(last_published_year())
    computed = compute_series(
        specs, year=year, don_codes=don_codes, zone_code=zone_code, min_year=min_year,
    )
    series = []
    for spec in specs:
        by_don = computed.get(spec.code)
        if not by_don:
            continue
        series.append(
            {
                "id": spec.code,
                "data": [
                    {
                        "id": don,
                        "data": [{"year": y, "value": obs.value} for y, obs in sorted(years.items())],
                    }
                    for don, years in sorted(by_don.items())
                ],
            }
        )
    return {"data": series}


def build_imputed_values(
    kpi_code: str,
    *,
    year: int,
    zone_code: Optional[str] = None,
) -> list[dict[str, Any]]:
    """The ``/api/KpiImputedValue`` list. ``source`` is ``"I"`` for imputed, ``"R"`` for reported."""
    spec = parse_code(kpi_code)
    if spec is None or spec.kind != KIND_TOTAL:
        return []
    spec = SeriesSpec(spec.code, KIND_TOTAL, candidates=spec.candidates, reported_only=False)
    by_don = compute_series([spec], year=year, zone_code=zone_code).get(spec.code, {})
    return [
        {"doncode": don, "value": years[year].value, "source": "I" if years[year].imputed else "R"}
        for don, years in sorted(by_don.items())
        if year in years
    ]


def build_kpi_values(kpi_codes: list[str], *, year: int) -> list[dict[str, Any]]:
    """The ``/api/kpivalue`` list, one row per society and KPI. Text answers stay text."""
    specs = []
    for code in kpi_codes:
        spec = parse_code(code)
        if spec is not None and spec.kind == KIND_TOTAL:
            specs.append(SeriesSpec(spec.code, KIND_TOTAL, candidates=spec.candidates, reported_only=False))
    if not specs:
        return []
    items_by_spec = _resolve_item_ids(specs)
    all_items = {i for ids in items_by_spec.values() for i in ids}
    rows_by_item: dict[int, list[Any]] = {}
    for row in _fetch_cells(all_items, year=year, don_codes=None, zone_code=None):
        rows_by_item.setdefault(row.form_item_id, []).append(row)
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for spec in specs:
        for item_id in items_by_spec.get(spec.code, []):
            for row in rows_by_item.get(item_id, []):
                if parse_period_year(row.period_name) != year:
                    continue
                number = _total_value(row)
                value = json_number(number) if number is not None else _clean_text(row.published_value)
                if value is None:
                    continue
                out[(row.don_code, spec.code)] = {
                    "doncode": row.don_code,
                    "DONCode": row.don_code,
                    "kpicode": spec.code,
                    "year": year,
                    "value": value,
                }
    return [out[key] for key in sorted(out)]


def three_year_rolling_mean(values_by_year: Mapping[int, float], end_year: int) -> Optional[float]:
    """Mean of ``end_year-2`` .. ``end_year``.

    A missing year uses the latest earlier value above zero. The society is left
    out when nothing positive was reported at or before the first year of the
    window. This is the figure the public site shows for volunteers.
    """
    filled: list[float] = []
    for year in (end_year - 2, end_year - 1, end_year):
        candidates = [reported_year for reported_year, value in values_by_year.items() if reported_year <= year and value > 0]
        if not candidates:
            return None
        filled.append(float(values_by_year[max(candidates)]))
    return sum(filled) / 3.0


def build_global_validation(year: int) -> list[dict[str, Any]]:
    """Volunteer headline per zone, plus a world row with an empty ``ZoneCode``.

    Each society contributes :func:`three_year_rolling_mean` of reported
    ``KPI_PeopleVol``. Imputed cells are not included. The single-year total
    stays on ``/api/data``.
    """
    spec = SeriesSpec(VOLUNTEER_BANK_CODE, KIND_TOTAL, candidates=(VOLUNTEER_BANK_CODE,), reported_only=True)
    by_don = compute_series([spec], year=None).get(spec.code, {})
    zone_by_don = {
        don: zone
        for don, zone in db.session.query(NationalSociety.fdrs_don_code, NationalSociety.fdrs_zone_code)
        .filter(NationalSociety.fdrs_don_code.in_(list(by_don)))
        .all()
    } if by_don else {}
    totals = {zone: 0.0 for zone in FDRS_ZONES}
    world = 0.0
    for don, years in by_don.items():
        numeric: dict[int, float] = {}
        for reported_year, observation in years.items():
            try:
                value = float(observation.value)
            except (TypeError, ValueError):
                continue
            if value > 0:
                numeric[reported_year] = value
        mean = three_year_rolling_mean(numeric, year)
        if mean is None:
            continue
        world += mean
        zone = zone_by_don.get(don)
        if zone in totals:
            totals[zone] += mean
    rows = [{"ZoneCode": zone, "IntValue": int(round(total))} for zone, total in totals.items()]
    rows.append({"ZoneCode": "", "IntValue": int(round(world))})
    return rows


def last_published_year() -> int:
    """Greatest year with any published public FDRS value (previous calendar year if none)."""
    periods = (
        db.session.query(AssignedForm.period_name)
        .select_from(FormData)
        .join(AssignmentEntityStatus, FormData.assignment_entity_status_id == AssignmentEntityStatus.id)
        .join(AssignedForm, AssignmentEntityStatus.assigned_form_id == AssignedForm.id)
        .join(FormItem, FormData.form_item_id == FormItem.id)
        .filter(AssignedForm.template_id == FDRS_TEMPLATE_ID)
        .filter(AssignmentEntityStatus.entity_type == "country")
        .filter(form_item_privacy_is_public_expr())
        .filter(published_snapshot_clause())
        .distinct()
        .all()
    )
    years = [y for y in (parse_period_year(p[0]) for p in periods) if y is not None]
    return max(years) if years else date.today().year - 1


# ---------------------------------------------------------------------------
# Society reference
# ---------------------------------------------------------------------------


def _legacy_date(value: Optional[date]) -> Optional[str]:
    return value.strftime("%d.%m.%Y") if value else None


def society_to_legacy(society: NationalSociety) -> dict[str, Any]:
    country = society.country
    zone_name = society.fdrs_zone_name or FDRS_ZONES.get(society.fdrs_zone_code or "")
    return {
        "KPI_DON_code": society.fdrs_don_code,
        "NSO_DON_name": society.name,
        "NSO_ZON_name": zone_name,
        "ZON_Code": society.fdrs_zone_code,
        "iso_3": country.iso3 if country else None,
        "iso_2": country.iso2 if country else None,
        "country": country.name if country else None,
        "lat": society.latitude,
        "long": society.longitude,
        "admission_date": _legacy_date(society.admission_date),
        "admission_date_note": society.admission_date_note,
        "cur_code": society.currency_code,
        "SearchString": society.fdrs_search_string,
        "Url": society.website,
        "url": society.website,
        "facebook": society.facebook,
        "twitter": society.twitter,
        "othersocial": society.other_social,
    }


def list_societies() -> list[dict[str, Any]]:
    rows = NationalSociety.query.filter(_with_don_code()).order_by(NationalSociety.name.asc()).all()
    return [society_to_legacy(s) for s in rows]


def get_society(don_code: str) -> Optional[dict[str, Any]]:
    society = NationalSociety.query.filter(
        func.upper(NationalSociety.fdrs_don_code) == (don_code or "").strip().upper(),
        NationalSociety.is_active.is_(True),
    ).first()
    return society_to_legacy(society) if society else None


# ---------------------------------------------------------------------------
# CSV / Excel download
# ---------------------------------------------------------------------------

EXPORT_COLUMNS = (
    "KPI_DON_code",
    "NSO_DON_name",
    "iso_3",
    "ZON_Code",
    "NSO_ZON_name",
    "year",
    "KPI_code",
    "Value",
    "source",
)
# One row per KPI code that appears on the Data sheet. ``KPI_code`` matches that sheet.
MAPPING_COLUMNS = (
    "KPI_code",
    "name",
    "definition",
    "type",
    "unit",
    "breakdown",
    "base_code",
    "area",
    "disaggregation_guidance",
    "data_source",
)
_SEX_LETTER = {"male": "M", "female": "F"}
_AGE_BAND_LABELS = {
    "6_12": "6-12",
    "13_17": "13-17",
    "18_29": "18-29",
    "30_39": "30-39",
    "40_49": "40-49",
    "50_59": "50-59",
    "60_69": "60-69",
    "70_79": "70-79",
    "80": "80 and over",
}
_DATA_COLUMN_WIDTHS = {
    "KPI_DON_code": 16,
    "NSO_DON_name": 28,
    "iso_3": 10,
    "ZON_Code": 12,
    "NSO_ZON_name": 18,
    "year": 10,
    "KPI_code": 32,
    "Value": 16,
    "source": 12,
}
_MAPPING_COLUMN_WIDTHS = {
    "KPI_code": 32,
    "name": 36,
    "definition": 64,
    "type": 18,
    "unit": 16,
    "breakdown": 24,
    "base_code": 24,
    "area": 28,
    "disaggregation_guidance": 40,
    "data_source": 28,
}


def _source_label(row: Any) -> str:
    if (row.published_source or "").strip().lower() == "imputed":
        return "imputed"
    return "reported"


def _codes_for_cell(row: Any, *, name_to_don: Mapping[str, str]) -> list[tuple[str, Any]]:
    """Legacy KPI strings this published cell contributes to the download."""
    item_id = row.form_item_id
    bank_code = (row.fdrs_kpi_code or "").strip()
    if bank_code:
        found: list[tuple[str, Any]] = []
        total = json_number(_total_value(row))
        if total is not None:
            found.append((bank_code, total))
        cells = breakdown_cells(row.published_disagg_data)
        sex_sums: dict[str, float] = {}
        for (sex, band), number in cells.items():
            letter = _SEX_LETTER.get(sex)
            if letter is None:
                continue
            sex_sums[letter] = sex_sums.get(letter, 0.0) + number
            if band:
                found.append((f"{bank_code}_{letter}_age_{band}", json_number(number)))
        for letter, total_for_sex in sex_sums.items():
            found.append((f"{bank_code}_{letter}", json_number(total_for_sex)))
        return found

    if item_id == INCOME_MATRIX_ITEM_ID:
        found_income: list[tuple[str, Any]] = []
        for code, label in INCOME_KPI_TO_ROW.items():
            amount = income_matrix_chf_amount(row.published_disagg_data, label, INCOME_MATRIX_COLUMN)
            if amount is None:
                continue
            found_income.append((code, amount))
        return found_income

    for code, question_item in QUESTION_KPI_TO_ITEM.items():
        if question_item == item_id:
            text = _clean_text(row.published_value)
            return [(code, text)] if text else []

    for prefix, (network_item, _column) in NETWORK_ITEMS.items():
        if network_item != item_id:
            continue
        spec = SeriesSpec(
            f"{prefix}1",
            KIND_NETWORK,
            reported_only=False,
            item_id=network_item,
            network_prefix=prefix,
            network_slot=1,
        )
        amounts = SeriesSpec(
            f"{prefix}1_amount",
            KIND_NETWORK,
            reported_only=False,
            item_id=network_item,
            network_prefix=prefix,
            network_slot=1,
            network_amount=True,
        )
        found = []
        for candidate in (spec, amounts):
            observation = observation_for(candidate, row, name_to_don=name_to_don)
            if observation is not None:
                found.append((candidate.code, observation.value))
        return found
    return []


def export_rows(*, don_code: Optional[str] = None) -> list[dict[str, Any]]:
    """One row per published public KPI, society and year. Later publications win a duplicate."""
    query = (
        db.session.query(
            FormData.id,
            FormData.form_item_id,
            FormData.published_value,
            FormData.published_numeric_value,
            FormData.published_disagg_data,
            FormData.published_source,
            AssignedForm.period_name,
            IndicatorBank.fdrs_kpi_code,
            NationalSociety.fdrs_don_code.label("don_code"),
            NationalSociety.name.label("name"),
            NationalSociety.fdrs_zone_code.label("zone_code"),
            NationalSociety.fdrs_zone_name.label("zone_name"),
            Country.iso3.label("country_iso3"),
        )
        .select_from(FormData)
        .join(AssignmentEntityStatus, FormData.assignment_entity_status_id == AssignmentEntityStatus.id)
        .join(AssignedForm, AssignmentEntityStatus.assigned_form_id == AssignedForm.id)
        .join(FormItem, FormData.form_item_id == FormItem.id)
        .outerjoin(IndicatorBank, IndicatorBank.id == FormItem.indicator_bank_id)
        .join(NationalSociety, and_(
            NationalSociety.country_id == AssignmentEntityStatus.entity_id,
            _with_don_code(),
        ))
        .outerjoin(Country, Country.id == NationalSociety.country_id)
        .filter(AssignedForm.template_id == FDRS_TEMPLATE_ID)
        .filter(AssignmentEntityStatus.entity_type == "country")
        .filter(form_item_privacy_is_public_expr())
        .filter(published_snapshot_clause())
        .order_by(FormData.published_at.asc().nullsfirst(), FormData.id.asc())
    )
    if don_code:
        query = query.filter(func.upper(NationalSociety.fdrs_don_code) == don_code.strip().upper())
    fetched = query.all()
    network_ids = {item_id for item_id, _column in NETWORK_ITEMS.values()}
    name_to_don = society_name_index() if any(row.form_item_id in network_ids for row in fetched) else {}

    chosen: dict[tuple[str, int, str], dict[str, Any]] = {}
    for row in fetched:
        year = parse_period_year(row.period_name)
        if year is None:
            continue
        identity = {
            "KPI_DON_code": row.don_code,
            "NSO_DON_name": row.name,
            "iso_3": row.country_iso3,
            "ZON_Code": row.zone_code,
            "NSO_ZON_name": row.zone_name or FDRS_ZONES.get(row.zone_code or ""),
            "year": year,
            "source": _source_label(row),
        }
        for code, value in _codes_for_cell(row, name_to_don=name_to_don):
            if value is None or not code:
                continue
            chosen[(row.don_code, year, code)] = {**identity, "KPI_code": code, "Value": value}
    return [chosen[key] for key in sorted(chosen)]


def _blank(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _area_text(area: Any, area_label: Any) -> Optional[str]:
    code = _blank(area)
    label = _blank(area_label)
    if code and label:
        return f"{code} — {label}"
    return code or label


def _base_code(spec: SeriesSpec) -> str:
    if spec.candidates:
        return spec.candidates[-1]
    return spec.code


def _breakdown_label(spec: SeriesSpec) -> str:
    if spec.kind == KIND_TOTAL:
        return "Total"
    if spec.kind == KIND_INCOME:
        return "Income source"
    if spec.kind == KIND_QUESTION:
        return "Question"
    if spec.kind == KIND_NETWORK:
        return "Amount" if spec.network_amount else "Counterpart societies"
    sex = "Male" if spec.sex == "male" else "Female" if spec.sex == "female" else ""
    if spec.kind == KIND_SEX:
        return sex
    band = _AGE_BAND_LABELS.get(spec.age or "", spec.age or "")
    return f"{sex}, age {band}".strip(", ")


def _mapping_sort_key(code: str) -> tuple:
    spec = parse_code(code)
    if spec is None:
        return (9, code.lower())
    if spec.kind == KIND_INCOME:
        order = list(INCOME_KPI_TO_ROW).index(code) if code in INCOME_KPI_TO_ROW else 99
        return (1, order, code.lower())
    if spec.kind == KIND_QUESTION:
        return (2, code.lower())
    if spec.kind == KIND_NETWORK:
        return (
            3,
            0 if spec.network_prefix == "supported" else 1,
            1 if spec.network_amount else 0,
            code.lower(),
        )
    age_index = AGE_BANDS.index(spec.age) if spec.age in AGE_BANDS else 99
    sex_index = {"male": 0, "female": 1}.get(spec.sex or "", 0)
    kind_index = {KIND_TOTAL: 0, KIND_SEX: 1, KIND_AGE: 2}.get(spec.kind, 9)
    return (0, _base_code(spec).lower(), kind_index, sex_index, age_index, code.lower())


def _bank_catalog(codes: set[str]) -> dict[str, dict[str, Any]]:
    """Public FDRS indicator-bank rows keyed by lower-case ``fdrs_kpi_code``."""
    wanted = {code.strip().lower() for code in codes if code and code.strip()}
    if not wanted:
        return {}
    fetched = (
        db.session.query(
            IndicatorBank.fdrs_kpi_code,
            IndicatorBank.name,
            IndicatorBank.definition,
            func.coalesce(IndicatorBankType.name, IndicatorBank.type).label("type_name"),
            func.coalesce(IndicatorBankUnit.name, IndicatorBank.unit).label("unit_name"),
            IndicatorBank.area,
            IndicatorBank.area_label,
            IndicatorBank.disaggregation_guidance,
            IndicatorBank.data_source,
        )
        .select_from(IndicatorBank)
        .outerjoin(IndicatorBankType, IndicatorBankType.id == IndicatorBank.indicator_type_id)
        .outerjoin(IndicatorBankUnit, IndicatorBankUnit.id == IndicatorBank.indicator_unit_id)
        .join(FormItem, FormItem.indicator_bank_id == IndicatorBank.id)
        .filter(FormItem.template_id == FDRS_TEMPLATE_ID)
        .filter(form_item_privacy_is_public_expr())
        .filter(func.lower(IndicatorBank.fdrs_kpi_code).in_(wanted))
        .all()
    )
    catalog: dict[str, dict[str, Any]] = {}
    for row in fetched:
        key = (row.fdrs_kpi_code or "").strip().lower()
        entry = {
            "name": _blank(row.name),
            "definition": _blank(row.definition),
            "type": _blank(row.type_name),
            "unit": _blank(row.unit_name),
            "area": _area_text(row.area, row.area_label),
            "disaggregation_guidance": _blank(row.disaggregation_guidance),
            "data_source": _blank(row.data_source),
        }
        current = catalog.get(key)
        if current is None or (not current.get("definition") and entry.get("definition")):
            catalog[key] = entry
    return catalog


def _question_catalog(codes: set[str]) -> dict[str, dict[str, Any]]:
    wanted = {code for code in codes if code in QUESTION_KPI_TO_ITEM}
    if not wanted:
        return {}
    items = {
        item.id: item
        for item in FormItem.query.filter(FormItem.id.in_([QUESTION_KPI_TO_ITEM[code] for code in wanted])).all()
    }
    catalog: dict[str, dict[str, Any]] = {}
    for code in wanted:
        item = items.get(QUESTION_KPI_TO_ITEM[code])
        if item is None:
            continue
        catalog[code] = {
            "name": _blank(item.label),
            "definition": _blank(item.definition),
            "type": _blank(item.type) or _blank(item.item_type),
        }
    return catalog


def _describe_code(
    code: str,
    banks: Mapping[str, Mapping[str, Any]],
    questions: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    row: dict[str, Any] = {column: None for column in MAPPING_COLUMNS}
    row["KPI_code"] = code
    row["base_code"] = code
    spec = parse_code(code)
    if spec is None:
        return row
    row["breakdown"] = _breakdown_label(spec)
    row["base_code"] = _base_code(spec)
    if spec.kind == KIND_INCOME:
        label = spec.income_row or code
        row["name"] = label
        row["definition"] = f"Funding from {label}, in Swiss francs."
        row["type"] = "number"
        row["unit"] = "CHF"
        row["base_code"] = code
        return row
    if spec.kind == KIND_QUESTION:
        meta = questions.get(code) or {}
        row["name"] = meta.get("name")
        row["definition"] = meta.get("definition")
        row["type"] = meta.get("type")
        row["base_code"] = code
        return row
    if spec.kind == KIND_NETWORK:
        given = spec.network_prefix == "supported"
        if spec.network_amount:
            row["name"] = "Support provided (CHF)" if given else "Support received (CHF)"
            row["definition"] = "Amounts in Swiss francs, in the same order as the counterpart DON codes."
            row["type"] = "number"
            row["unit"] = "CHF"
        else:
            row["name"] = "National Societies supported" if given else "National Societies support was received from"
            row["definition"] = (
                "DON codes of the National Societies this society supported, separated by commas."
                if given
                else "DON codes of the National Societies this society received support from, separated by commas."
            )
            row["type"] = "text"
        row["base_code"] = code
        return row
    meta = None
    for candidate in spec.candidates:
        meta = banks.get(candidate.lower())
        if meta:
            row["base_code"] = candidate
            break
    if meta:
        for field in ("name", "definition", "type", "unit", "area", "disaggregation_guidance", "data_source"):
            row[field] = meta.get(field)
    return row


def export_mapping(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Indicator dictionary for the KPI codes present in ``rows``.

    Sex and age codes inherit the parent indicator's name, definition, type, and unit.
    Income sources, questions, and network-support codes are described from the FDRS
    form, because they are not indicator-bank rows.
    """
    codes = {str(row.get("KPI_code") or "").strip() for row in rows}
    codes.discard("")
    if not codes:
        return []
    bank_codes: set[str] = set()
    question_codes: set[str] = set()
    for code in codes:
        spec = parse_code(code)
        if spec is None:
            continue
        if spec.kind == KIND_QUESTION:
            question_codes.add(code)
        bank_codes.update(spec.candidates)
    banks = _bank_catalog(bank_codes)
    questions = _question_catalog(question_codes)
    return [_describe_code(code, banks, questions) for code in sorted(codes, key=_mapping_sort_key)]


def export_csv(rows: Iterable[Mapping[str, Any]]) -> bytes:
    """Data table only. A CSV has one table; the indicator dictionary is the workbook's second sheet."""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS, extrasaction="ignore", lineterminator="\r\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: "" if row.get(column) is None else row[column] for column in EXPORT_COLUMNS})
    return buffer.getvalue().encode("utf-8-sig")


def _fill_sheet(sheet: Any, columns: tuple[str, ...], records: Iterable[Mapping[str, Any]], widths: Mapping[str, int]) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheet.append(list(columns))
    for record in records:
        sheet.append([record.get(column) for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    if sheet.dimensions:
        sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 18)


def export_xlsx(
    rows: Iterable[Mapping[str, Any]],
    mapping: Optional[Iterable[Mapping[str, Any]]] = None,
) -> bytes:
    """Workbook with a Data sheet and an Indicators sheet for the codes in that data.

    Pass ``mapping`` when the caller already built the indicator dictionary, so a
    batch of society files does not look the definitions up again.
    """
    from openpyxl import Workbook

    materialised = list(rows)
    indicator_rows = list(mapping) if mapping is not None else export_mapping(materialised)
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    _fill_sheet(data, EXPORT_COLUMNS, materialised, _DATA_COLUMN_WIDTHS)
    indicators = workbook.create_sheet("Indicators")
    _fill_sheet(indicators, MAPPING_COLUMNS, indicator_rows, _MAPPING_COLUMN_WIDTHS)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
