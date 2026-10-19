"""Thumbnails for PDFs are covered elsewhere. These cover Office files and video links."""
import io
import zipfile

import pytest
from docx import Document
from PIL import Image

pytestmark = [pytest.mark.unit]


def _png(width=200, height=200, color=(10, 20, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, "PNG")
    return buf.getvalue()


def _docx_with_picture() -> bytes:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as package:
        package.writestr("word/media/image1.png", _png())
    return archive.getvalue()


def test_youtube_id_accepts_short_and_watch_links():
    from app.services.content.thumbnail_service import ThumbnailService

    assert ThumbnailService.youtube_video_id("https://youtu.be/s8gWS22wpro") == "s8gWS22wpro"
    assert ThumbnailService.youtube_video_id("https://www.youtube.com/watch?v=s8gWS22wpro") == "s8gWS22wpro"
    assert ThumbnailService.youtube_video_id("https://evil.example/watch?v=s8gWS22wpro") is None


def test_word_file_uses_its_embedded_picture(app):
    from app.services.content.thumbnail_service import ThumbnailService

    stored = {}

    def upload(category, rel_path, data):
        stored["path"] = rel_path
        stored["png"] = data[:8]
        return rel_path

    with app.app_context():
        import app.services.platform.storage_service as storage
        original = storage.upload
        storage.upload = upload
        try:
            path = ThumbnailService.generate_from_document_bytes(
                _docx_with_picture(), "guide.docx", "folder-1", language_code="fr", title="Guide",
            )
        finally:
            storage.upload = original

    assert path == "folder-1/thumbnails/thumbnail_fr.png"
    assert stored["png"].startswith(b"\x89PNG")


def test_word_file_without_a_picture_gets_a_text_cover(app):
    from app.services.content.thumbnail_service import ThumbnailService

    document = Document()
    document.add_paragraph("How to add a user")
    document.add_paragraph("Open the administration screen.")
    buf = io.BytesIO()
    document.save(buf)

    stored = {}

    def upload(category, rel_path, data):
        stored["png"] = data
        return rel_path

    with app.app_context():
        import app.services.platform.storage_service as storage
        original = storage.upload
        storage.upload = upload
        try:
            path = ThumbnailService.generate_from_document_bytes(
                buf.getvalue(), "howto.docx", "folder-2", language_code="en", title="How to add a user",
            )
        finally:
            storage.upload = original

    assert path.endswith("thumbnail_en.png")
    assert stored["png"].startswith(b"\x89PNG")


def test_video_thumbnail_uses_the_youtube_preview(app, monkeypatch):
    from app.services.content.thumbnail_service import ThumbnailService

    def fake_download(url, allowed_hosts=None):
        return _png()

    monkeypatch.setattr(ThumbnailService, "_download_https", staticmethod(fake_download))
    stored = {}

    def upload(category, rel_path, data):
        stored["url_path"] = rel_path
        return rel_path

    with app.app_context():
        import app.services.platform.storage_service as storage
        original = storage.upload
        storage.upload = upload
        try:
            path = ThumbnailService.generate_from_youtube_url(
                "https://youtu.be/s8gWS22wpro", "folder-3", language_code="es",
            )
        finally:
            storage.upload = original

    assert path == "folder-3/thumbnails/thumbnail_es.png"
