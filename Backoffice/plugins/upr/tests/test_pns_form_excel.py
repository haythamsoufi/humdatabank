"""Workbook mapping for the PNS Planning and PNS Reporting Excel forms."""

from __future__ import annotations

import os
import sys

BACKOFFICE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
PLUGIN_SCRIPTS = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts"))
for _path in (BACKOFFICE_DIR, PLUGIN_SCRIPTS):
    if _path not in sys.path:
        sys.path.insert(0, _path)

from openpyxl import load_workbook  # noqa: E402

from pns_form_excel import (  # noqa: E402
    default_planning_columns,
    planning_column_slots,
    planning_funding_cells,
    prepare_planning_workbook,
    read_planning_data_rows,
    read_planning_grid_overrides,
    read_reporting_funding_rows,
    read_staff_rows,
    reporting_column_map,
    reporting_funding_cells,
    rewrite_planning_year_formula,
    staff_cells,
    workbook_has_planning_structure,
    workbook_has_reporting_structure,
    write_planning_funding_rows,
)

T22_COLUMNS = [
    *(f"{code}_[assignment_year]" for code in ("SP1", "SP2", "EFs", "fr_tot", "confirmedfn")),
    *(f"{code}_[[assignment_year]+1]" for code in ("SP1", "SP2", "EFs", "fr_tot", "confirmedfn")),
    *(f"{code}_[[assignment_year]+2]" for code in ("SP1", "SP2", "EFs", "fr_tot", "confirmedfn")),
]

PLANNING = os.path.join(
    os.path.dirname(__file__), "..", "static", "templates", "pns_planning_form.xlsx"
)
REPORTING = os.path.join(
    os.path.dirname(__file__), "..", "static", "templates", "pns_reporting_form.xlsx"
)


def test_planning_template_has_funding_tables():
    wb = load_workbook(PLANNING, read_only=False, data_only=False)
    try:
        assert workbook_has_planning_structure(wb)
        formula = wb["Funding Requirement"]["D12"].value
        assert "SP1" in formula
        assert '"2026"' in formula
    finally:
        wb.close()


def test_rewrite_planning_years_shifts_literals_without_collisions():
    formula = 'SUMIFS(PNS_data[Value], PNS_data[Year], "2026") + SUMIFS(PNS_data[Value], PNS_data[Year], "2028")'
    updated = rewrite_planning_year_formula(formula, 2030)
    assert '"2030"' in updated
    assert '"2032"' in updated
    assert '"2026"' not in updated
    assert '"2028"' not in updated


def test_planning_data_rows_become_country_matrix_cells():
    wb = load_workbook(PLANNING, read_only=False, data_only=False)
    try:
        rows = read_planning_data_rows(wb)
    finally:
        wb.close()
    sample = [row for row in rows if row["ns"] == "American Red Cross" and row["country"] == "Bangladesh" and row["year"] == 2026]
    cells, warnings = planning_funding_cells(
        sample,
        ns_name="American Red Cross",
        year=2026,
        country_name_to_id={"bangladesh": 50},
        form_columns=["SP1", "EFs"],
    )
    assert warnings == []
    assert cells["50_EFs"] == 30000
    assert cells["50_SP1"] == 400000
    assert "50_SP2" not in cells


def test_column_slots_resolve_assignment_year_placeholders():
    slots = planning_column_slots(T22_COLUMNS, 2026)
    assert slots[("SP1", 2026)] == "SP1_2026"
    assert slots[("EFs", 2027)] == "EFs_2027"
    assert slots[("confirmedfn", 2028)] == "confirmedfn_2028"
    assert not any(attribute == "Total" for attribute, _year in slots)
    assert planning_column_slots(None, 2026) == planning_column_slots(default_planning_columns(), 2026)


def test_planning_rows_for_later_years_use_year_columns():
    rows = [
        {"ns": "American Red Cross", "country": "Bangladesh", "year": 2026, "attribute": "SP1", "value": 10},
        {"ns": "American Red Cross", "country": "Bangladesh", "year": 2027, "attribute": "SP1", "value": 20},
        {"ns": "American Red Cross", "country": "Bangladesh", "year": 2028, "attribute": "EFs", "value": 30},
        {"ns": "American Red Cross", "country": "Bangladesh", "year": 2029, "attribute": "SP1", "value": 40},
        {"ns": "American Red Cross", "country": "Bangladesh", "year": 2027, "attribute": "Total", "value": 50},
    ]
    cells, warnings = planning_funding_cells(
        rows,
        ns_name="American Red Cross",
        year=2026,
        country_name_to_id={"bangladesh": 50},
        form_columns=T22_COLUMNS,
    )
    assert warnings == []
    assert cells == {"50_SP1_2026": 10, "50_SP1_2027": 20, "50_EFs_2028": 30}


def test_planning_grid_reads_confirmed_funding_per_year_block():
    wb = load_workbook(PLANNING, read_only=False, data_only=False)
    try:
        prepare_planning_workbook(wb, ns_name="American Red Cross", year=2026)
        ws = wb["Funding Requirement"]
        assert ws["C12"].value == "Angola"
        ws["K12"] = 100
        ws["S12"] = 200
        ws["AA12"] = 300
        rows = read_planning_grid_overrides(wb)
    finally:
        wb.close()
    cells, _warnings = planning_funding_cells(
        rows,
        ns_name="American Red Cross",
        year=2026,
        country_name_to_id={"angola": 7},
        form_columns=T22_COLUMNS,
    )
    assert cells == {"7_confirmedfn_2026": 100, "7_confirmedfn_2027": 200, "7_confirmedfn_2028": 300}


def test_planning_export_writes_each_year_and_confirmed_funding():
    wb = load_workbook(PLANNING, read_only=False, data_only=False)
    try:
        prepare_planning_workbook(wb, ns_name="American Red Cross", year=2026)
        write_planning_funding_rows(
            wb,
            {
                "7_SP1_2026": 10,
                "7_EFs_2027": 20,
                "7_SP2_2028": {"original": 5, "modified": 30, "isModified": True},
                "7_SP2_2026": {"original": 5, "modified": "", "isModified": True},
                "7_fr_tot_2026": 10,
                "7_confirmedfn_2027": 40,
            },
            ns_name="American Red Cross",
            year=2026,
            country_id_to_name={7: "Angola"},
            form_columns=T22_COLUMNS,
        )
        data = wb["FR data"]
        written = {
            (data.cell(row, 3).value, data.cell(row, 4).value): data.cell(row, 5).value
            for row in range(2, 6)
            if data.cell(row, 1).value
        }
        confirmed = wb["Funding Requirement"]["S12"].value
    finally:
        wb.close()
    assert written == {(2026, "SP1"): 10, (2027, "EFs"): 20, (2028, "SP2"): 30}
    assert confirmed == 40


def test_staff_rows_use_host_ns_ids():
    wb = load_workbook(PLANNING, read_only=False, data_only=False)
    try:
        ws = wb["PNS Staff Presence"]
        ws["D2"] = 3
        ws["F2"] = 1
        rows = read_staff_rows(wb)
    finally:
        wb.close()
    cells, warnings = staff_cells(rows, iso3_to_ns_id={"AGO": 80}, form_columns=["intl_delegates_hns", "intl_delegates_ifrc"])
    assert warnings == []
    assert cells["80_intl_delegates_hns"] == 3
    assert cells["80_intl_delegates_ifrc"] == 1


def test_reporting_template_maps_funding_columns():
    wb = load_workbook(REPORTING, read_only=False, data_only=False)
    try:
        assert workbook_has_reporting_structure(wb)
        ws = wb["Funding data"]
        ws["C15"] = "Angola"
        ws["E15"] = 10
        ws["F15"] = 4
        rows = read_reporting_funding_rows(wb)
    finally:
        wb.close()
    column_map = reporting_column_map(["Total Funding", "Total Expenditure", "Total Transferred to HNS"])
    cells, warnings = reporting_funding_cells(
        rows,
        country_name_to_ns_id={"angola": 12},
        column_map=column_map,
    )
    assert warnings == []
    assert cells["12_Total Funding"] == 10
    assert cells["12_Total Expenditure"] == 4
