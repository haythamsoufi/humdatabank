"""Single entry point for opening user-supplied spreadsheet files.

``openpyxl.load_workbook`` / ``pandas.read_excel`` happily inflate a tiny
zip-bomb ``.xlsx`` into gigabytes of XML or a sheet dimensioned to millions of
rows. Every code path that opens an uploaded (or otherwise untrusted) workbook
must use :func:`load_workbook_safe`, :func:`read_excel_safe` or
:func:`inspect_xlsx`, which enforce, *before* openpyxl parses anything:

* a maximum raw size,
* a maximum member count and total uncompressed size,
* a maximum per-member compression ratio,
* maximum declared sheet dimensions (rows / columns) and sheet count,

and then default to ``read_only=True`` so memory stays bounded. Trusted,
server-owned template files are not covered and may keep using openpyxl
directly (see docs/DEVELOPER-HANDBOOK.md "Safe primitives").

Limits can be overridden with ``WORKBOOK_MAX_*`` Flask config keys.
"""

from __future__ import annotations

import io
import os
import re
import zipfile
from dataclasses import dataclass
from typing import Any, BinaryIO, Iterator, Optional, Union

WorkbookSource = Union[bytes, bytearray, str, "os.PathLike[str]", BinaryIO, Any]

_DIMENSION_RE = re.compile(rb'<dimension\s+ref="([A-Za-z]+)(\d+)(?::([A-Za-z]+)(\d+))?"')
_SHEET_MEMBER_RE = re.compile(r"^xl/worksheets/[^/]+\.xml$")
_HEAD_BYTES = 16384


class UnsafeWorkbookError(ValueError):
    """The spreadsheet is malformed or exceeds a configured safety limit.

    ``str(exc)`` is a stable, user-presentable message that never echoes file content.
    """


@dataclass(frozen=True)
class WorkbookLimits:
    max_upload_bytes: int = 25 * 1024 * 1024
    max_uncompressed_bytes: int = 250 * 1024 * 1024
    max_members: int = 2000
    max_member_ratio: int = 200
    ratio_min_member_bytes: int = 1024 * 1024
    max_sheets: int = 100
    max_rows: int = 200_000
    max_cols: int = 1024

    @classmethod
    def from_config(cls) -> "WorkbookLimits":
        try:
            from flask import current_app, has_app_context

            cfg = current_app.config if has_app_context() else {}
        except Exception:
            cfg = {}
        base = cls()

        def pick(key: str, default: int) -> int:
            try:
                value = int(cfg.get(key, default))
            except (TypeError, ValueError):
                return default
            return value if value > 0 else default

        return cls(
            max_upload_bytes=pick("WORKBOOK_MAX_UPLOAD_BYTES", base.max_upload_bytes),
            max_uncompressed_bytes=pick("WORKBOOK_MAX_UNCOMPRESSED_BYTES", base.max_uncompressed_bytes),
            max_members=pick("WORKBOOK_MAX_MEMBERS", base.max_members),
            max_member_ratio=pick("WORKBOOK_MAX_MEMBER_RATIO", base.max_member_ratio),
            ratio_min_member_bytes=base.ratio_min_member_bytes,
            max_sheets=pick("WORKBOOK_MAX_SHEETS", base.max_sheets),
            max_rows=pick("WORKBOOK_MAX_ROWS", base.max_rows),
            max_cols=pick("WORKBOOK_MAX_COLS", base.max_cols),
        )


def _col_to_index(letters: bytes) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ch - 64)
    return n


def _read_source_bytes(source: WorkbookSource, limits: WorkbookLimits) -> bytes:
    cap = limits.max_upload_bytes
    if isinstance(source, (bytes, bytearray)):
        data = bytes(source)
    elif isinstance(source, (str, os.PathLike)) and not hasattr(source, "read"):
        path = os.fspath(source)
        try:
            size = os.path.getsize(path)
        except OSError as exc:
            raise UnsafeWorkbookError("The spreadsheet file could not be read.") from exc
        if size > cap:
            raise UnsafeWorkbookError("The spreadsheet file is too large.")
        with open(path, "rb") as fh:
            data = fh.read(cap + 1)
    else:
        stream = source if hasattr(source, "read") else getattr(source, "stream", source)
        try:
            position = stream.tell()
        except Exception:
            position = None
        try:
            data = stream.read(cap + 1)
        except TypeError:
            data = stream.read()
        if position is not None:
            try:
                stream.seek(position)
            except Exception:
                pass
        if isinstance(data, str):
            raise UnsafeWorkbookError("The spreadsheet file is not a valid .xlsx workbook.")
    if len(data) > cap:
        raise UnsafeWorkbookError("The spreadsheet file is too large.")
    if not data:
        raise UnsafeWorkbookError("The spreadsheet file is empty.")
    return data


def inspect_xlsx(data: bytes, limits: Optional[WorkbookLimits] = None) -> None:
    """Validate the zip container and declared sheet dimensions of ``data``.

    Raises :class:`UnsafeWorkbookError`; never parses workbook XML with openpyxl.
    """
    limits = limits or WorkbookLimits.from_config()
    if len(data) > limits.max_upload_bytes:
        raise UnsafeWorkbookError("The spreadsheet file is too large.")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, OSError, EOFError, ValueError) as exc:
        raise UnsafeWorkbookError("The spreadsheet file is not a valid .xlsx workbook.") from exc
    with zf:
        infos = zf.infolist()
        if len(infos) > limits.max_members:
            raise UnsafeWorkbookError("The spreadsheet contains too many internal files.")
        total = 0
        sheet_members = []
        for info in infos:
            if info.flag_bits & 0x1:
                raise UnsafeWorkbookError("Encrypted spreadsheets are not supported.")
            total += info.file_size
            if total > limits.max_uncompressed_bytes:
                raise UnsafeWorkbookError("The spreadsheet expands to too much data.")
            if (
                info.file_size >= limits.ratio_min_member_bytes
                and info.file_size > max(info.compress_size, 1) * limits.max_member_ratio
            ):
                raise UnsafeWorkbookError("The spreadsheet has an abnormal compression ratio.")
            if _SHEET_MEMBER_RE.match(info.filename):
                sheet_members.append(info)
        if len(sheet_members) > limits.max_sheets:
            raise UnsafeWorkbookError("The spreadsheet contains too many sheets.")
        for info in sheet_members:
            with zf.open(info) as member:
                head = member.read(_HEAD_BYTES)
            match = _DIMENSION_RE.search(head)
            if not match:
                continue
            first_col, first_row, last_col, last_row = match.groups()
            rows = int(last_row or first_row)
            cols = _col_to_index(last_col or first_col)
            if rows > limits.max_rows or cols > limits.max_cols:
                raise UnsafeWorkbookError("The spreadsheet has too many rows or columns.")


def load_workbook_safe(
    source: WorkbookSource,
    *,
    read_only: bool = True,
    data_only: bool = True,
    limits: Optional[WorkbookLimits] = None,
    **kwargs: Any,
):
    """Open an untrusted ``.xlsx`` with size/zip/dimension checks. Returns an openpyxl Workbook.

    ``source`` may be bytes, a filesystem path, a file-like object or a werkzeug
    ``FileStorage``. Pass ``read_only=False`` only when the caller must mutate cells
    or access random cell coordinates; dimensions are still bounded up front.
    """
    import openpyxl

    limits = limits or WorkbookLimits.from_config()
    data = _read_source_bytes(source, limits)
    inspect_xlsx(data, limits)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=read_only, data_only=data_only, **kwargs
        )
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise UnsafeWorkbookError("The spreadsheet file is not a valid .xlsx workbook.") from exc
    except Exception as exc:
        if exc.__class__.__name__ == "InvalidFileException":
            raise UnsafeWorkbookError("The spreadsheet file is not a valid .xlsx workbook.") from exc
        raise
    if len(workbook.sheetnames) > limits.max_sheets:
        workbook.close()
        raise UnsafeWorkbookError("The spreadsheet contains too many sheets.")
    return workbook


def read_excel_safe(
    source: WorkbookSource,
    *,
    limits: Optional[WorkbookLimits] = None,
    **kwargs: Any,
):
    """``pandas.read_excel`` behind the same guards; rows are capped via ``nrows``."""
    import pandas as pd

    limits = limits or WorkbookLimits.from_config()
    data = _read_source_bytes(source, limits)
    inspect_xlsx(data, limits)
    kwargs.setdefault("engine", "openpyxl")
    requested = kwargs.get("nrows")
    kwargs["nrows"] = limits.max_rows if requested is None else min(int(requested), limits.max_rows)
    try:
        return pd.read_excel(io.BytesIO(data), **kwargs)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise UnsafeWorkbookError("The spreadsheet file is not a valid .xlsx workbook.") from exc


def safe_iter_rows(worksheet, limits: Optional[WorkbookLimits] = None, **kwargs: Any) -> Iterator[Any]:
    """``worksheet.iter_rows`` that raises past ``limits.max_rows`` rows or ``limits.max_cols`` columns.

    Guards read-only sheets whose ``<dimension>`` is missing or lies.
    """
    limits = limits or WorkbookLimits.from_config()
    for count, row in enumerate(worksheet.iter_rows(**kwargs), start=1):
        if count > limits.max_rows or len(row) > limits.max_cols:
            raise UnsafeWorkbookError("The spreadsheet has too many rows or columns.")
        yield row
