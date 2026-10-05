"""Download external Everyone Counts sources into gitignored JSON snapshots.

Names from the security-unit register and Motiro survey rows are not saved.
OCAC output is the baseline volunteer-management score only, not the full workbook.
"""

from __future__ import annotations

import csv
import io
import json
import subprocess
import tempfile
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from plugins.fdrs.services.ecr_external import (
    MEDIAN_AGE_NAME,
    fill_log_gaps,
    interpolate_median_age,
    mechanism_cause,
    mechanism_group,
)
from plugins.fdrs.services.external_snapshots import save_snapshot
from plugins.fdrs.services.world_bank_snapshot import USER_AGENT, _iso3

ROOT = "https://api.worldbank.org/v2"
GO_ROOT = "https://goadmin.ifrc.org/api/v2/appeal/?format=json&limit=500"
ILO_MIN = "https://rplumber.ilo.org/data/indicator/?id=EAR_INEE_CUR_NB_A&lang=en"
ILO_MEAN = "https://rplumber.ilo.org/data/indicator/?id=EAR_EMTA_SEX_CUR_NB_A&lang=en"
REPO = "stevepowell99/ECR2026-data-and-code"
AGE_BANDS = (
    "0004", "0509", "1014", "1519", "2024", "2529", "3034", "3539",
    "4044", "4549", "5054", "5559", "6064", "6569", "7074", "7579", "80UP",
)
AGE_SPEC = tuple(
    (band, lower, 10 if band == "80UP" else 5)
    for band, lower in zip(AGE_BANDS, list(range(0, 80, 5)) + [80])
)
NAME_ALIASES = {
    "united states of america": "USA",
    "united states": "USA",
    "united kingdom": "GBR",
    "uk": "GBR",
    "czech republic": "CZE",
    "czechia": "CZE",
    "myanmar": "MMR",
    "burma": "MMR",
    "united arab emirates": "ARE",
    "russia": "RUS",
    "russian federation": "RUS",
    "south korea": "KOR",
    "korea, republic of": "KOR",
    "republic of korea": "KOR",
    "north korea": "PRK",
    "vietnam": "VNM",
    "viet nam": "VNM",
    "laos": "LAO",
    "lao pdr": "LAO",
    "syria": "SYR",
    "syrian arab republic": "SYR",
    "iran": "IRN",
    "iran, islamic republic of": "IRN",
    "bolivia": "BOL",
    "bolivia (plurinational state of)": "BOL",
    "venezuela": "VEN",
    "venezuela (bolivarian republic of)": "VEN",
    "tanzania": "TZA",
    "united republic of tanzania": "TZA",
    "dr congo": "COD",
    "democratic republic of the congo": "COD",
    "congo, dem. rep.": "COD",
    "congo, democratic republic of the": "COD",
    "republic of the congo": "COG",
    "congo": "COG",
    "ivory coast": "CIV",
    "cote d'ivoire": "CIV",
    "côte d'ivoire": "CIV",
    "palestine": "PSE",
    "west bank and gaza": "PSE",
    "taiwan": "TWN",
    "hong kong": "HKG",
    "hong kong sar, china": "HKG",
    "macao": "MAC",
    "swaziland": "SWZ",
    "eswatini": "SWZ",
    "cape verde": "CPV",
    "cabo verde": "CPV",
    "kyrgyzstan": "KGZ",
    "slovak republic": "SVK",
    "slovakia": "SVK",
    "türkiye": "TUR",
    "turkey": "TUR",
    "moldova": "MDA",
    "republic of moldova": "MDA",
    "gambia": "GMB",
    "the gambia": "GMB",
    "bahamas": "BHS",
    "the bahamas": "BHS",
    "central africa republic": "CAF",
    "central african republic": "CAF",
    "columbia": "COL",
    "colombia": "COL",
    "congo brazzaville": "COG",
    "egypt": "EGY",
    "egypt arab rep": "EGY",
    "state of libya": "LBY",
    "libya": "LBY",
    "palestinian territories": "PSE",
    "palestinian territory": "PSE",
    "state of palestine": "PSE",
    "puerto rico": "PRI",
    "somalia": "SOM",
    "trinidad and tobago": "TTO",
    "yemen": "YEM",
    "yemen rep": "YEM",
}


def _get(url: str, timeout: int = 90) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read()


def _json(url: str) -> dict | list:
    return json.loads(_get(url).decode("utf-8"))


def _gh(path: str) -> bytes:
    return subprocess.check_output(
        ["gh", "api", "-H", "Accept: application/vnd.github.raw", f"repos/{REPO}/contents/{path}"],
    )


def _normalise_name(value: str) -> str:
    text = value.lower().replace("’", "'").replace(".", "").replace(",", "").replace("&", " and ")
    return " ".join(text.split())


def _country_maps() -> tuple[dict[str, str], dict[str, str]]:
    payload = _json(f"{ROOT}/country?format=json&per_page=400")
    rows = payload[1] if isinstance(payload, list) and len(payload) > 1 else []
    by_name = {}
    by_iso2 = {}
    for row in rows:
        iso3 = _iso3(row.get("id"))
        iso2 = str(row.get("iso2Code") or "").strip().upper()
        if not iso3 or str((row.get("region") or {}).get("id") or "") in {"", "NA"}:
            continue
        by_name[_normalise_name(str(row.get("name") or ""))] = iso3
        if len(iso2) == 2:
            by_iso2[iso2] = iso3
    by_name.update({_normalise_name(name): iso3 for name, iso3 in NAME_ALIASES.items()})
    return by_name, by_iso2


def build_go() -> None:
    appeals = []
    url = GO_ROOT
    while url:
        page = _json(url)
        for row in page.get("results") or []:
            country = row.get("country") or {}
            dtype = row.get("dtype") or {}
            start = str(row.get("start_date") or "")
            year = int(start[:4]) if len(start) >= 4 and start[:4].isdigit() else None
            iso3 = _iso3(country.get("iso3"))
            if not iso3 or year is None:
                continue
            atype = row.get("atype_display") or row.get("atype")
            if isinstance(atype, dict):
                atype = atype.get("display") or atype.get("name") or ""
            appeals.append(
                {
                    "iso3": iso3,
                    "country_name": str(country.get("name") or ""),
                    "year": year,
                    "atype": str(atype or ""),
                    "dtype": str(dtype.get("name") or "") if isinstance(dtype, dict) else str(dtype or ""),
                }
            )
        url = page.get("next")
        print(f"  GO page, {len(appeals)} appeals", flush=True)
    save_snapshot("go_appeals", {"source": "https://goadmin.ifrc.org/api/v2/appeal/", "appeals": appeals})
    print(f"GO appeals: {len(appeals)}")


def build_wgi(by_name: dict[str, str]) -> None:
    years = [2010, 2011, 2012, 2013, 2014, 2015, 2016, 2017, 2018, 2021, 2022]
    rows = []
    unmatched = set()
    for year in years:
        raw = _gh(f"raw-data/external/WGI (1)/{year}.csv").decode("utf-8-sig")
        for record in csv.DictReader(io.StringIO(raw), delimiter=";"):
            country = record.get("Country") or ""
            iso3 = by_name.get(_normalise_name(country))
            if not iso3:
                unmatched.add(country)
                continue

            def score(label: str) -> float | None:
                text = str(record.get(label) or "").replace("%", "").strip()
                try:
                    return float(text)
                except ValueError:
                    return None

            rows.append(
                {
                    "iso3": iso3,
                    "year": year,
                    "helping": score("Helping a stranger Score"),
                    "donating": score("Donating money Score"),
                    "volunteering": score("Volunteering time Score"),
                }
            )
    save_snapshot(
        "world_giving_index",
        {
            "source": "Charities Aid Foundation World Giving Index country files, 2010-2018 and 2021-2022",
            "rows": rows,
        },
    )
    print(f"WGI rows: {len(rows)}; unmatched: {sorted(unmatched)}")


def build_median_age() -> None:
    totals: dict[tuple[str, int, str], float] = defaultdict(float)
    for band in AGE_BANDS:
        for sex in ("MA", "FE"):
            url = f"{ROOT}/country/all/indicator/SP.POP.{band}.{sex}?format=json&per_page=20000&date=2010:2024"
            payload = _json(url)
            for row in payload[1] if isinstance(payload, list) and len(payload) > 1 else []:
                iso3 = _iso3(row.get("countryiso3code"))
                try:
                    year = int(row.get("date"))
                    value = float(row.get("value"))
                except (TypeError, ValueError):
                    continue
                if iso3 and value > 0:
                    totals[(iso3, year, band)] += value
            print(f"  age {band} {sex}", flush=True)
    grouped: dict[tuple[str, int], dict[str, float]] = defaultdict(dict)
    for (iso3, year, band), value in totals.items():
        grouped[(iso3, year)][band] = value
    rows = []
    for (iso3, year), bands in grouped.items():
        if len(bands) < 10:
            continue
        age = interpolate_median_age(bands, AGE_SPEC)
        if age is None:
            continue
        rows.append({"iso3": iso3, "year": year, "median_age": round(age, 2)})
    rows.sort(key=lambda item: (item["iso3"], item["year"]))
    save_snapshot(
        MEDIAN_AGE_NAME,
        {"source": "World Bank population by five-year age band, SP.POP.<band>.MA and .FE", "rows": rows},
    )
    print(f"Median age rows: {len(rows)}")


def build_awsd() -> None:
    raw = _gh("raw-data/external/aid_worker_security/AWSD_global_security_incidents.csv").decode("utf-8-sig")
    counts: dict[tuple[int, str], int] = defaultdict(int)
    for record in csv.DictReader(io.StringIO(raw)):
        try:
            year = int(record.get("Year") or "")
            killed = int(float(record.get("Total killed") or 0))
        except ValueError:
            continue
        if year < 2016 or year > 2025 or killed <= 0:
            continue
        movement = float(record.get("ICRC") or 0) + float(record.get("NRCS and IFRC") or 0)
        group = "movement" if movement > 0 else "other"
        counts[(year, group)] += killed
    rows = [{"year": year, "group": group, "killed": killed} for (year, group), killed in sorted(counts.items())]
    save_snapshot(
        "awsd_killed",
        {
            "source": "Humanitarian Outcomes Aid Worker Security Database, yearly totals only",
            "rows": rows,
        },
    )
    print(f"AWSD year rows: {len(rows)}")


def _header_map(sheet) -> dict[str, int]:
    headers = {}
    for index, cell in enumerate(next(sheet.iter_rows(max_row=1, values_only=True))):
        headers[str(cell or "").strip().lower()] = index
    return headers


def build_deaths() -> None:
    destination = Path(tempfile.gettempdir()) / "ecr_deaths_source.xlsx"
    destination.write_bytes(_gh("processed-data/inputs/RCRC Deaths 2012-2026.xlsx"))
    workbook = load_workbook(destination, read_only=True, data_only=True)
    sheet = workbook["DATA"]
    headers = _header_map(sheet)
    counts: dict[tuple[int, str, str, str], int] = defaultdict(int)
    year_fill = None
    for values in sheet.iter_rows(min_row=2, values_only=True):
        def cell(name: str):
            index = headers.get(name)
            return values[index] if index is not None and index < len(values) else None

        raw_year = cell("year")
        if isinstance(raw_year, (int, float)):
            year_fill = int(raw_year)
        date = cell("date")
        year = None
        if hasattr(date, "year"):
            year = int(date.year)
        elif year_fill is not None:
            year = year_fill
        if year is None or not cell("country"):
            continue
        killed = cell("killed")
        injured = cell("injured")
        killed_n = int(killed or 0) if isinstance(killed, (int, float)) else 0
        injured_n = int(injured or 0) if isinstance(injured, (int, float)) else 0
        if killed_n + injured_n <= 0:
            continue
        role_raw = str(cell("staff / volunteer") or "").strip()
        if role_raw in {"Staff", "Volunteer"}:
            role = role_raw
        elif role_raw == "Mixed (see names)":
            role = "Mixed"
        else:
            role = "Unknown"
        mechanism = mechanism_group(cell("mechanism"))
        counts[(year, role, mechanism, mechanism_cause(mechanism))] += killed_n
    workbook.close()
    rows = [
        {"year": year, "role": role, "mechanism": mechanism, "cause": cause, "killed": killed}
        for (year, role, mechanism, cause), killed in sorted(counts.items())
        if killed
    ]
    save_snapshot(
        "rcrc_deaths_aggregate",
        {"source": "IFRC security-unit register, counts only; names are not stored", "rows": rows},
    )
    try:
        destination.unlink(missing_ok=True)
    except PermissionError:
        print("Deaths source file is still open; it is in the temp folder and was not copied into the databank.")
    print(f"Death aggregate rows: {len(rows)}")


def build_ocac(by_iso2: dict[str, str]) -> None:
    # The export's dimension record is 1x1, so read_only mode drops the assessment columns.
    workbook = load_workbook(io.BytesIO(_gh("raw-data/OCAC_10062026135045.xlsx")), data_only=True)
    sheet = workbook.active
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    code_col = [row[0] if row else None for row in grid]
    name_col = [row[1] if row and len(row) > 1 else None for row in grid]

    def find_row(predicate) -> list:
        for index, (code, name) in enumerate(zip(code_col, name_col)):
            if predicate(code, name):
                return list(grid[index][2:])
        return []

    iso2s = [str(value or "").strip().upper() for value in find_row(lambda _code, name: name == "iso")]
    years = []
    for value in find_row(lambda _code, name: name == "Year"):
        try:
            years.append(int(value))
        except (TypeError, ValueError):
            years.append(None)

    def attribute(number: int) -> list[float | None]:
        raw = find_row(lambda code, _name: str(code or "").strip() == f"OCAC{number}")
        parsed = []
        for value in raw:
            letter = str(value or "").strip().upper()
            parsed.append({"A": 1, "B": 2, "C": 3, "D": 4, "E": 5}.get(letter))
        return parsed

    recruit = attribute(65)
    policy = attribute(66)
    best: dict[str, tuple[int, float]] = {}
    width = min(len(iso2s), len(years), len(recruit), len(policy))
    for index in range(width):
        iso3 = by_iso2.get(iso2s[index])
        year = years[index]
        scores = [value for value in (recruit[index], policy[index]) if value is not None]
        if not iso3 or year is None or not scores:
            continue
        mean = sum(scores) / len(scores)
        current = best.get(iso3)
        if current is None or year < current[0]:
            best[iso3] = (year, mean)
    rows = [{"iso3": iso3, "year": year, "volmgmt": round(score, 3)} for iso3, (year, score) in sorted(best.items())]
    save_snapshot(
        "ocac_volunteer_management",
        {
            "source": "OCAC baseline, mean of volunteer recruitment/retention and volunteering strategy only",
            "rows": rows,
        },
    )
    print(f"OCAC societies: {len(rows)}")


def _ilo_rows(url: str, sex_total: bool) -> dict[tuple[str, str], tuple[int, float]]:
    raw = _get(url, timeout=120).decode("utf-8-sig")
    latest: dict[tuple[str, str], tuple[int, float]] = {}
    for record in csv.DictReader(io.StringIO(raw)):
        if sex_total and record.get("sex") not in {None, "", "SEX_T"} and "sex" in record:
            continue
        currency = record.get("classif1") or ""
        if currency not in {"CUR_TYPE_USD", "CUR_TYPE_PPP"}:
            continue
        iso3 = _iso3(record.get("ref_area"))
        try:
            year = int(str(record.get("time") or "")[:4])
            value = float(record.get("obs_value"))
        except (TypeError, ValueError):
            continue
        if not iso3 or value <= 0:
            continue
        key = (iso3, "usd" if currency == "CUR_TYPE_USD" else "ppp")
        current = latest.get(key)
        if current is None or year > current[0]:
            latest[key] = (year, value)
    return latest


def build_wages() -> None:
    minimum = _ilo_rows(ILO_MIN, sex_total=False)
    average = _ilo_rows(ILO_MEAN, sex_total=True)
    gni_ppp_raw = _json(f"{ROOT}/country/all/indicator/NY.GNP.PCAP.PP.CD?format=json&per_page=20000&date=2015:2025")
    gni_ppp_year: dict[str, tuple[int, float]] = {}
    for row in gni_ppp_raw[1] if isinstance(gni_ppp_raw, list) and len(gni_ppp_raw) > 1 else []:
        iso3 = _iso3(row.get("countryiso3code"))
        try:
            year = int(row.get("date"))
            value = float(row.get("value"))
        except (TypeError, ValueError):
            continue
        if iso3 and value > 0:
            current = gni_ppp_year.get(iso3)
            if current is None or year >= current[0]:
                gni_ppp_year[iso3] = (year, value)
    gni_ppp = {iso3: value for iso3, (_year, value) in gni_ppp_year.items()}
    from plugins.fdrs.services.world_bank_snapshot import SNAPSHOT_NAME, load_snapshot

    world_bank = load_snapshot(SNAPSHOT_NAME) or {}
    gni_usd = {
        row.get("iso3"): row.get("gni_per_capita_usd")
        for row in world_bank.get("gni_per_capita") or []
    }
    countries = set(gni_ppp) | set(gni_usd) | {iso3 for iso3, _cur in minimum} | {iso3 for iso3, _cur in average}
    rows = []
    for iso3 in sorted(countries):
        if not iso3:
            continue
        rows.append(
            {
                "iso3": iso3,
                "min_usd": (minimum.get((iso3, "usd")) or (None, None))[1],
                "min_ppp": (minimum.get((iso3, "ppp")) or (None, None))[1],
                "mean_usd": (average.get((iso3, "usd")) or (None, None))[1],
                "mean_ppp": (average.get((iso3, "ppp")) or (None, None))[1],
                "gni_usd": gni_usd.get(iso3),
                "gni_ppp": gni_ppp.get(iso3),
            }
        )
    for value_key, predictor in (
        ("min_usd", "gni_usd"),
        ("min_ppp", "gni_ppp"),
        ("mean_usd", "gni_usd"),
        ("mean_ppp", "gni_ppp"),
    ):
        fill_log_gaps(rows, value_key, predictor)
    saved = []
    for row in rows:
        row["min_imputed"] = bool(row.get("min_usd_imputed") or row.get("min_ppp_imputed"))
        row["mean_imputed"] = bool(row.get("mean_usd_imputed") or row.get("mean_ppp_imputed"))
        if row.get("min_usd") or row.get("mean_usd"):
            saved.append(
                {
                    "iso3": row["iso3"],
                    "min_usd": row.get("min_usd"),
                    "min_ppp": row.get("min_ppp"),
                    "mean_usd": row.get("mean_usd"),
                    "mean_ppp": row.get("mean_ppp"),
                    "min_imputed": row["min_imputed"],
                    "mean_imputed": row["mean_imputed"],
                }
            )
    save_snapshot("ilo_wages", {"source": "ILOSTAT monthly wages, gaps filled from GNI", "rows": saved})
    print(f"ILO wage countries: {len(saved)}")


def build_hours() -> None:
    raw = _gh("processed-data/inputs/ifrc_baseline_volunteer_hours_americas.csv").decode("utf-8-sig")
    rows = []
    for record in csv.DictReader(io.StringIO(raw)):
        iso3 = _iso3(record.get("iso_3"))
        try:
            hours_mid = float(record.get("hours_mid") or "")
        except ValueError:
            continue
        if not iso3:
            continue
        rows.append(
            {
                "iso3": iso3,
                "country": record.get("country") or iso3,
                "hours_bin": record.get("hours_bin") or "",
                "hours_mid": hours_mid,
            }
        )
    save_snapshot(
        "americas_volunteer_hours",
        {"source": "IFRC Americas volunteer-hours study, monthly hours band midpoints", "rows": rows},
    )
    print(f"Hours societies: {len(rows)}")


def main() -> None:
    print(datetime.now(timezone.utc).isoformat(), "country maps", flush=True)
    by_name, by_iso2 = _country_maps()
    print("GO", flush=True)
    build_go()
    print("WGI", flush=True)
    build_wgi(by_name)
    print("median age", flush=True)
    build_median_age()
    print("AWSD", flush=True)
    build_awsd()
    print("deaths aggregate", flush=True)
    build_deaths()
    print("OCAC", flush=True)
    build_ocac(by_iso2)
    print("wages", flush=True)
    build_wages()
    print("hours", flush=True)
    build_hours()
    print("done")


if __name__ == "__main__":
    main()
