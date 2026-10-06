"""Round-trip helpers for the PNS Planning (T22) and PNS Reporting (T23) workbooks.

The workbooks are the IFRC PNS collection forms. Values the assignment form
stores are written into the data tables; display sheets keep their formulas.
Sample rows for other National Societies are cleared on export.
"""

from __future__ import annotations

import os
import re
import sys
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

_HERE = os.path.abspath(__file__)
_SCRIPT_DIR = os.path.dirname(_HERE)
_BACKOFFICE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(_SCRIPT_DIR)))
_CORE_IMPORTS = os.path.join(_BACKOFFICE_DIR, "scripts", "imports")
for _path in (_BACKOFFICE_DIR, _SCRIPT_DIR, _CORE_IMPORTS):
    if _path not in sys.path:
        sys.path.insert(0, _path)

from import_upr_excel_data import (  # noqa: E402
    STAFF_INDICATOR_COLUMNS,
    T22_BREAKDOWN_AREAS,
    T22_ROW_TOTAL_COLUMN,
)

PLANNING_FUNDING_SHEET = "Funding Requirement"
PLANNING_DATA_SHEET = "FR data"
PLANNING_DATA_TABLE = "PNS_data"
ACTIVITIES_SHEET = "Activities data"
STAFF_SHEET = "PNS Staff Presence"
REPORTING_FUNDING_SHEET = "Funding data"
REPORTING_DATA_TABLE = "Data_PNS"
REPORTING_REFERENCE_SHEET = "Funding Requirement by Country"
REPORTING_PLAN_SHEET = "FR"

PLANNING_NS_CELL = "C3"
REPORTING_NS_CELL = "D3"
PLANNING_YEAR_HEADER_CELLS = ("D8", "L8", "T8")
# First-year block on Funding Requirement. Attribute codes match the SUMIFS
# literals (column D is SP1 even though its header text says Disasters).
PLANNING_GRID_COLUMNS = {
    "D": "SP1",
    "E": "SP2",
    "F": "SP3",
    "G": "SP4",
    "H": "SP5",
    "I": "EFs",
    "J": T22_ROW_TOTAL_COLUMN,
}

REPORTING_EXCEL_COLUMNS = {
    "Funding": ("total funding", "funding"),
    "Expenditure": ("total expenditure", "expenditure"),
    "Transferred": ("total transferred to hns", "transferred", "total transferred"),
    "SP1": ("sp1",),
    "SP2": ("sp2",),
    "SP3": ("sp3",),
    "SP4": ("sp4",),
    "SP5": ("sp5",),
    "EFs": ("efs", "ef", "enabling functions"),
    "Comments": ("comments", "comments/notes", "notes"),
}

_YEAR_LITERALS = ("2026", "2027", "2028")


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _as_number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, dict):
        raw = value.get("modified")
        if raw in ("", None) and value.get("isModified"):
            return None
        if raw in ("", None):
            raw = value.get("original")
        return _as_number(raw)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text or text in {"-", "—"}:
            return None
        try:
            value = float(text)
        except ValueError:
            return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _as_int_year(value: Any) -> Optional[int]:
    number = _as_number(value)
    if number is None:
        return None
    year = int(number)
    if 1990 <= year <= 2100:
        return year
    return None


def _json_number(value: float) -> int | float:
    if float(value).is_integer():
        return int(value)
    return float(value)


def _is_formula(value: Any) -> bool:
    if isinstance(value, str) and value.startswith("="):
        return True
    return type(value).__name__ == "ArrayFormula"


def _clear_constants(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.value is None or _is_formula(cell.value):
                continue
            cell.value = None


def _header_map(ws, row: int, max_col: Optional[int] = None) -> Dict[str, int]:
    found: Dict[str, int] = {}
    limit = max_col or (ws.max_column or 1)
    for col in range(1, limit + 1):
        key = _norm(ws.cell(row, col).value)
        if key and key not in found:
            found[key] = col
    return found


def rewrite_planning_year_formula(formula: Any, base_year: int) -> Any:
    """Point hardcoded PNS_data year literals at the assignment year and the two that follow."""
    if not isinstance(formula, str) or "PNS_data[Year]" not in formula:
        return formula
    placeholders = {
        f'"{_YEAR_LITERALS[2]}"': "__PNS_YEAR_2__",
        f'"{_YEAR_LITERALS[1]}"': "__PNS_YEAR_1__",
        f'"{_YEAR_LITERALS[0]}"': "__PNS_YEAR_0__",
    }
    updated = formula
    for literal, token in placeholders.items():
        updated = updated.replace(literal, token)
    updated = (
        updated.replace("__PNS_YEAR_0__", f'"{base_year}"')
        .replace("__PNS_YEAR_1__", f'"{base_year + 1}"')
        .replace("__PNS_YEAR_2__", f'"{base_year + 2}"')
    )
    return updated


def point_planning_selected_name(wb) -> None:
    """The template's ``selected`` name is ``#REF!``; formulas filter on it."""
    target = f"'{PLANNING_FUNDING_SHEET}'!${PLANNING_NS_CELL[0]}${PLANNING_NS_CELL[1:]}"
    defn = wb.defined_names.get("selected")
    if defn is not None:
        defn.attr_text = target
        return
    from openpyxl.workbook.defined_name import DefinedName

    wb.defined_names.add(DefinedName("selected", attr_text=target))


def sheet_country_index(wb) -> Dict[str, Dict[str, str]]:
    """Map a country display name to ISO3 and region using the staff list."""
    if STAFF_SHEET not in wb.sheetnames:
        return {}
    ws = wb[STAFF_SHEET]
    headers = _header_map(ws, 1)
    iso_col = headers.get("iso3")
    name_col = headers.get("country")
    region_col = headers.get("region")
    if not iso_col or not name_col:
        return {}
    index: Dict[str, Dict[str, str]] = {}
    for row in range(2, (ws.max_row or 1) + 1):
        iso3 = str(ws.cell(row, iso_col).value or "").strip().upper()
        name = str(ws.cell(row, name_col).value or "").strip()
        if not iso3 or not name:
            continue
        region = str(ws.cell(row, region_col).value or "").strip() if region_col else ""
        index[_norm(name)] = {"iso3": iso3, "country": name, "region": region}
    return index


def read_planning_data_rows(wb) -> List[Dict[str, Any]]:
    ws = wb[PLANNING_DATA_SHEET]
    headers = _header_map(ws, 1, 5)
    ns_col = headers.get("national society name")
    country_col = headers.get("country")
    year_col = headers.get("year")
    attr_col = headers.get("attribute")
    value_col = headers.get("value")
    rows: List[Dict[str, Any]] = []
    if not all((ns_col, country_col, year_col, attr_col, value_col)):
        return rows
    for row in range(2, (ws.max_row or 1) + 1):
        ns_name = str(ws.cell(row, ns_col).value or "").strip()
        country = str(ws.cell(row, country_col).value or "").strip()
        year = _as_int_year(ws.cell(row, year_col).value)
        attribute = str(ws.cell(row, attr_col).value or "").strip()
        value = _as_number(ws.cell(row, value_col).value)
        if not ns_name or not country or year is None or not attribute or value is None:
            continue
        rows.append(
            {
                "ns": ns_name,
                "country": country,
                "year": year,
                "attribute": "EFs" if _norm(attribute) == "ef" else attribute.strip(),
                "value": value,
            }
        )
    return rows


def read_planning_grid_overrides(wb) -> List[Dict[str, Any]]:
    """Numbers typed over the first-year formulas, used when FR data is empty."""
    if PLANNING_FUNDING_SHEET not in wb.sheetnames:
        return []
    ws = wb[PLANNING_FUNDING_SHEET]
    ns_name = str(ws[PLANNING_NS_CELL].value or "").strip()
    year = _as_int_year(ws["D8"].value)
    rows: List[Dict[str, Any]] = []
    for row in range(12, (ws.max_row or 11) + 1):
        country = str(ws.cell(row, 3).value or "").strip()
        if not country or _is_formula(ws.cell(row, 3).value):
            continue
        for column, attribute in PLANNING_GRID_COLUMNS.items():
            cell = ws[f"{column}{row}"]
            if _is_formula(cell.value):
                continue
            value = _as_number(cell.value)
            if value is None:
                continue
            rows.append(
                {
                    "ns": ns_name,
                    "country": country,
                    "year": year,
                    "attribute": attribute,
                    "value": value,
                }
            )
    return rows


def planning_funding_cells(
    rows: Sequence[Mapping[str, Any]],
    *,
    ns_name: str,
    year: int,
    country_name_to_id: Mapping[str, int],
    form_columns: Optional[Sequence[str]] = None,
) -> Tuple[Dict[str, Any], List[str]]:
    warnings: List[str] = []
    cells: Dict[str, Any] = {}
    allowed = {name.strip().lower(): name for name in (form_columns or list(T22_BREAKDOWN_AREAS) + [T22_ROW_TOTAL_COLUMN])}
    ns_key = _norm(ns_name)
    for row in rows:
        if ns_key and _norm(row.get("ns")) not in {ns_key, ""}:
            continue
        if row.get("year") != year:
            continue
        attribute = str(row.get("attribute") or "").strip()
        column = allowed.get(attribute.lower())
        if not column:
            continue
        country_id = country_name_to_id.get(_norm(row.get("country")))
        if not country_id:
            warnings.append(f"No country match for planning funding row {row.get('country')!r}.")
            continue
        cells[f"{country_id}_{column}"] = _json_number(float(row["value"]))
    return cells, warnings


def read_staff_rows(wb) -> List[Dict[str, Any]]:
    if STAFF_SHEET not in wb.sheetnames:
        return []
    ws = wb[STAFF_SHEET]
    headers = _header_map(ws, 1)
    iso_col = headers.get("iso3")
    if not iso_col:
        return []
    column_by_header = {_norm(header): key for header, key in STAFF_INDICATOR_COLUMNS.items()}
    excel_columns = {
        col: column_by_header[header]
        for header, col in headers.items()
        if header in column_by_header
    }
    rows: List[Dict[str, Any]] = []
    for row in range(2, (ws.max_row or 1) + 1):
        iso3 = str(ws.cell(row, iso_col).value or "").strip().upper()
        if not iso3:
            continue
        values = {}
        for col, key in excel_columns.items():
            number = _as_number(ws.cell(row, col).value)
            if number is not None:
                values[key] = number
        if values:
            rows.append({"iso3": iso3, "values": values})
    return rows


def staff_cells(
    rows: Sequence[Mapping[str, Any]],
    *,
    iso3_to_ns_id: Mapping[str, int],
    form_columns: Optional[Sequence[str]] = None,
) -> Tuple[Dict[str, Any], List[str]]:
    warnings: List[str] = []
    cells: Dict[str, Any] = {}
    allowed = {name.strip().lower(): name for name in (form_columns or STAFF_INDICATOR_COLUMNS.values())}
    for row in rows:
        ns_id = iso3_to_ns_id.get(str(row.get("iso3") or "").upper())
        if not ns_id:
            warnings.append(f"No National Society for staff row {row.get('iso3')!r}.")
            continue
        for key, value in (row.get("values") or {}).items():
            column = allowed.get(str(key).lower())
            if not column:
                continue
            cells[f"{ns_id}_{column}"] = _json_number(float(value))
    return cells, warnings


def reporting_column_map(form_columns: Sequence[str]) -> Dict[str, str]:
    """Map Data_PNS header (Funding, SP1, ...) to the published matrix column name."""
    available = {name.strip().lower(): name for name in form_columns}
    mapped: Dict[str, str] = {}
    for excel_header, needles in REPORTING_EXCEL_COLUMNS.items():
        for needle in needles:
            if needle in available:
                mapped[excel_header] = available[needle]
                break
    return mapped


def read_reporting_funding_rows(wb) -> List[Dict[str, Any]]:
    ws = wb[REPORTING_FUNDING_SHEET]
    if REPORTING_DATA_TABLE not in ws.tables:
        return []
    headers = _header_map(ws, 14, 14)
    country_col = headers.get("country")
    if not country_col:
        return []
    value_columns = {
        header: col
        for header, col in headers.items()
        if header in {_norm(name) for name in REPORTING_EXCEL_COLUMNS}
    }
    header_by_norm = {_norm(name): name for name in REPORTING_EXCEL_COLUMNS}
    rows: List[Dict[str, Any]] = []
    for row in range(15, 71):
        country_cell = ws.cell(row, country_col)
        if _is_formula(country_cell.value):
            continue
        country = str(country_cell.value or "").strip()
        if not country:
            continue
        values = {}
        for header, col in value_columns.items():
            cell = ws.cell(row, col)
            if _is_formula(cell.value):
                continue
            number = _as_number(cell.value)
            text = None if number is not None else str(cell.value or "").strip()
            excel_name = header_by_norm.get(header)
            if not excel_name:
                continue
            if number is not None:
                values[excel_name] = number
            elif excel_name == "Comments" and text:
                values[excel_name] = text
        if values:
            rows.append({"country": country, "values": values})
    return rows


def reporting_funding_cells(
    rows: Sequence[Mapping[str, Any]],
    *,
    country_name_to_ns_id: Mapping[str, int],
    column_map: Mapping[str, str],
) -> Tuple[Dict[str, Any], List[str]]:
    warnings: List[str] = []
    cells: Dict[str, Any] = {}
    for row in rows:
        ns_id = country_name_to_ns_id.get(_norm(row.get("country")))
        if not ns_id:
            warnings.append(f"No National Society for reporting country {row.get('country')!r}.")
            continue
        for excel_name, value in (row.get("values") or {}).items():
            column = column_map.get(excel_name)
            if not column:
                continue
            if isinstance(value, str):
                cells[f"{ns_id}_{column}"] = value
            else:
                cells[f"{ns_id}_{column}"] = _json_number(float(value))
    return cells, warnings


def _matrix_pairs(cells: Mapping[str, Any], form_columns: Iterable[str]) -> List[Tuple[str, str, float]]:
    columns = {name.strip().lower(): name for name in form_columns}
    pairs: List[Tuple[str, str, float]] = []
    for key, raw in cells.items():
        if not isinstance(key, str) or "_" not in key or key.startswith("_"):
            continue
        row_id, column = key.split("_", 1)
        if column.strip().lower() not in columns:
            continue
        number = _as_number(raw)
        if number is None:
            continue
        pairs.append((row_id, columns[column.strip().lower()], number))
    return pairs


def prepare_planning_workbook(wb, *, ns_name: str, year: Optional[int]) -> None:
    point_planning_selected_name(wb)
    funding = wb[PLANNING_FUNDING_SHEET]
    funding[PLANNING_NS_CELL] = ns_name
    if year is not None:
        for offset, coord in enumerate(PLANNING_YEAR_HEADER_CELLS):
            funding[coord] = year + offset
        title = funding["B1"].value
        if isinstance(title, str):
            updated = re.sub(r"\d{4}\s*[-–]\s*\d{4}", f"{year}-{year + 2}", title)
            funding["B1"] = updated
        for row in funding.iter_rows(min_row=1, max_row=funding.max_row or 1, max_col=funding.max_column or 1):
            for cell in row:
                cell.value = rewrite_planning_year_formula(cell.value, year)
    data = wb[PLANNING_DATA_SHEET]
    _clear_constants(data, 2, data.max_row or 2, 1, 5)
    if ACTIVITIES_SHEET in wb.sheetnames:
        activities = wb[ACTIVITIES_SHEET]
        _clear_constants(activities, 2, activities.max_row or 2, 1, activities.max_column or 1)
    _clear_staff_values(wb)


def write_planning_funding_rows(
    wb,
    cells: Mapping[str, Any],
    *,
    ns_name: str,
    year: int,
    country_id_to_name: Mapping[int, str],
    form_columns: Sequence[str],
) -> None:
    ws = wb[PLANNING_DATA_SHEET]
    row_number = 2
    for row_id, column, value in _matrix_pairs(cells, form_columns):
        if column == T22_ROW_TOTAL_COLUMN:
            continue
        country = country_id_to_name.get(int(row_id))
        if not country:
            continue
        ws.cell(row_number, 1).value = ns_name
        ws.cell(row_number, 2).value = country
        ws.cell(row_number, 3).value = year
        ws.cell(row_number, 4).value = column
        ws.cell(row_number, 5).value = _json_number(value)
        row_number += 1


def _clear_staff_values(wb) -> None:
    if STAFF_SHEET not in wb.sheetnames:
        return
    ws = wb[STAFF_SHEET]
    headers = _header_map(ws, 1)
    columns = [
        col
        for header, col in headers.items()
        if header in {_norm(name) for name in STAFF_INDICATOR_COLUMNS}
    ]
    if not columns:
        return
    _clear_constants(ws, 2, ws.max_row or 2, min(columns), max(columns))


def write_staff_rows(
    wb,
    cells: Mapping[str, Any],
    *,
    ns_id_to_iso3: Mapping[int, str],
    form_columns: Sequence[str],
) -> None:
    if STAFF_SHEET not in wb.sheetnames:
        return
    ws = wb[STAFF_SHEET]
    headers = _header_map(ws, 1)
    iso_col = headers.get("iso3")
    if not iso_col:
        return
    column_by_header = {_norm(header): key for header, key in STAFF_INDICATOR_COLUMNS.items()}
    excel_col_by_key = {}
    allowed = {name.lower(): name for name in form_columns}
    for header, col in headers.items():
        key = column_by_header.get(header)
        if key and key.lower() in allowed:
            excel_col_by_key[allowed[key.lower()]] = col
    iso_rows = {}
    for row in range(2, (ws.max_row or 1) + 1):
        iso3 = str(ws.cell(row, iso_col).value or "").strip().upper()
        if iso3 and iso3 not in iso_rows:
            iso_rows[iso3] = row
    for row_id, column, value in _matrix_pairs(cells, form_columns):
        iso3 = ns_id_to_iso3.get(int(row_id))
        excel_row = iso_rows.get((iso3 or "").upper())
        excel_col = excel_col_by_key.get(column)
        if not excel_row or not excel_col:
            continue
        ws.cell(excel_row, excel_col).value = _json_number(value)


def prepare_reporting_workbook(wb, *, ns_name: str) -> None:
    funding = wb[REPORTING_FUNDING_SHEET]
    funding[REPORTING_NS_CELL] = ns_name
    headers = _header_map(funding, 14, 14)
    input_cols = [
        col
        for header, col in headers.items()
        if header in {_norm(name) for name in REPORTING_EXCEL_COLUMNS} or header in {"region", "country"}
    ]
    if input_cols:
        _clear_constants(funding, 15, 70, min(input_cols), max(input_cols))
    _clear_staff_values(wb)
    if REPORTING_PLAN_SHEET in wb.sheetnames:
        plan = wb[REPORTING_PLAN_SHEET]
        _clear_constants(plan, 2, plan.max_row or 2, 1, 4)
    if REPORTING_REFERENCE_SHEET in wb.sheetnames:
        reference = wb[REPORTING_REFERENCE_SHEET]
        _clear_constants(reference, 9, reference.max_row or 9, 2, 5)


def write_reporting_funding_rows(
    wb,
    cells: Mapping[str, Any],
    *,
    ns_id_to_country: Mapping[int, Dict[str, str]],
    column_map: Mapping[str, str],
) -> None:
    ws = wb[REPORTING_FUNDING_SHEET]
    headers = _header_map(ws, 14, 14)
    country_col = headers.get("country")
    region_col = headers.get("region")
    excel_col = {}
    for excel_name, form_name in column_map.items():
        col = headers.get(_norm(excel_name))
        if col:
            excel_col[form_name] = col
    grouped: Dict[int, Dict[str, Any]] = {}
    for key, raw in cells.items():
        if not isinstance(key, str) or "_" not in key or key.startswith("_"):
            continue
        row_id, column = key.split("_", 1)
        if column not in excel_col:
            continue
        try:
            ns_id = int(row_id)
        except ValueError:
            continue
        grouped.setdefault(ns_id, {})[column] = raw
    excel_row = 15
    for ns_id, values in grouped.items():
        if excel_row > 70:
            break
        country = ns_id_to_country.get(ns_id) or {}
        if country_col and country.get("country"):
            ws.cell(excel_row, country_col).value = country["country"]
        if region_col and country.get("region"):
            ws.cell(excel_row, region_col).value = country["region"]
        for column, value in values.items():
            number = _as_number(value)
            ws.cell(excel_row, excel_col[column]).value = value if number is None else _json_number(number)
        excel_row += 1


def selected_national_society(wb, *, planning: bool) -> str:
    sheet = PLANNING_FUNDING_SHEET if planning else REPORTING_FUNDING_SHEET
    cell = PLANNING_NS_CELL if planning else REPORTING_NS_CELL
    if sheet not in wb.sheetnames:
        return ""
    return str(wb[sheet][cell].value or "").strip()


def workbook_has_planning_structure(wb) -> bool:
    if PLANNING_FUNDING_SHEET not in wb.sheetnames or PLANNING_DATA_SHEET not in wb.sheetnames:
        return False
    return PLANNING_DATA_TABLE in wb[PLANNING_DATA_SHEET].tables


def workbook_has_reporting_structure(wb) -> bool:
    if REPORTING_FUNDING_SHEET not in wb.sheetnames:
        return False
    return REPORTING_DATA_TABLE in wb[REPORTING_FUNDING_SHEET].tables


def names_match(left: str, right: str) -> bool:
    return bool(_norm(left)) and _norm(left) == _norm(right)
