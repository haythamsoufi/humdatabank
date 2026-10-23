"""Formatted Excel export of the UPR emergencies catalogue."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

NAVY = "011E41"
WHITE = "FFFFFF"
BAND = "F3F4F6"
GROUP = "FFF7ED"
GROUP_INK = "7C2D12"
INK = "111827"
MUTED = "6B7280"
LINE = "E5E7EB"
GREEN = "F0FDF4"
GREEN_INK = "166534"
AMBER = "FFF7ED"
AMBER_INK = "9A3412"
BLUE = "EFF6FF"
BLUE_INK = "1E40AF"
ZEBRA = "F9FAFB"

STATUS_LABELS = {
    "covered": "Covered",
    "missing": "Not covered",
    "not_applicable": "Not applicable",
    "no_longer": "No longer applicable",
    "not_entered": "Not reported",
    "partial": "Partial",
    "entered": "Entered",
    "not_annualised": "Not annualised",
}

_THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def build_emergencies_workbook(payload: dict[str, Any], *, round_label: str = "") -> bytes:
    """Workbook matching the emergencies table: singles, then grouped appeals."""
    sections = [str(name) for name in (payload.get("sections") or []) if name]
    rows = [row for row in (payload.get("rows") or []) if isinstance(row, dict)]
    columns = _columns(sections, plan=_show_plan_columns(payload, rows))
    organised = _organise(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Emergencies"

    title = "UPR emergencies"
    if round_label:
        title = f"{title} — {round_label}"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(name="Calibri", bold=True, size=16, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    note = sheet.cell(
        2,
        1,
        "ERP is the appeal record: name, code, duration, and Total Multiyear Requested. Databank is what focal points entered on the assignment. Annualised funding compares the two.",
    )
    note.font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    note.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    sheet.row_dimensions[2].height = 20

    header_row = 4
    source_row = 5
    for index, column in enumerate(columns, 1):
        header = sheet.cell(header_row, index, column["label"])
        header.font = Font(name="Calibri", bold=True, size=11, color=WHITE)
        header.fill = PatternFill("solid", fgColor=NAVY)
        header.alignment = Alignment(vertical="center", wrap_text=True)
        source = sheet.cell(source_row, index, column["source"])
        source.font = Font(name="Calibri", bold=True, size=9, color=_source_ink(column["source"]))
        source.fill = PatternFill("solid", fgColor=_source_fill(column["source"]))
        source.alignment = Alignment(horizontal="center", vertical="center")
        header.border = _THIN
        source.border = _THIN
    sheet.row_dimensions[header_row].height = 22
    sheet.row_dimensions[source_row].height = 18
    sheet.freeze_panes = "A6"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:5"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 110

    cursor = 6
    if organised["singles"]:
        cursor = _band(
            sheet,
            cursor,
            len(columns),
            "Single appeals  ·  " + _count_phrase(len(organised["singles"]), "appeal", "appeals"),
        )
        for row in organised["singles"]:
            cursor = _data_row(sheet, cursor, columns, row, child=False)
    if organised["groups"]:
        grouped_appeals = sum(len(group["parents"]) + len(group["children"]) for group in organised["groups"])
        cursor = _band(
            sheet,
            cursor,
            len(columns),
            "Grouped appeals  ·  "
            + _count_phrase(len(organised["groups"]), "group", "groups")
            + "  ·  "
            + _count_phrase(grouped_appeals, "appeal", "appeals"),
        )
        for group in organised["groups"]:
            cursor = _group_row(sheet, cursor, columns, group)
            for row in group["parents"] + group["children"]:
                cursor = _data_row(sheet, cursor, columns, row, child=True)

    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = column["width"]

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(round_label: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", round_label or "round").strip("-").lower()
    return f"upr-emergencies-{slug or 'round'}.xlsx"


def _show_plan_columns(payload: dict[str, Any], rows: list[dict[str, Any]]) -> bool:
    if payload.get("plan_columns"):
        return True
    return any(row.get("funding_check") or row.get("duration_years") or row.get("requested") for row in rows)


def _section_rank(name: str) -> int:
    text = name.lower()
    if "people to be reached" in text or "emergency appeal" in text:
        return 0
    if "funding requirement" in text:
        return 1
    return 2


def _columns(sections: list[str], *, plan: bool) -> list[dict[str, str]]:
    columns = [
        {"key": "country", "label": "Country", "source": "Databank", "width": 24},
        {"key": "emergency", "label": "Emergency", "source": "ERP", "width": 42},
        {"key": "code", "label": "Code", "source": "ERP", "width": 16},
        {"key": "group", "label": "Grouped appeal", "source": "ERP", "width": 32},
        {"key": "coverage", "label": "Coverage", "source": "Databank", "width": 18},
    ]
    for section in sorted(sections, key=_section_rank):
        columns.append({"key": f"section:{section}", "label": section, "source": "Databank", "width": 28})
    if plan:
        columns.append({"key": "requested", "label": "Total Multiyear Requested", "source": "ERP", "width": 28})
        columns.append({"key": "duration", "label": "Duration", "source": "ERP", "width": 22})
        columns.append({"key": "funding", "label": "Annualised funding", "source": "Databank · ERP", "width": 22})
    return columns


def _organise(rows: list[dict[str, Any]]) -> dict[str, Any]:
    parent_names: dict[str, str] = {}
    for row in rows:
        parent = str(row.get("group_code") or row.get("part_of") or "").strip()
        code = str(row.get("code") or "").strip()
        if parent and code != parent:
            parent_names[parent] = str(row.get("group_name") or parent_names.get(parent) or parent)
        if row.get("appeal_kind") == "group" and code:
            parent_names[code] = str(row.get("group_name") or row.get("name") or parent_names.get(code) or code)
    singles = []
    groups: dict[str, dict[str, Any]] = {}
    for row in rows:
        parent = str(row.get("group_code") or row.get("part_of") or "").strip()
        code = str(row.get("code") or "").strip()
        is_parent = row.get("appeal_kind") == "group" or (code in parent_names and code != parent)
        if not parent and not is_parent:
            singles.append(row)
            continue
        key = code if is_parent else parent
        group = groups.setdefault(key, {"code": key, "name": parent_names.get(key) or key, "parents": [], "children": []})
        if row.get("group_name"):
            group["name"] = row["group_name"]
        elif is_parent and row.get("name"):
            group["name"] = row["name"]
        if is_parent:
            group["parents"].append(row)
        else:
            group["children"].append(row)
    ordered = sorted(groups.values(), key=lambda group: str(group["name"]).lower())
    return {"singles": singles, "groups": ordered}


def _count_phrase(count: int, singular: str, plural: str) -> str:
    return f"{count} {singular if count == 1 else plural}"


def _band(sheet, row_index: int, width: int, label: str) -> int:
    sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=width)
    cell = sheet.cell(row_index, 1, label.upper())
    cell.font = Font(name="Calibri", bold=True, size=11, color="374151")
    cell.fill = PatternFill("solid", fgColor=BAND)
    cell.alignment = Alignment(vertical="center", indent=1)
    for column in range(1, width + 1):
        sheet.cell(row_index, column).fill = PatternFill("solid", fgColor=BAND)
        sheet.cell(row_index, column).border = _THIN
    sheet.row_dimensions[row_index].height = 22
    return row_index + 1


def _group_requested(group: dict[str, Any]) -> float | None:
    for row in group["parents"] + group["children"]:
        amount = row.get("group_requested")
        if isinstance(amount, (int, float)) and amount > 0:
            return float(amount)
        if row.get("appeal_kind") == "group":
            own = row.get("requested")
            if isinstance(own, (int, float)) and own > 0:
                return float(own)
    return None


def _group_row(sheet, row_index: int, columns: list[dict[str, str]], group: dict[str, Any]) -> int:
    label = group["name"]
    requested = _group_requested(group)
    for index, column in enumerate(columns, 1):
        value = None
        if column["key"] == "emergency":
            value = label
        elif column["key"] == "code":
            value = group["code"]
        elif column["key"] == "requested":
            value = requested
        cell = sheet.cell(row_index, index, value)
        cell.font = Font(name="Calibri", bold=True, size=11, color=GROUP_INK)
        cell.fill = PatternFill("solid", fgColor=GROUP)
        cell.border = _THIN
        cell.alignment = Alignment(vertical="center")
        if column["key"] == "requested" and value is not None:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row_index].height = 22
    return row_index + 1


def _data_row(sheet, row_index: int, columns: list[dict[str, str]], row: dict[str, Any], *, child: bool) -> int:
    fill = PatternFill("solid", fgColor=ZEBRA if row_index % 2 == 0 else WHITE)
    for index, column in enumerate(columns, 1):
        value, kind = _cell_value(column["key"], row)
        cell = sheet.cell(row_index, index, value)
        cell.font = Font(name="Calibri", size=10, color=INK)
        cell.fill = fill
        cell.border = _THIN
        cell.alignment = Alignment(vertical="center", indent=2 if child and column["key"] == "emergency" else 0)
        if kind == "money" and value is not None:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if kind == "mark":
            cell.font = Font(name="Calibri", size=14, bold=True, color=_MARK_INK.get(value, INK))
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif column["key"] in {"funding"} or column["key"].startswith("section:"):
            cell.alignment = Alignment(horizontal="center", vertical="center")
            _paint_status(cell, value)
    sheet.row_dimensions[row_index].height = 18
    return row_index + 1


def _cell_value(key: str, row: dict[str, Any]) -> tuple[Any, str]:
    if key == "country":
        return row.get("country_name") or "", "text"
    if key == "emergency":
        return row.get("name") or row.get("code") or "", "text"
    if key == "code":
        return row.get("code") or "", "text"
    if key == "group":
        if row.get("appeal_kind") == "single" or not row.get("group_code"):
            return "", "text"
        name = row.get("group_name") or row.get("group_code")
        return f"{name} ({row.get('group_code')})", "text"
    if key.startswith("section:"):
        section = key.split(":", 1)[1]
        figures = row.get("figures") if isinstance(row.get("figures"), dict) else None
        mark = (row.get("marks") or {}).get(section)
        if figures is not None and section in figures and mark in {"covered", "missing"}:
            amount = figures.get(section)
            if isinstance(amount, (int, float)) and amount > 0:
                return amount, "money"
            return "Not reported", "text"
        return STATUS_LABELS.get(mark, ""), "text"
    if key == "coverage":
        return _COVERAGE_MARK.get(_coverage_state(row), ""), "mark"
    if key == "duration":
        return _duration(row), "text"
    if key == "requested":
        return row.get("requested"), "money"
    if key == "funding":
        return STATUS_LABELS.get(row.get("funding_check"), ""), "text"
    return "", "text"


def _duration(row: dict[str, Any]) -> str:
    years = row.get("duration_years")
    if not isinstance(years, int) or years <= 0:
        return ""
    label = "1 year" if years == 1 else f"{years} years"
    start = row.get("duration_from")
    end = row.get("duration_to")
    if start and end and str(start) != str(end):
        return f"{label} · {start}–{end}"
    if start:
        return f"{label} · {start}"
    return label


_COVERAGE_MARK = {
    "covered": "✓",
    "missing": "✕",
    "partial": "–",
}
_MARK_INK = {
    "✓": "16A34A",
    "✕": "DC2626",
    "–": "CA8A04",
}


def _coverage_state(row: dict[str, Any]) -> str:
    values = set((row.get("marks") or {}).values())
    if "covered" in values and "missing" in values:
        return "partial"
    return str(row.get("status") or "")


def _paint_status(cell, value: str) -> None:
    if value in {"Covered", "Entered"}:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN_INK)
    elif value in {"Not covered", "Not annualised", "Not reported", "No longer applicable"}:
        cell.fill = PatternFill("solid", fgColor=AMBER)
        cell.font = Font(name="Calibri", size=10, bold=True, color=AMBER_INK)
    elif value in {"Not applicable"}:
        cell.font = Font(name="Calibri", size=10, color=MUTED)


def _source_fill(source: str) -> str:
    if source == "GO":
        return BLUE
    if "ERP" in source and "Databank" in source:
        return "F5F3FF"
    if source == "ERP":
        return AMBER
    return GREEN


def _source_ink(source: str) -> str:
    if source == "GO":
        return BLUE_INK
    if "ERP" in source and "Databank" in source:
        return "5B21B6"
    if source == "ERP":
        return AMBER_INK
    return GREEN_INK
