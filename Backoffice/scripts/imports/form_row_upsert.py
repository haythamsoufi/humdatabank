"""Shared form_data row upsert used by FDRS and UPR Excel import pipelines.

Lives in core ``scripts/imports/`` so plugin-local scripts can import it without
depending on each other.
"""

from __future__ import annotations

import json
import logging
import re
import time
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Callable, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)


class ImportRowsCancelled(Exception):
    """Raised when an in-flight form-data import is cancelled by the user."""


# Backward-compatible alias used by FDRS callers and tests.
FdrsSyncCancelled = ImportRowsCancelled

COL_ASSIGNMENT = "assignment_entity_status_id"
COL_PUBLIC = "public_submission_id"
COL_ITEM = "item_id"
COL_ITEM_ALT = "form_item_id"
COL_VALUE = "value"
COL_DISAGG = "disagg_data"
COL_DATA_NA = "data_not_available"
COL_NA = "not_applicable"
COL_PREFILLED = "prefilled_value"
COL_IMPUTED = "imputed_value"
COL_SUBMITTED = "submitted_at"

ALL_COLUMNS = (
    COL_ASSIGNMENT,
    COL_PUBLIC,
    COL_ITEM,
    COL_VALUE,
    COL_DISAGG,
    COL_DATA_NA,
    COL_NA,
    COL_PREFILLED,
    COL_IMPUTED,
    COL_SUBMITTED,
)


def _normalize_headers(row: Dict[str, str]) -> Dict[str, str]:
    """Normalize keys to strip whitespace and match expected names."""
    return {k.strip(): v for k, v in row.items()}


def _parse_submitted_at(raw: Optional[str]) -> Optional[datetime]:
    """Parse submitted_at from Power Query style (e.g. 17/09/2025  14:56:03)."""
    if not raw or not str(raw).strip():
        return None
    raw = str(raw).strip()
    for fmt in ("%d/%m/%Y  %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _parse_json_field(raw: Optional[str]) -> Optional[Any]:
    """Parse a structured JSON field such as disagg_data."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    s = str(raw).strip()
    if s.lower() in ("null", "none", ""):
        return None
    try:
        return json.loads(s)
    except json.JSONDecodeError:
        return None


def _parse_scalar_field(raw: Optional[str]) -> Optional[str]:
    """Parse scalar auxiliary values into the same text shape as form_data.value."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    s = str(raw).strip()
    if s.lower() in ("null", "none", "undefined", ""):
        return None
    try:
        parsed = json.loads(s)
    except json.JSONDecodeError:
        return s

    parsed = _normalize_json_numbers_to_ints(parsed)
    if parsed is None:
        return None
    if isinstance(parsed, dict):
        return None
    if isinstance(parsed, list):
        return json.dumps(parsed, ensure_ascii=False)
    if isinstance(parsed, bool):
        return "true" if parsed else "false"
    return str(parsed)


def _disagg_data_for_db(val: Any) -> Optional[Any]:
    """Return value to store in disagg_data column; use None (DB NULL) when empty or 'null'."""
    if val is None:
        return None
    if isinstance(val, str):
        if not val.strip() or val.strip().lower() == "null":
            return None
        return val
    if isinstance(val, dict) and not val:
        return None
    return val


def _parse_bool(raw: Optional[str]) -> bool:
    """Parse boolean (data_not_available, not_applicable). Defaults to False when absent."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return False
    s = str(raw).strip().lower()
    if s in ("1", "true", "yes", "on"):
        return True
    if s in ("0", "false", "no", "off"):
        return False
    return False


_THOUSANDS_GROUPING_RE = re.compile(r"^\s*[-+]?\d{1,3}(,\d{3})+(\.\d+)?\s*$")


def _normalize_numeric_string_for_parse(s: str) -> str:
    t = (s or "").strip()
    if not t:
        return t
    if _THOUSANDS_GROUPING_RE.match(t) or ("," in t and "." in t):
        return t.replace(",", "")
    return t


def _to_whole_number_int(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, bool):
        return val
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        try:
            d = Decimal(str(val))
            if d.is_finite():
                return int(d.to_integral_value(rounding=ROUND_HALF_UP))
        except Exception as e:
            logger.debug("_parse_value_int fallback: %s", e)
            return int(Decimal(val).to_integral_value(rounding=ROUND_HALF_UP))
        return val
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return val
        s = _normalize_numeric_string_for_parse(s)
        try:
            d = Decimal(s)
        except InvalidOperation:
            return val
        if not d.is_finite():
            return val
        return int(d.to_integral_value(rounding=ROUND_HALF_UP))
    return val


def _normalize_json_numbers_to_ints(obj: Any) -> Any:
    if obj is None:
        return None
    if isinstance(obj, (bool, int)):
        return obj
    if isinstance(obj, (float, str)):
        return _to_whole_number_int(obj)
    if isinstance(obj, list):
        return [_normalize_json_numbers_to_ints(x) for x in obj]
    if isinstance(obj, dict):
        return {k: _normalize_json_numbers_to_ints(v) for k, v in obj.items()}
    return obj


def _excel_cell_value(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, bool):
        return val
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        if val != val or val in (float("inf"), float("-inf")):
            return val
        if val == int(val):
            return int(val)
        return val
    if isinstance(val, Decimal):
        if not val.is_finite():
            return val
        try:
            as_int = int(val.to_integral_value(rounding=ROUND_HALF_UP))
            if Decimal(as_int) == val:
                return as_int
        except Exception as e:
            logger.debug("_excel_cell_value decimal: %s", e)
        return float(val)
    if isinstance(val, str):
        s = val.strip()
        if not s or s.lower() in ("null", "none"):
            return ""
        s_norm = _normalize_numeric_string_for_parse(s)
        try:
            d = Decimal(s_norm)
        except InvalidOperation:
            return val
        if not d.is_finite():
            return val
        try:
            as_int = int(d.to_integral_value(rounding=ROUND_HALF_UP))
            if Decimal(as_int) == d:
                return as_int
        except Exception as e:
            logger.debug("_excel_cell_value str decimal: %s", e)
        return val
    return val


def _coerce_value(raw: Any) -> Optional[str]:
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s.lower() in ("null", "none"):
        return None

    s_norm = _normalize_numeric_string_for_parse(s)
    try:
        d = Decimal(s_norm)
        if d.is_finite():
            s = str(int(d.to_integral_value(rounding=ROUND_HALF_UP)))
        else:
            return None
    except InvalidOperation:
        pass

    if len(s) > 255:
        return s[:255]
    return s if s else None


def write_rows_to_excel(
    rows: List[Dict[str, str]],
    path: str,
    columns: Optional[Tuple[str, ...]] = None,
    extra_sheets: Optional[List[Tuple[str, List[Dict[str, Any]], Optional[Tuple[str, ...]]]]] = None,
    sheet_title: str = "Ready to import",
) -> None:
    """Write list of row dicts to Excel; first row = headers. Optionally add extra sheets."""
    try:
        import openpyxl
    except ImportError:
        import sys
        sys.exit("Excel support requires openpyxl: pip install openpyxl")
    cols = columns or (ALL_COLUMNS if rows else ())
    if rows and not cols:
        cols = tuple(rows[0].keys())
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = (sheet_title or "Ready to import")[:31]
    for c, key in enumerate(cols, 1):
        sheet.cell(row=1, column=c, value=key)
    for r, row in enumerate(rows, 2):
        for c, key in enumerate(cols, 1):
            val = row.get(key, "")
            if isinstance(val, (dict, list)):
                val = json.dumps(_normalize_json_numbers_to_ints(val))
            else:
                val = _excel_cell_value(val)
            sheet.cell(row=r, column=c, value=val)
    for sheet_name, extra_rows, extra_cols in extra_sheets or []:
        if not extra_rows and not extra_cols:
            continue
        ws = wb.create_sheet(title=sheet_name[:31])
        ec = extra_cols or (tuple(extra_rows[0].keys()) if extra_rows else ())
        for c, key in enumerate(ec, 1):
            ws.cell(row=1, column=c, value=key)
        for r, row in enumerate(extra_rows, 2):
            for c, key in enumerate(ec, 1):
                val = row.get(key, "")
                if isinstance(val, (dict, list)):
                    val = json.dumps(_normalize_json_numbers_to_ints(val))
                else:
                    val = _excel_cell_value(val)
                ws.cell(row=r, column=c, value=val)
    wb.save(path)


def row_to_payload(row: Dict[str, str]) -> Tuple[Optional[int], Optional[int], Optional[int], Dict[str, Any]]:
    aes_raw = row.get(COL_ASSIGNMENT)
    item_raw = row.get(COL_ITEM) or row.get(COL_ITEM_ALT)
    if not aes_raw and not row.get(COL_PUBLIC):
        return None, None, None, {}
    try:
        assignment_entity_status_id = int(aes_raw) if aes_raw else None
    except (ValueError, TypeError):
        assignment_entity_status_id = None
    try:
        form_item_id = int(item_raw) if item_raw else None
    except (ValueError, TypeError):
        form_item_id = None
    public_raw = row.get(COL_PUBLIC)
    try:
        public_submission_id = int(public_raw) if public_raw else None
    except (ValueError, TypeError):
        public_submission_id = None

    if not form_item_id:
        return None, None, None, {}
    if not assignment_entity_status_id and not public_submission_id:
        return None, None, None, {}

    disagg_data = _parse_json_field(row.get(COL_DISAGG))
    prefilled_value = _parse_scalar_field(row.get(COL_PREFILLED))
    imputed_value = _parse_scalar_field(row.get(COL_IMPUTED))

    payload = {
        "value": _coerce_value(row.get(COL_VALUE)),
        "disagg_data": _normalize_json_numbers_to_ints(disagg_data),
        "data_not_available": _parse_bool(row.get(COL_DATA_NA)),
        "not_applicable": _parse_bool(row.get(COL_NA)),
        "prefilled_value": prefilled_value,
        "imputed_value": imputed_value,
        "submitted_at": _parse_submitted_at(row.get(COL_SUBMITTED)),
        "disagg_type": (row.get("_debug_disagg_type") or "").strip() or None,
    }
    return assignment_entity_status_id, public_submission_id, form_item_id, payload


def flask_app_for_import():
    """Reuse the running Flask app when a worker already pushed a context."""
    from flask import current_app, has_app_context

    if has_app_context():
        return current_app._get_current_object()
    from app import create_app

    return create_app()


def _commit_upsert_and_yield() -> None:
    from app.extensions import db

    db.session.commit()
    db.session.expire_all()
    time.sleep(0.05)


def upsert_form_data_rows(
    rows: List[Dict[str, str]],
    *,
    dry_run: bool = False,
    batch_size: int = 1000,
    valid_form_item_ids: Optional[set] = None,
    valid_assignment_entity_status_ids: Optional[set] = None,
    progress_cb: Optional[Callable[[Dict[str, Any]], None]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    progress_start_pct: float = 20.0,
    progress_end_pct: float = 100.0,
    stats: Optional[Dict[str, int]] = None,
    change_recorder: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, int]:
    """Upsert ready-to-import rows into form_data (shared by FDRS and UPR Excel pipelines)."""
    from sqlalchemy import tuple_

    from app.extensions import db
    from app.models.assignments import AssignmentEntityStatus
    from app.models.forms import FormData

    if stats is None:
        stats = {"loaded": 0, "skipped": 0, "inserted": 0, "updated": 0, "errors": 0}
    stats["loaded"] = len(rows)

    def _check_cancel() -> None:
        if cancel_check and cancel_check():
            try:
                db.session.rollback()
            except Exception as e:
                logger.debug("rollback on cancel failed: %s", e)
            raise ImportRowsCancelled()

    if valid_assignment_entity_status_ids is not None:
        valid_aes_ids = valid_assignment_entity_status_ids
    else:
        valid_aes_ids = set(aid for (aid,) in db.session.query(AssignmentEntityStatus.id).all())
    total_rows = len(rows)
    span = max(progress_end_pct - progress_start_pct, 0.0)

    if progress_cb:
        progress_cb({
            "stage": "upsert",
            "message": "Starting upsert...",
            "current": 0,
            "total": total_rows,
            "percent": progress_start_pct,
            "stats": dict(stats),
        })

    def _maybe_report(i: int, row: Dict[str, Any]) -> None:
        if i == 1 or i % 50 == 0 or i == total_rows:
            _check_cancel()
        if not progress_cb:
            return
        if not (i == 1 or i % 250 == 0 or i == total_rows):
            return
        pct = progress_start_pct + (span * (i / total_rows)) if total_rows else progress_end_pct
        kpi = (row.get("_debug_kpi_code") or "").strip()
        iso3 = (row.get("_debug_iso3") or "").strip()
        yr = (row.get("_debug_year") or "").strip()
        details = " ".join(b for b in (iso3, yr, kpi) if b).strip()
        msg = f"Processing {i}/{total_rows} ({pct:.1f}%)" + (f" - {details}" if details else "")
        progress_cb({
            "stage": "upsert",
            "message": msg,
            "current": i,
            "total": total_rows,
            "percent": pct,
            "stats": dict(stats),
        })
        time.sleep(0.02)

    prefetch_size = max(2000, int(batch_size or 1000))

    for batch_start in range(0, total_rows, prefetch_size):
        _check_cancel()
        batch = rows[batch_start: batch_start + prefetch_size]

        parsed_batch = []
        aes_pairs_set = set()
        pub_pairs_set = set()
        for r in batch:
            try:
                p = row_to_payload(r)
            except Exception as e:
                logger.debug("row_to_payload failed: %s", e)
                p = (None, None, None, {})
            parsed_batch.append(p)
            aes_id, pub_id, item_id, _ = p
            if not item_id:
                continue
            if aes_id:
                aes_pairs_set.add((int(aes_id), int(item_id)))
            elif pub_id:
                pub_pairs_set.add((int(pub_id), int(item_id)))

        existing_by_aes: Dict[Tuple[int, int], FormData] = {}
        existing_by_pub: Dict[Tuple[int, int], FormData] = {}
        if aes_pairs_set:
            q = FormData.query.filter(
                tuple_(FormData.assignment_entity_status_id, FormData.form_item_id).in_(list(aes_pairs_set))
            )
            for fd in q.all():
                key = (int(fd.assignment_entity_status_id), int(fd.form_item_id))
                existing_by_aes.setdefault(key, fd)
        if pub_pairs_set:
            q = FormData.query.filter(
                tuple_(FormData.public_submission_id, FormData.form_item_id).in_(list(pub_pairs_set))
            )
            for fd in q.all():
                key = (int(fd.public_submission_id), int(fd.form_item_id))
                existing_by_pub.setdefault(key, fd)

        for j_rel, (row, (assignment_entity_status_id, public_submission_id, form_item_id, payload)) in enumerate(
            zip(batch, parsed_batch)
        ):
            j = batch_start + j_rel + 1
            if not form_item_id or (not assignment_entity_status_id and not public_submission_id):
                stats["skipped"] += 1
                _maybe_report(j, row)
                continue
            if valid_form_item_ids is not None and form_item_id not in valid_form_item_ids:
                stats["skipped"] += 1
                _maybe_report(j, row)
                continue
            if assignment_entity_status_id and assignment_entity_status_id not in valid_aes_ids:
                stats["skipped"] += 1
                _maybe_report(j, row)
                continue

            if assignment_entity_status_id:
                existing = existing_by_aes.get((int(assignment_entity_status_id), int(form_item_id)))
            else:
                existing = existing_by_pub.get((int(public_submission_id), int(form_item_id)))

            if change_recorder:
                try:
                    change_recorder({
                        "op": "update" if existing else "insert",
                        "aes_id": assignment_entity_status_id,
                        "item_id": form_item_id,
                        "iso3": (row.get("_debug_iso3") or "").strip() or None,
                        "year": (row.get("_debug_year") or "").strip() or None,
                        "kpi": (row.get("_debug_kpi_code") or "").strip() or None,
                        "old_value": getattr(existing, "value", None) if existing else None,
                        "old_disagg": getattr(existing, "disagg_data", None) if existing else None,
                        "new_value": payload.get("value"),
                        "new_disagg": payload.get("disagg_data"),
                    })
                except Exception:
                    logger.debug("import change recorder failed", exc_info=True)

            if dry_run:
                if existing:
                    stats["updated"] += 1
                else:
                    stats["inserted"] += 1
                _maybe_report(j, row)
                continue

            try:
                disagg_for_db = _disagg_data_for_db(payload["disagg_data"])
                if disagg_for_db is None:
                    disagg_for_db = db.null()
                prefilled_for_db = FormData._coerce_scalar_text_value(payload["prefilled_value"])
                imputed_for_db = FormData._coerce_scalar_text_value(payload["imputed_value"])
                disagg_type = payload.get("disagg_type")
                if not disagg_type:
                    if isinstance(payload.get("disagg_data"), dict) and payload["disagg_data"].get("mode"):
                        disagg_type = "standard_disagg"
                    elif payload.get("value") not in (None, ""):
                        disagg_type = "simple"
                    elif isinstance(payload.get("disagg_data"), dict) and payload["disagg_data"]:
                        disagg_type = "matrix"
                if existing:
                    existing.value = payload["value"]
                    existing._sync_numeric_value_from_string()
                    existing.disagg_data = disagg_for_db
                    existing.disagg_type = disagg_type
                    existing.data_not_available = payload["data_not_available"]
                    existing.not_applicable = payload["not_applicable"]
                    existing.prefilled_value = prefilled_for_db
                    FormData.sync_imputed_numeric_value(existing, imputed_for_db)
                    if payload["submitted_at"] is not None:
                        existing.submitted_at = payload["submitted_at"]
                    db.session.add(existing)
                    stats["updated"] += 1
                else:
                    entry = FormData(
                        assignment_entity_status_id=assignment_entity_status_id,
                        public_submission_id=public_submission_id,
                        form_item_id=form_item_id,
                        value=payload["value"],
                        disagg_data=disagg_for_db,
                        disagg_type=disagg_type,
                        data_not_available=payload["data_not_available"],
                        not_applicable=payload["not_applicable"],
                        prefilled_value=prefilled_for_db,
                        submitted_at=payload["submitted_at"],
                    )
                    entry._sync_numeric_value_from_string()
                    FormData.sync_imputed_numeric_value(entry, imputed_for_db)
                    db.session.add(entry)
                    stats["inserted"] += 1
                    if assignment_entity_status_id:
                        existing_by_aes[(int(assignment_entity_status_id), int(form_item_id))] = entry
                    else:
                        existing_by_pub[(int(public_submission_id), int(form_item_id))] = entry
            except Exception as e:
                stats["errors"] += 1
                if j < 5 or stats["errors"] <= 3:
                    logger.error("Row %d error: %s", j, e)

            if batch_size and ((stats["inserted"] + stats["updated"]) % batch_size == 0) and (stats["inserted"] + stats["updated"]) > 0:
                _commit_upsert_and_yield()
            _maybe_report(j, row)

    if not dry_run and (stats["inserted"] + stats["updated"]) > 0:
        _commit_upsert_and_yield()

    return stats
