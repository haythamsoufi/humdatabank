"""Durable FDRS Tools artifacts (verification + document-status Excel).

Workbooks used to live in OS temp files / ``instance/``, which App Service
wipes on deploy. These helpers write the latest workbook + a small JSON
sidecar through ``storage_service`` (filesystem locally, Azure Blob in prod).
"""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone
from typing import Any, BinaryIO, Dict, List, Optional, Union

from app.services.platform import storage_service

logger = logging.getLogger(__name__)

STORAGE_CATEGORY = "fdrs_tools"

Source = Union[str, bytes, BinaryIO]


def utc_iso_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def verify_xlsx_rel(template_id: int) -> str:
    return f"verify/{int(template_id)}/latest.xlsx"


def verify_meta_rel(template_id: int) -> str:
    return f"verify/{int(template_id)}/latest.json"


def documents_xlsx_rel() -> str:
    return "documents/latest.xlsx"


def documents_meta_rel() -> str:
    return "documents/latest.json"


def exists(rel_path: str) -> bool:
    return storage_service.exists(STORAGE_CATEGORY, rel_path)


def verify_latest_exists(template_id: int) -> bool:
    return exists(verify_xlsx_rel(template_id)) and exists(verify_meta_rel(template_id))


def documents_latest_exists() -> bool:
    return exists(documents_xlsx_rel()) and exists(documents_meta_rel())


def upload_bytes(rel_path: str, data: bytes) -> str:
    return storage_service.upload(STORAGE_CATEGORY, rel_path, data)


def download_bytes(rel_path: str) -> bytes:
    return storage_service.download(STORAGE_CATEGORY, rel_path)


def stream_xlsx(rel_path: str, filename: str):
    return storage_service.stream_response(
        STORAGE_CATEGORY,
        rel_path,
        filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
    )


def load_json(rel_path: str) -> Optional[Dict[str, Any]]:
    if not exists(rel_path):
        return None
    try:
        raw = download_bytes(rel_path)
        data = json.loads(raw.decode("utf-8"))
    except Exception:
        logger.exception("Failed to read FDRS tools artifact JSON: %s", rel_path)
        return None
    return data if isinstance(data, dict) else None


def persist_json(rel_path: str, payload: Dict[str, Any]) -> str:
    body = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")
    return upload_bytes(rel_path, body)


def persist_xlsx_and_meta(
    *,
    xlsx_rel: str,
    meta_rel: str,
    local_xlsx_path: str,
    meta: Dict[str, Any],
) -> Dict[str, Any]:
    """Upload a local workbook + sidecar JSON. Returns the stored meta."""
    with open(local_xlsx_path, "rb") as handle:
        upload_bytes(xlsx_rel, handle.read())
    stored = dict(meta)
    stored.setdefault("completed_at", utc_iso_now())
    persist_json(meta_rel, stored)
    return stored


def persist_verify_latest(
    *,
    template_id: int,
    local_xlsx_path: str,
    stats: Optional[Dict[str, Any]] = None,
    extra: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    meta: Dict[str, Any] = {
        "kind": "fdrs.sync_verify",
        "template_id": int(template_id),
        "stats": dict(stats or {}),
        "completed_at": utc_iso_now(),
    }
    if extra:
        meta.update(extra)
    return persist_xlsx_and_meta(
        xlsx_rel=verify_xlsx_rel(template_id),
        meta_rel=verify_meta_rel(template_id),
        local_xlsx_path=local_xlsx_path,
        meta=meta,
    )


def persist_documents_latest(
    *,
    local_xlsx_path: str,
    stats: Optional[Dict[str, Any]] = None,
    extra: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    meta: Dict[str, Any] = {
        "kind": "fdrs.document_status",
        "stats": dict(stats or {}),
        "completed_at": utc_iso_now(),
    }
    if extra:
        meta.update(extra)
    return persist_xlsx_and_meta(
        xlsx_rel=documents_xlsx_rel(),
        meta_rel=documents_meta_rel(),
        local_xlsx_path=local_xlsx_path,
        meta=meta,
    )


def load_verify_latest_meta(template_id: int) -> Optional[Dict[str, Any]]:
    if not verify_latest_exists(template_id):
        return None
    return load_json(verify_meta_rel(template_id))


def load_documents_latest_meta() -> Optional[Dict[str, Any]]:
    if not documents_latest_exists():
        return None
    return load_json(documents_meta_rel())


def _cell_to_json(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _as_filter_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip().lower()


def _open_workbook(source: Source):
    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        return openpyxl.load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    if hasattr(source, "read"):
        data = source.read()
        if isinstance(data, str):
            data = data.encode("utf-8")
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return openpyxl.load_workbook(source, read_only=True, data_only=True)


def read_xlsx_sheet(
    source: Source,
    *,
    sheet: str,
    aliases: Optional[Dict[str, str]] = None,
    page: int = 1,
    per_page: int = 200,
    filters: Optional[Dict[str, str]] = None,
    search: Optional[str] = None,
    search_columns: Optional[List[str]] = None,
) -> Dict[str, Any]:
    """Read one sheet as a paginated JSON table.

    *filters* is ``{column_name_lower: expected_value_lower}`` and matches
    case-insensitively (booleans become ``true``/``false``). *search* is a
    substring match across *search_columns* (or every column when omitted).
    ``per_page <= 0`` returns every matching row (capped at 50_000) so
    client-side grids can filter locally.
    """
    page = max(1, int(page or 1))
    try:
        requested = int(per_page) if per_page is not None else 200
    except (TypeError, ValueError):
        requested = 200
    return_all = requested <= 0
    per_page = 50000 if return_all else min(500, max(1, requested))
    wanted = (sheet or "").strip()
    if aliases:
        wanted = aliases.get(wanted.lower(), wanted)
    active_filters = {
        str(key).strip().lower(): str(val).strip().lower()
        for key, val in (filters or {}).items()
        if val is not None and str(val).strip() != ""
    }
    query = (search or "").strip().lower() or None
    search_cols = {c.strip().lower() for c in (search_columns or []) if c and c.strip()}

    wb = _open_workbook(source)
    try:
        sheet_names = list(wb.sheetnames)
        if wanted not in wb.sheetnames:
            wanted = sheet_names[0] if sheet_names else wanted
        ws = wb[wanted] if wanted in wb.sheetnames else None
        if ws is None:
            return {
                "sheet": wanted,
                "sheets": sheet_names,
                "columns": [],
                "rows": [],
                "page": page,
                "per_page": per_page,
                "total_rows": 0,
                "filtered_rows": 0,
            }

        rows_iter = ws.iter_rows(values_only=True)
        raw_header = next(rows_iter, None) or ()
        columns = [str(c) if c is not None else "" for c in raw_header]
        col_index = {col.lower(): i for i, col in enumerate(columns) if col}

        matched: List[Dict[str, Any]] = []
        total_rows = 0
        for raw in rows_iter:
            total_rows += 1
            skip = False
            for col_key, expected in active_filters.items():
                idx = col_index.get(col_key)
                cell = raw[idx] if idx is not None and idx < len(raw) else None
                if _as_filter_text(cell) != expected:
                    skip = True
                    break
            if skip:
                continue
            if query:
                haystack_idxs = (
                    [col_index[c] for c in search_cols if c in col_index]
                    if search_cols
                    else range(len(columns))
                )
                found = False
                for idx in haystack_idxs:
                    cell = raw[idx] if idx < len(raw) else None
                    if query in _as_filter_text(cell):
                        found = True
                        break
                if not found:
                    continue
            row: Dict[str, Any] = {}
            for i, col in enumerate(columns):
                if not col:
                    continue
                row[col] = _cell_to_json(raw[i] if i < len(raw) else None)
            matched.append(row)

        filtered_rows = len(matched)
        if return_all:
            start = 0
            page_rows = matched[:per_page]
            page = 1
        else:
            start = (page - 1) * per_page
            page_rows = matched[start:start + per_page]
        return {
            "sheet": wanted,
            "sheets": sheet_names,
            "columns": [c for c in columns if c],
            "rows": page_rows,
            "page": page,
            "per_page": per_page if not return_all else len(page_rows),
            "total_rows": total_rows,
            "filtered_rows": filtered_rows,
        }
    finally:
        wb.close()
