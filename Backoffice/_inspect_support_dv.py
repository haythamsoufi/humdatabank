import openpyxl

path = r"app/static/templates/unified_country_plan.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["Planned Bilateral Support"]
for dv in ws.data_validations.dataValidation:
    print(f"{dv.sqref}: {dv.formula1}")

# Table9 NS names count
from openpyxl.utils import range_boundaries
td = wb["TemplateData"]
tbl = td.tables["Table9"]
ref = tbl.ref if hasattr(tbl, "ref") else tbl
min_col, min_row, max_col, max_row = range_boundaries(ref)
ns_names = []
for row in range(min_row + 1, max_row + 1):
    ns_names.append(td.cell(row, min_col + 1).value)
print("Table9 NS count:", len(ns_names))
print("sample:", ns_names[:3])

wb.close()
