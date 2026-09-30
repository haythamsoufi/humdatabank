import io
import zipfile

import openpyxl
import pytest
from werkzeug.datastructures import FileStorage

from app.utils.file_parsing import (
    MAX_CSV_ROWS,
    UploadLimitError,
    load_json_upload,
    parse_csv_to_rows,
    parse_excel_to_rows,
    read_stream_capped,
)
from app.utils.safe_workbook import (
    UnsafeWorkbookError,
    WorkbookLimits,
    inspect_xlsx,
    load_workbook_safe,
    read_excel_safe,
    read_workbook_file_bytes,
    safe_iter_rows,
)


def _xlsx_bytes(rows=((1, "a"), (2, "b"))):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_with(members):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in members.items():
            zf.writestr(name, payload)
    return buf.getvalue()


def test_loads_valid_workbook_from_bytes_stream_trusted_path_and_filestorage(tmp_path):
    data = _xlsx_bytes()
    path = tmp_path / "w.xlsx"
    path.write_bytes(data)
    for source in (
        data,
        io.BytesIO(data),
        read_workbook_file_bytes(str(path)),
        FileStorage(io.BytesIO(data), filename="w.xlsx"),
    ):
        wb = load_workbook_safe(source)
        assert [tuple(r) for r in wb.active.iter_rows(values_only=True)] == [(1, "a"), (2, "b")]
        wb.close()


def test_stream_position_is_restored():
    data = _xlsx_bytes()
    stream = io.BytesIO(data)
    load_workbook_safe(stream).close()
    assert stream.tell() == 0


def test_defaults_to_read_only():
    wb = load_workbook_safe(_xlsx_bytes())
    assert wb.read_only is True
    wb.close()
    wb = load_workbook_safe(_xlsx_bytes(), read_only=False)
    assert wb.read_only is False


@pytest.mark.parametrize("payload", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_rejects_non_xlsx(payload):
    with pytest.raises(UnsafeWorkbookError):
        load_workbook_safe(payload)


def test_rejects_oversized_upload():
    limits = WorkbookLimits(max_upload_bytes=100)
    with pytest.raises(UnsafeWorkbookError, match="too large"):
        load_workbook_safe(_xlsx_bytes(), limits=limits)


def test_rejects_high_compression_ratio_zip_bomb():
    bomb = _zip_with({"xl/workbook.xml": b"<x/>", "xl/worksheets/sheet1.xml": b"0" * (8 * 1024 * 1024)})
    assert len(bomb) < 100_000
    with pytest.raises(UnsafeWorkbookError, match="compression ratio"):
        load_workbook_safe(bomb)


def test_rejects_total_uncompressed_size():
    limits = WorkbookLimits(max_uncompressed_bytes=1000)
    with pytest.raises(UnsafeWorkbookError, match="expands"):
        inspect_xlsx(_xlsx_bytes(), limits)


def test_rejects_too_many_members():
    members = {f"xl/media/f{i}.txt": b"x" for i in range(12)}
    with pytest.raises(UnsafeWorkbookError, match="too many internal"):
        inspect_xlsx(_zip_with(members), WorkbookLimits(max_members=10))


def test_rejects_too_many_sheets():
    members = {f"xl/worksheets/sheet{i}.xml": b"<worksheet/>" for i in range(5)}
    with pytest.raises(UnsafeWorkbookError, match="too many sheets"):
        inspect_xlsx(_zip_with(members), WorkbookLimits(max_sheets=3))


def test_rejects_huge_declared_dimension_without_parsing_rows():
    sheet = b'<?xml version="1.0"?><worksheet><dimension ref="A1:C5000000"/><sheetData/></worksheet>'
    with pytest.raises(UnsafeWorkbookError, match="too many rows or columns"):
        inspect_xlsx(_zip_with({"xl/worksheets/sheet1.xml": sheet}))
    wide = b'<worksheet><dimension ref="A1:ZZZZ2"/></worksheet>'
    with pytest.raises(UnsafeWorkbookError):
        inspect_xlsx(_zip_with({"xl/worksheets/sheet1.xml": wide}))


def test_normal_dimension_is_accepted():
    ok = b'<worksheet><dimension ref="A1:D200"/></worksheet>'
    inspect_xlsx(_zip_with({"xl/worksheets/sheet1.xml": ok}))


def test_encrypted_member_is_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", b"<x/>")
    raw = bytearray(buf.getvalue())
    raw[6] |= 0x01
    central = raw.rfind(b"PK\x01\x02")
    raw[central + 8] |= 0x01
    with pytest.raises(UnsafeWorkbookError, match="Encrypted"):
        inspect_xlsx(bytes(raw))


def test_safe_iter_rows_caps_rows():
    wb = load_workbook_safe(_xlsx_bytes(rows=[(i,) for i in range(20)]))
    with pytest.raises(UnsafeWorkbookError):
        list(safe_iter_rows(wb.active, WorkbookLimits(max_rows=5), values_only=True))
    wb.close()


def test_read_excel_safe_reads_and_caps_rows():
    df = read_excel_safe(_xlsx_bytes(rows=[("h1", "h2"), (1, 2), (3, 4)]))
    assert list(df.columns) == ["h1", "h2"]
    assert len(df) == 2
    many = _xlsx_bytes(rows=[("h",)] + [(i,) for i in range(50)])
    with pytest.raises(UnsafeWorkbookError):
        read_excel_safe(many, limits=WorkbookLimits(max_rows=10))
    assert len(read_excel_safe(many, nrows=5)) == 5


def test_read_excel_safe_rejects_bomb():
    bomb = _zip_with({"xl/worksheets/sheet1.xml": b"0" * (8 * 1024 * 1024)})
    with pytest.raises(UnsafeWorkbookError):
        read_excel_safe(bomb)


def test_limits_from_config_override(app):
    app.config["WORKBOOK_MAX_ROWS"] = 7
    app.config["WORKBOOK_MAX_COLS"] = "bad"
    with app.app_context():
        limits = WorkbookLimits.from_config()
    assert limits.max_rows == 7
    assert limits.max_cols == WorkbookLimits().max_cols


def test_parse_excel_to_rows_uses_safe_loader():
    file = FileStorage(io.BytesIO(_xlsx_bytes(rows=[("a", "b"), (1, 2)])), filename="x.xlsx")
    columns, rows = parse_excel_to_rows(file)
    assert columns == ["a", "b"]
    assert rows == [{"a": 1, "b": 2}]
    bomb = _zip_with({"xl/worksheets/sheet1.xml": b"0" * (8 * 1024 * 1024)})
    with pytest.raises(ValueError):
        parse_excel_to_rows(FileStorage(io.BytesIO(bomb), filename="x.xlsx"))


def test_parse_csv_handles_quoted_newlines_and_caps():
    csv_bytes = b'a,b\r\n1,"line1\nline2"\r\n'
    columns, rows = parse_csv_to_rows(FileStorage(io.BytesIO(csv_bytes), filename="x.csv"))
    assert columns == ["a", "b"]
    assert rows == [{"a": "1", "b": "line1\nline2"}]

    big = b"a\n" + b"1\n" * (MAX_CSV_ROWS + 1)
    with pytest.raises(UploadLimitError):
        parse_csv_to_rows(FileStorage(io.BytesIO(big), filename="x.csv"))


def test_read_stream_capped_and_json_upload():
    assert read_stream_capped(io.BytesIO(b"abc"), 3) == b"abc"
    with pytest.raises(UploadLimitError):
        read_stream_capped(io.BytesIO(b"abcd"), 3)
    assert load_json_upload(FileStorage(io.BytesIO(b'{"a": 1}'))) == {"a": 1}
    with pytest.raises(UploadLimitError):
        load_json_upload(FileStorage(io.BytesIO(b'{"a": "' + b"x" * 100 + b'"}')), max_bytes=10)
    with pytest.raises(ValueError):
        load_json_upload(FileStorage(io.BytesIO(b"{not json")))
    with pytest.raises(ValueError):
        load_json_upload(FileStorage(io.BytesIO(b"[" * 100000)))


_TRUSTED_WORKBOOK_READERS = {
    "app/utils/safe_workbook.py",
    "plugins/fdrs/scripts/import_fdrs_form_data.py",
    "plugins/upr/scripts/upr_country_reporting_excel_template.py",
    "plugins/upr/scripts/unified_country_plan_excel_template.py",
    "plugins/pb_progress/db_source.py",
    "plugins/pb_progress/visuals/scripts/pb_figures/data.py",
    "plugins/pb_progress/visuals/scripts/pb_figures/translations.py",
    "plugins/pb_progress/visuals/scripts/patch_footnote_placeholders.py",
}


def test_no_unguarded_workbook_loaders_in_upload_paths():
    import re
    from pathlib import Path

    root = Path(__file__).resolve().parents[3]
    pattern = re.compile(r"(?<![\w.])(?:openpyxl\.)?load_workbook\(|\bpd\.read_excel\(|\bpd\.ExcelFile\(")
    offenders = []
    for base in ("app", "plugins"):
        for path in (root / base).rglob("*.py"):
            rel = path.relative_to(root).as_posix()
            if "/tests/" in rel or rel in _TRUSTED_WORKBOOK_READERS:
                continue
            for lineno, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
                stripped = line.strip()
                if stripped.startswith("#") or "def load_workbook(" in stripped:
                    continue
                if pattern.search(line) and "load_workbook_safe(" not in line:
                    offenders.append(f"{rel}:{lineno}")
    assert not offenders, "Use app.utils.safe_workbook for untrusted workbooks: " + ", ".join(offenders)


def test_path_sources_are_rejected_by_the_loader(tmp_path):
    path = tmp_path / "w.xlsx"
    path.write_bytes(_xlsx_bytes())
    with pytest.raises(TypeError):
        load_workbook_safe(str(path))


def test_read_workbook_file_bytes_enforces_size_cap_and_missing_file(tmp_path):
    path = tmp_path / "w.xlsx"
    path.write_bytes(_xlsx_bytes())
    with pytest.raises(UnsafeWorkbookError):
        read_workbook_file_bytes(str(path), WorkbookLimits(max_upload_bytes=10))
    with pytest.raises(UnsafeWorkbookError):
        read_workbook_file_bytes(str(tmp_path / "missing.xlsx"))
