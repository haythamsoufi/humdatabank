"""Formatted workbook for the UPR emergencies catalogue."""

from io import BytesIO

from openpyxl import load_workbook

from plugins.upr.validation.emergency_excel import build_emergencies_workbook, export_filename


def test_workbook_groups_children_and_labels_each_column_source():
    content = build_emergencies_workbook(
        {
            "sections": ["People to be reached in 2027", "Funding Requirements for 2027"],
            "plan_columns": True,
            "rows": [
                {
                    "country_name": "Afghanistan",
                    "name": "Earthquake",
                    "code": "MDRAF019",
                    "appeal_kind": "single",
                    "status": "covered",
                    "marks": {
                        "People to be reached in 2027": "covered",
                        "Funding Requirements for 2027": "covered",
                    },
                    "figures": {
                        "People to be reached in 2027": 12000,
                        "Funding Requirements for 2027": 16000000,
                    },
                    "requested": 27000000,
                    "duration_years": 3,
                    "duration_from": 2025,
                    "duration_to": 2027,
                    "funding_check": "entered",
                },
                {
                    "country_name": "DR Congo",
                    "name": "Ebola outbreak",
                    "code": "MDRCD051",
                    "appeal_kind": "child",
                    "group_code": "MDRS1007",
                    "group_name": "EVD Regional Appeal",
                    "group_requested": 800000000,
                    "status": "missing",
                    "marks": {
                        "People to be reached in 2027": "missing",
                        "Funding Requirements for 2027": "covered",
                    },
                    "figures": {
                        "People to be reached in 2027": None,
                        "Funding Requirements for 2027": None,
                    },
                    "requested": 5000000,
                    "duration_years": 3,
                    "duration_from": 2026,
                    "duration_to": 2028,
                    "funding_check": "not_annualised",
                },
            ],
        },
        round_label="2027 planning",
    )

    sheet = load_workbook(BytesIO(content)).active
    headers = []
    column = 1
    while sheet.cell(4, column).value:
        headers.append(sheet.cell(4, column).value)
        column += 1
    sources = [sheet.cell(5, index + 1).value for index in range(len(headers))]

    assert sheet["A1"].value == "UPR emergencies — 2027 planning"
    assert headers[:4] == ["Country", "Emergency", "Code", "Grouped appeal"]
    assert "2027" not in headers
    assert "2028" not in headers
    assert "2029" not in headers
    assert "Total Multiyear Requested" in headers
    assert sources[headers.index("Total Multiyear Requested")] == "ERP"
    assert sources[headers.index("Emergency")] == "ERP"
    assert sources[headers.index("Code")] == "ERP"
    assert sources[headers.index("Duration")] == "ERP"
    coverage = headers.index("Coverage")
    people_index = headers.index("People to be reached in 2027")
    funding_index = headers.index("Funding Requirements for 2027")
    assert coverage < people_index < funding_index < headers.index("Total Multiyear Requested") < headers.index("Duration")
    people = headers.index("People to be reached in 2027") + 1
    funding = headers.index("Funding Requirements for 2027") + 1
    assert sheet.cell(7, people).value == 12000
    assert sheet.cell(7, funding).value == 16000000
    assert sheet.cell(7, headers.index("Coverage") + 1).value == "✓"
    assert sheet["A6"].value == "SINGLE APPEALS  ·  1 APPEAL"
    values = [sheet.cell(row, 1).value for row in range(6, sheet.max_row + 1)]
    grouped = next(value for value in values if str(value).startswith("GROUPED APPEALS"))
    assert grouped == "GROUPED APPEALS  ·  1 GROUP  ·  1 APPEAL"
    group_row = values.index(grouped) + 7
    assert sheet.cell(group_row, 2).value == "EVD Regional Appeal"
    assert sheet.cell(group_row, 3).value == "MDRS1007"
    assert sheet.cell(group_row, headers.index("Total Multiyear Requested") + 1).value == 800000000
    child_row = group_row + 1
    assert sheet.cell(child_row, headers.index("Total Multiyear Requested") + 1).value == 5000000
    assert sheet.cell(child_row, people).value == "Not reported"
    assert sheet.cell(child_row, funding).value == "Not reported"
    assert sheet.cell(child_row, headers.index("Coverage") + 1).value == "–"
    assert export_filename("2027 planning") == "upr-emergencies-2027-planning.xlsx"
