"""NSD master import: upload, preview, run, status."""
import io
import os
import re
import threading
import uuid
from pathlib import Path

from flask import current_app, render_template, request, send_file
from flask_login import current_user

from app import db
from app.routes.admin.shared import admin_permission_required
from app.utils.advanced_validation import validate_upload_extension_and_mime
from app.utils.api_helpers import get_json_safe
from app.utils.api_responses import json_bad_request, json_ok
from app.utils.file_parsing import EXCEL_EXTENSIONS

from plugins.nsd_report_forms import bp
from plugins.nsd_report_forms.plugin import NSD_ADMIN_PERMISSION
from plugins.nsd_report_forms.importer import apply as importer_apply
from plugins.nsd_report_forms.importer import jobs, layout, workbook

MAX_UPLOAD_BYTES = 40 * 1024 * 1024
_FILE_ID = re.compile(r'^[0-9a-f]{32}$')


def _upload_dir() -> Path:
    d = Path(current_app.instance_path) / 'nsd_import_uploads'
    d.mkdir(parents=True, exist_ok=True)
    return d


def _stored_path(file_id: str):
    if not file_id or not _FILE_ID.match(file_id):
        return None
    p = _upload_dir() / f'{file_id}.xlsx'
    return p if p.is_file() else None


@bp.route('/', methods=['GET'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def import_page():
    return render_template('plugins/nsd_report_forms/nsd_import/import.html', sheets=layout.SHEETS, required_sheets=layout.REQUIRED_SHEETS,
                           template_names=layout.TEMPLATE_NAMES)


@bp.route('/sample', methods=['GET'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def download_blank():
    """An empty workbook with the right sheets and column headers."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in layout.SHEETS.items():
        ws = wb.create_sheet(name)
        for i, c in enumerate(cols, 1):
            h = ws.cell(row=1, column=i, value=c.name)
            h.font = Font(bold=True, color='FFFFFF')
            h.fill = PatternFill('solid', fgColor='DC3545' if c.required else '6C757D')
            ws.cell(row=2, column=i, value=c.help).font = Font(italic=True, color='6C757D', size=9)
            ws.column_dimensions[h.column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='NSD_master_import_blank.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/upload', methods=['POST'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def upload():
    f = request.files.get('file')
    if not f or not f.filename:
        return json_bad_request('No file selected')
    valid, error_msg, _ext = validate_upload_extension_and_mime(f, EXCEL_EXTENSIONS)
    if not valid:
        return json_bad_request(error_msg or 'Invalid file type')
    data = f.read()
    if len(data) > MAX_UPLOAD_BYTES:
        return json_bad_request('File too large (max 40 MB)')
    file_id = uuid.uuid4().hex
    path = _upload_dir() / f'{file_id}.xlsx'
    path.write_bytes(data)

    book = workbook.parse(str(path))
    result = {'success': True, 'file_id': file_id, 'filename': f.filename, 'errors': book.errors,
              'warnings': book.warnings, 'sheets': book.sheets_found}
    if not book.errors:
        result['analysis'] = importer_apply.analyze(book)
    return json_ok(**result)


@bp.route('/run', methods=['POST'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def run_import():
    data = get_json_safe()
    path = _stored_path(str(data.get('file_id') or ''))
    if path is None:
        return json_bad_request('Uploaded file not found. Please upload it again.')
    if jobs.any_running():
        return json_bad_request('Another NSD import is still running.')
    overwrite = bool(data.get('overwrite_reports', False))
    user_id = current_user.id
    app = current_app._get_current_object()
    job_id = jobs.create(user_id)

    def work():
        with app.app_context():
            try:
                book = workbook.parse(str(path))
                if book.errors:
                    jobs.update(job_id, state='failed', error='; '.join(book.errors[:5]), finished=_now())
                    return
                result = importer_apply.run(
                    book, user_id, overwrite,
                    progress=lambda done, total, cur: jobs.update(job_id, done=done, total=total, current=cur),
                    cancelled=lambda: jobs.is_cancelled(job_id))
                state = 'cancelled' if result.get('cancelled') else 'done'
                jobs.update(job_id, state=state, result=result, finished=_now())
            except Exception as exc:  # report, never crash the worker silently
                current_app.logger.exception('NSD import failed')
                jobs.update(job_id, state='failed', error=f'{type(exc).__name__}: {str(exc)[:300]}', finished=_now())
            finally:
                db.session.remove()

    threading.Thread(target=work, daemon=True, name=f'nsd-import-{job_id[:8]}').start()
    return json_ok(success=True, job_id=job_id)


def _now():
    import time
    return time.time()


@bp.route('/status/<job_id>', methods=['GET'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def status(job_id):
    job = jobs.get(job_id)
    if not job:
        return json_bad_request('Unknown job')
    job.pop('user_id', None)
    return json_ok(success=True, job=job)


@bp.route('/cancel/<job_id>', methods=['POST'])
@admin_permission_required(NSD_ADMIN_PERMISSION)
def cancel(job_id):
    if not jobs.get(job_id):
        return json_bad_request('Unknown job')
    jobs.request_cancel(job_id)
    return json_ok(success=True)
